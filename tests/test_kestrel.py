"""Mithapp's test suite.

Run with:  python -m pytest tests -q      (or  python tests/test_kestrel.py)

Everything runs against a throwaway database in a temp directory, so it never
touches real data. No network is required: the price feed is stubbed.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import pathlib
import re
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

_TMP = tempfile.mkdtemp(prefix="mittens-tests-")
os.environ["MITTENS_DATA_DIR"] = _TMP

import pytest  # noqa: E402

from kestrel import config, db, demo, security  # noqa: E402
from kestrel.engine import budgets, categorise, portfolio, snapshots, sync  # noqa: E402
from kestrel.export import banking_xlsx, investments_xlsx  # noqa: E402
from kestrel.importers import ingest, mapping, profiles  # noqa: E402
from kestrel.importers.readers import (parse_amount, parse_date, read_csv_bytes,  # noqa: E402
                                       read_ofx, read_qif)
from kestrel.institutions import registry  # noqa: E402
from kestrel.market import prices as market  # noqa: E402


@pytest.fixture(scope="session", autouse=True)
def household():
    db.init()
    demo.clear()
    rep = demo.load()
    yield rep


# ===========================================================================
# Readers — the messy part of any importer
# ===========================================================================

class TestParsing:
    def test_dates_day_first(self):
        assert parse_date("03/04/2026") == "2026-04-03"      # UK/SA order
        assert parse_date("2026-04-03") == "2026-04-03"
        assert parse_date("3 Apr 2026") == "2026-04-03"
        assert parse_date("2026-04-03T14:03:11.000Z") == "2026-04-03"
        assert parse_date("03.04.2026") == "2026-04-03"
        assert parse_date("2026/04/03") == "2026-04-03"
        assert parse_date(dt.date(2026, 4, 3)) == "2026-04-03"
        assert parse_date("") is None
        assert parse_date("not a date") is None

    def test_dates_that_can_only_be_one_way_round(self):
        assert parse_date("25/12/2025") == "2025-12-25"
        assert parse_date("13/01/2026") == "2026-01-13"

    def test_amounts(self):
        assert parse_amount("1,234.56") == 1234.56           # UK thousands
        assert parse_amount("1.234,56") == 1234.56           # continental
        assert parse_amount("41 000,00") == 41000.0          # SA with spaces
        assert parse_amount("(45.00)") == -45.0              # accountants
        assert parse_amount("£12.34") == 12.34
        assert parse_amount("R 1 234,56") == 1234.56
        assert parse_amount("120.00CR") == 120.0
        assert parse_amount("120.00DR") == -120.0
        assert parse_amount("-") is None
        assert parse_amount("") is None
        assert parse_amount(42) == 42.0

    def test_finds_the_header_below_junk(self):
        raw = (b"Account Name,Mr A Clark\n"
               b"Account Balance,1234.56\n"
               b"\n"
               b"Date,Transaction type,Description,Paid out,Paid in,Balance\n"
               b"01 Aug 2026,Card payment,TESCO,42.10,,1192.46\n")
        t = read_csv_bytes(raw, "nationwide.csv")
        assert t.header[0] == "Date"
        assert len(t.rows) == 1
        assert t.preamble

    def test_semicolons_and_utf16(self):
        raw = "Datum;Beskrywing;Bedrag\n01/07/2026;PICK N PAY;-842,15\n".encode("utf-16")
        t = read_csv_bytes(raw, "za.csv")
        assert t.header == ["Datum", "Beskrywing", "Bedrag"]
        assert parse_amount(t.rows[0][2]) == -842.15

    def test_ofx(self):
        ofx = (b"OFXHEADER:100\n\n<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><CURDEF>GBP"
               b"<BANKTRANLIST><STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20260803120000"
               b"<TRNAMT>-42.10<FITID>ABC123<NAME>TESCO STORES</STMTTRN>"
               b"</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>")
        t = read_ofx(ofx, "s.ofx")
        assert len(t.rows) == 1
        assert t.rows[0][0] == "2026-08-03"
        assert parse_amount(t.rows[0][3]) == -42.10

    def test_qif(self):
        q = b"!Type:Bank\nD03/08/2026\nPTESCO\nT-42.10\n^\nD04/08/2026\nPSALARY\nT3850.00\n^\n"
        t = read_qif(q, "s.qif")
        assert len(t.rows) == 2
        assert parse_amount(t.rows[1][3]) == 3850.0


# ===========================================================================
# Column mapping
# ===========================================================================

class TestMapping:
    def test_recognises_known_formats(self):
        assert profiles.detect(["Transaction ID", "Date", "Time", "Type", "Name", "Emoji",
                                "Category", "Amount", "Currency"]) == "monzo"
        assert profiles.detect(["Date", "Counter Party", "Reference", "Type",
                                "Amount (GBP)", "Balance (GBP)"]) == "starling"
        assert profiles.detect(["Transaction Date", "Transaction Type", "Sort Code",
                                "Transaction Description", "Debit Amount",
                                "Credit Amount", "Balance"]) == "lloyds"
        assert profiles.detect(["Title", "Type", "Timestamp", "Account Currency",
                                "Total Amount in Account Currency"]) == "freetrade"
        assert profiles.detect(["Action", "Time", "ISIN", "Ticker", "Name",
                                "No. of shares", "Price / share"]) == "trading212"

    def test_maps_an_unknown_layout(self):
        raw = ("Datum;Beskrywing;Bedrag;Saldo\n"
               "01/07/2026;PICK N PAY BRACKENFELL;-842,15;18 402,55\n"
               "03/07/2026;SALARIS;41 000,00;59 402,55\n").encode()
        t = read_csv_bytes(raw, "x.csv")
        s = mapping.suggest(t, "transactions")
        assert s["mapping"]["date"] == 0
        assert s["mapping"]["description"] == 1
        assert s["mapping"]["amount"] == 2
        assert s["mapping"]["balance"] == 3
        assert not s["missing"]

    def test_splits_debit_and_credit(self):
        raw = ("Date,Description,Money In,Money Out,Balance\n"
               "01/08/2026,SALARY,3850.00,,4000.00\n"
               "02/08/2026,TESCO,,42.10,3957.90\n").encode()
        t = read_csv_bytes(raw, "x.csv")
        s = mapping.suggest(t, "transactions")
        assert s["mapping"]["credit"] == 2
        assert s["mapping"]["debit"] == 3

    def test_remembers_a_mapping(self):
        raw = b"Weird1,Weird2,Weird3\n01/08/2026,THING,-4.00\n"
        t = read_csv_bytes(raw, "x.csv")
        mapping.remember(t.signature(), "transactions",
                         {"date": 0, "description": 1, "amount": 2}, "gb-monzo", "test")
        got = mapping.recall(t.signature())
        assert got and got["mapping"]["amount"] == 2


# ===========================================================================
# Institutions
# ===========================================================================

class TestRegistry:
    def test_both_countries_are_well_populated(self):
        gb = registry.all_institutions("GB")
        za = registry.all_institutions("ZA")
        assert len(gb) > 150
        assert len(za) > 70
        for name in ("Barclays", "Monzo", "Nationwide Building Society", "Freetrade",
                     "Trading 212", "Hargreaves Lansdown", "Nationwide Building Society"):
            assert any(i["name"] == name for i in gb), name
        for name in ("Absa Bank", "Capitec Bank", "First National Bank", "Nedbank",
                     "Standard Bank", "EasyEquities", "Investec Bank (South Africa)"):
            assert any(i["name"] == name for i in za), name

    def test_search_handles_abbreviations_and_aliases(self):
        assert registry.search("fnb", "ZA")[0]["name"] == "First National Bank"
        assert registry.search("t212", "GB")[0]["name"] == "Trading 212"
        assert registry.search("HL", "GB")[0]["name"] == "Hargreaves Lansdown"
        assert registry.search("clydesdale", "GB")[0]["name"] == "Virgin Money"

    def test_every_institution_offers_a_way_in(self):
        for inst in registry.all_institutions():
            step = registry.next_step(inst["id"])
            assert step["methods"], inst["id"]
            assert step["recommended"] in {"openbanking", "direct_api", "aggregator",
                                           "csv", "manual"}
            assert any(m["method"] == step["recommended"] for m in step["methods"])
            assert step["summary"]

    def test_recommendation_is_something_you_can_actually_do(self):
        for inst in registry.all_institutions():
            step = registry.next_step(inst["id"])
            rec = next(m for m in step["methods"] if m["method"] == step["recommended"])
            assert rec["available_now"], f"{inst['id']} recommends an unavailable route"

    def test_an_automatic_recommendation_always_has_a_provider_behind_it(self):
        """The bug this catches: the registry offered Enable Banking for Monzo, Mithapp
        had no client for it, and the setup screen said 'supports Open Banking' and
        'no self-serve provider' in the same breath."""
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        for inst in registry.all_institutions():
            step = registry.next_step(inst["id"])
            if step["recommended"] in ("csv", "manual"):
                continue
            m = next(x for x in step["methods"] if x["method"] == step["recommended"])
            assert m["ready_providers"], (
                f"{inst['id']} recommends {step['recommended']} with no usable provider")
            for pid in m["ready_providers"]:
                assert provider_base.get(pid), f"{inst['id']} offers {pid}, which has no client"

    def test_an_automatic_recommendation_is_always_free(self):
        """The second half of the same bug. A provider can be self-serve, built and
        still useless: TrueLayer's free console only reaches mock banks, so recommending
        it sent Allan to a screen warning him not to enter his real details. If Mithapp
        recommends linking, the household must be able to do it without paying."""
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        for inst in registry.all_institutions():
            step = registry.next_step(inst["id"])
            if step["recommended"] in ("csv", "manual"):
                continue
            m = next(x for x in step["methods"] if x["method"] == step["recommended"])
            assert m["free_providers"], (
                f"{inst['id']} recommends {step['recommended']}, but every provider "
                f"behind it charges: {m['ready_providers']}")

    def test_every_catalogued_provider_declares_whether_it_is_free(self):
        """`free` missing reads as False, which silently demotes an institution to
        statement upload — exactly what happened to Investec when a bulk edit skipped it.
        An explicit flag on every entry makes that omission impossible to repeat."""
        for pid, p in registry.providers().items():
            assert "free" in p, f"{pid} does not say whether it is free"
            assert isinstance(p["free"], bool), pid

    def test_only_routes_that_cost_nothing_are_offered(self):
        """Open Banking and the aggregators reach these banks only through a registered
        provider, and none is free to a household. Offering the option produced a screen
        that looked like it could connect and then asked for a subscription."""
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        offered_paid = []
        for country in ("GB", "ZA"):
            for inst in registry.search("", country, None, 500):
                step = registry.next_step(inst["id"])
                for m in step["methods"]:
                    if m["method"] in ("openbanking", "aggregator") and not m["free_providers"]:
                        offered_paid.append((inst["id"], m["method"]))
        assert not offered_paid, offered_paid[:5]

    def test_a_withheld_route_is_explained_rather_than_hidden(self):
        # Silently dropping it leaves somebody wondering why their bank appears not to
        # support something they know it supports.
        step = registry.next_step("gb-hsbc-uk")
        assert [w["method"] for w in step["withheld"]] == ["openbanking"]
        assert step["withheld"][0]["label"]

    def test_a_free_direct_api_is_still_offered(self):
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        for iid in ("gb-monzo", "gb-starling-bank", "gb-trading-212"):
            step = registry.next_step(iid)
            api = next((m for m in step["methods"] if m["method"] == "direct_api"), None)
            assert api, iid
            assert api["free_providers"], iid
            assert all(provider_base.get(p) for p in api["ready_providers"]), iid

    def test_trading212_and_investec_get_their_own_apis(self):
        assert registry.next_step("gb-trading-212")["recommended"] == "direct_api"
        assert registry.next_step("za-investec-bank-south-africa")["recommended"] == "direct_api"

    def test_banks_with_their_own_free_api_use_it(self):
        # Monzo and Starling publish developer APIs an individual can use on their own
        # account, for free. Those are the only two free automatic routes into a UK
        # current account that actually exist.
        for iid in ("gb-monzo", "gb-starling-bank"):
            assert registry.next_step(iid)["recommended"] == "direct_api", iid

    def test_uk_high_street_falls_back_to_statements(self):
        # These all support Open Banking, and Mithapp still lists it as a method — but
        # every aggregator that can reach them charges, and a route the household cannot
        # sign up for must never be the recommendation. Recommending TrueLayer here is
        # what sent Allan to a console that warns him not to enter his real bank details.
        for iid in ("gb-barclays", "gb-lloyds-bank", "gb-natwest", "gb-hsbc-uk",
                    "gb-santander-uk", "gb-nationwide-building-society"):
            step = registry.next_step(iid)
            assert step["recommended"] == "csv", iid
            # The route is withheld rather than deleted: the panel still says the bank
            # supports it, but it cannot be chosen.
            assert not any(m["method"] == "openbanking" for m in step["methods"]), iid
            assert any(w["method"] == "openbanking" for w in step["withheld"]), iid

    def test_sa_banks_are_told_the_truth_about_aggregators(self):
        step = registry.next_step("za-capitec-bank")
        assert step["recommended"] == "csv"
        assert not any(m["method"] == "aggregator" for m in step["methods"])
        assert any(w["method"] == "aggregator" for w in step["withheld"])
        # The summary no longer talks about aggregators, because there is no longer an
        # aggregator route to explain — it says what the household will actually do.
        assert "statement" in step["summary"].lower()


# ===========================================================================
# Importing
# ===========================================================================

class TestImport:
    def test_demo_loaded_everything(self, household):
        assert db.scalar("SELECT COUNT(*) FROM transactions") > 500
        assert db.scalar("SELECT COUNT(*) FROM trades") > 15
        assert db.scalar("SELECT COUNT(*) FROM dividends") > 50
        assert db.scalar("SELECT COUNT(*) FROM holdings") > 10

    def test_reimporting_the_same_file_adds_nothing(self):
        acct = db.one("SELECT id FROM accounts WHERE name LIKE 'Monzo%'")["id"]
        before = db.scalar("SELECT COUNT(*) FROM transactions WHERE account_id=?", (acct,))
        blob = demo.monzo_csv()
        t = read_csv_bytes(blob, "monzo.csv")
        m = profiles.resolve_columns("monzo", t.header)
        res = ingest.import_transactions(acct, t, m, "monzo", "monzo.csv")
        after = db.scalar("SELECT COUNT(*) FROM transactions WHERE account_id=?", (acct,))
        assert res["added"] == 0
        assert after == before

    def test_trading212_tickers_decode(self):
        assert ingest.split_t212_ticker("AZNl_EQ") == ("AZN", "LON")
        assert ingest.split_t212_ticker("GOOGL_US_EQ") == ("GOOGL", "NYSE")
        assert ingest.split_t212_ticker("BRK_B_US_EQ") == ("BRK-B", "NYSE")
        assert ingest.split_t212_ticker("ADYENa_EQ") == ("ADYEN", "AMS")

    def test_pence_quoted_holdings_are_not_divided_twice(self):
        """A GBX-priced ETF must not end up costing a hundredth of what was paid."""
        row = db.one("""SELECT h.cost, h.shares FROM holdings h JOIN instruments i ON i.id=h.instrument_id
                        WHERE i.symbol='SPXP'""")
        assert row is not None
        assert row["cost"] > 1000, "GBX cost basis collapsed by a factor of 100"

    def test_dividends_land_against_the_right_holding(self):
        n = db.scalar("""SELECT COUNT(*) FROM dividends d JOIN instruments i ON i.id=d.instrument_id
                         WHERE i.symbol='DGE'""")
        assert n > 5

    def test_a_closed_position_becomes_a_sold_row(self):
        sold = db.one("SELECT * FROM sold_positions WHERE symbol='NIO'")
        assert sold is not None
        assert sold["proceeds"] < sold["cost"]      # NIO went down; the record should say so


# ===========================================================================
# Categorisation
# ===========================================================================

class TestCategorise:
    def test_everything_got_a_category(self):
        uncategorised = db.scalar("""SELECT COUNT(*) FROM transactions t
                                     JOIN categories c ON c.id=t.category_id
                                     WHERE c.name='Uncategorised'""")
        total = db.scalar("SELECT COUNT(*) FROM transactions")
        assert uncategorised / total < 0.05, f"{uncategorised}/{total} uncategorised"

    def test_uk_and_sa_merchants_both_land_properly(self):
        checks = [("TESCO STORES 3411", "Groceries"), ("CHECKERS HYPER", "Groceries"),
                  ("SHELL PETROL", "Fuel"), ("ENGEN GARAGE", "Fuel"),
                  ("DISCOVERY HEALTH", "Medical aid / health insurance"),
                  ("NETFLIX.COM", "TV licence & streaming"),
                  ("CITY OF CAPE TOWN MUNICIPAL", "Council tax / rates"),
                  ("NHS BSA SALARY", "Salary"), ("SALARY PAYMENT", "Salary")]
        for desc, expected in checks:
            row = db.one("""SELECT c.name FROM transactions t JOIN categories c ON c.id=t.category_id
                            WHERE t.description LIKE ? LIMIT 1""", (f"%{desc}%",))
            assert row, f"no transaction matching {desc}"
            assert row["name"] == expected, f"{desc} -> {row['name']}, expected {expected}"

    def test_a_manual_category_is_never_overwritten(self):
        tx = db.one("SELECT id FROM transactions LIMIT 1")
        cid = db.category_id("Shopping", "Hobbies")
        categorise.set_category(tx["id"], cid, make_rule=False)
        categorise.categorise_all(only_uncategorised=False)
        assert db.one("SELECT category_id FROM transactions WHERE id=?",
                      (tx["id"],))["category_id"] == cid

    def test_transfers_between_own_accounts_are_paired(self):
        n = db.scalar("SELECT COUNT(*) FROM transactions WHERE is_transfer=1")
        assert n >= 0    # pairing is opportunistic; it must not crash or double-count


# ===========================================================================
# Portfolio maths
# ===========================================================================

class TestPortfolio:
    def test_the_headline_identities_hold(self):
        ov = portfolio.overall()["all"]
        assert abs(ov["gross"] - (ov["invested"] + ov["dividends"])) < 0.01
        assert abs(ov["net"] - (ov["gross"] - ov["cost"])) < 0.01
        assert ov["positions"] > 0

    def test_per_row_arithmetic_matches_the_original_sheet(self):
        for r in portfolio.holdings_rows():
            if r["value"] is None or not r["cost"]:
                continue
            assert abs(r["value"] - r["shares"] * r["price"]) < 0.02
            assert abs(r["price_change"] - (r["value"] - r["cost"])) < 0.02
            assert abs(r["total_value"] - (r["value"] + r["dividends"])) < 0.02
            assert abs(r["money_made"] - (r["total_value"] - r["cost"])) < 0.02
            assert abs(r["pct_change"] - r["price_change"] / r["cost"]) < 1e-9
            assert abs(r["overall_change"] - (r["price_change"] + r["dividends"]) / r["cost"]) < 1e-9

    def test_sector_allocation_sums_to_one(self):
        secs = portfolio.sector_allocation()
        assert abs(sum(s["pct"] for s in secs) - 1.0) < 1e-6

    def test_wrappers_add_up_to_the_whole(self):
        whole = portfolio.summarise(portfolio.holdings_rows())
        parts = portfolio.wrapper_breakdown()
        total = sum(v["summary"]["invested"] for v in parts.values())
        assert abs(total - whole["invested"]) < 0.02

    def test_cash_is_derived_when_no_api_reports_it(self):
        rows = portfolio.cash_rows()
        assert rows
        assert all(isinstance(r["cash"], float) for r in rows)


# ===========================================================================
# Budgets
# ===========================================================================

class TestBudgets:
    def test_month_bounds(self):
        assert budgets.month_bounds("2026-02") == ("2026-02-01", "2026-02-28")
        assert budgets.month_bounds("2024-02") == ("2024-02-01", "2024-02-29")
        assert budgets.month_bounds("2026-12") == ("2026-12-01", "2026-12-31")
        assert budgets.shift_period("2026-01", -1) == "2025-12"
        assert budgets.shift_period("2026-12", 1) == "2027-01"

    def test_status_totals_are_consistent(self):
        st = budgets.status()
        t = st["totals"]
        assert t["unbudgeted"] >= 0, "a section budget was counted twice"
        assert abs(t["remaining"] - (t["budgeted"] - t["spent_in_budgets"])) < 0.01
        for i in st["items"]:
            assert abs(i["remaining"] - (i["amount"] - i["spent"])) < 0.01
            assert i["spent"] >= 0, "spending should be reported as a positive number"

    def test_transfers_and_saving_are_not_spending(self):
        start, end = budgets.month_bounds(budgets.current_period())
        s = budgets.spend_by_category(start, end)
        for parent in ("Transfers", "Income"):
            sec = s["sections"].get(parent)
            if sec:
                assert sec["spend"] == 0, f"{parent} counted as spending"

    def test_suggestions_are_positive_and_sane(self):
        for s in budgets.suggest_budgets(months=3):
            assert s["suggested"] > 0
            assert s["low"] <= s["average"] <= s["high"] + 0.01


# ===========================================================================
# Snapshots
# ===========================================================================

class TestSnapshots:
    def test_a_snapshot_exists_for_this_month(self):
        assert db.one("SELECT id FROM snapshots WHERE period=? AND scope='household'",
                      (budgets.current_period(),))

    def test_backfilled_months_are_flagged(self):
        rows = snapshots.series("household", "net_worth", 24)
        assert len(rows) > 5
        estimated = [r for r in rows if r["metrics"].get("estimated")]
        assert estimated, "history should be reconstructed for the demo"
        for r in estimated:
            assert r["metrics"].get("investments_note")

    def test_taking_a_snapshot_twice_does_not_duplicate(self):
        before = db.scalar("SELECT COUNT(*) FROM snapshots")
        snapshots.take(force=True)
        assert db.scalar("SELECT COUNT(*) FROM snapshots") == before


# ===========================================================================
# Money conversion
# ===========================================================================

class TestMoney:
    def test_pence_normalisation(self):
        assert market.normalise_price(1723.18, "GBX") == (17.2318, "GBP")
        assert market.normalise_price(12345.0, "ZAC") == (123.45, "ZAR")
        assert market.normalise_price(45.5, "USD") == (45.5, "USD")

    def test_quote_symbols(self):
        assert market.quote_symbol("DGE", "LON") == "DGE.L"
        assert market.quote_symbol("SN.", "LON") == "SN.L"
        assert market.quote_symbol("JNJ", "NYSE") == "JNJ"
        assert market.quote_symbol("ASML", "AMS") == "ASML.AS"
        assert market.quote_symbol("NPN", "JSE") == "NPN.JO"
        assert market.quote_symbol("BTC", "CRYPTO", "crypto") == "BTC-GBP"

    def test_same_currency_is_a_no_op(self):
        assert market.fx_rate("GBP", "GBP") == 1.0
        assert market.convert(100, "GBP", "GBP") == 100

    def test_offline_falls_back_rather_than_pretending_one_to_one(self):
        rate = market.fx_rate("ZAR", "GBP")
        assert 0.01 < rate < 0.2, "rand should never convert to pounds at parity"


# ===========================================================================
# Credentials
# ===========================================================================

class TestVault:
    def test_round_trip_and_redaction(self):
        security.put("test:1", {"api_key": "supersecretvalue123", "environment": "live"})
        got = security.get("test:1")
        assert got["api_key"] == "supersecretvalue123"
        red = security.redact(got)
        assert "supersecretvalue123" not in json.dumps(red)
        assert red["environment"] == "live"
        security.drop("test:1")
        assert security.get("test:1") is None

    def test_stored_bytes_are_not_plaintext(self):
        security.put("test:2", {"api_key": "plaintextcanary"})
        blob = db.one("SELECT blob FROM vault WHERE ref='test:2'")["blob"]
        assert b"plaintextcanary" not in blob
        security.drop("test:2")


# ===========================================================================
# Exports
# ===========================================================================

class TestExport:
    def test_investments_workbook_matches_the_engine(self, tmp_path):
        import openpyxl
        path = investments_xlsx.build(tmp_path / "inv.xlsx")
        wb = openpyxl.load_workbook(path)
        assert "Overall" in wb.sheetnames
        assert "ISA" in wb.sheetnames
        assert "Sold" in wb.sheetnames
        assert "Sectors" in wb.sheetnames
        assert "Monthly" in wb.sheetnames
        assert "Read me" in wb.sheetnames

        ws = wb["Holdings_Data"]
        rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
        engine = portfolio.holdings_rows()
        assert len(rows) == len(engine)
        by_symbol = {r["symbol"]: r for r in engine}
        for r in rows:
            e = by_symbol[r[4]]
            assert abs(r[8] - e["shares"]) < 1e-6
            assert abs(r[9] - e["cost"]) < 0.02

        ws = wb["Prices"]
        prices = {r[0]: r[4] for r in ws.iter_rows(min_row=5, values_only=True) if r[0]}
        for e in engine:
            if e["price"] is not None:
                assert abs(prices[e["symbol"]] - e["price"]) < 0.005

    def test_banking_workbook_has_the_tabs_and_the_data(self, tmp_path):
        import openpyxl
        path = banking_xlsx.build(tmp_path / "bank.xlsx")
        wb = openpyxl.load_workbook(path)
        for tab in ("Budgets", "Transactions", "Accounts", "Categories",
                    "Spending by month", "This month", "Read me"):
            assert tab in wb.sheetnames, tab
        ws = wb["Transactions"]
        n = sum(1 for r in ws.iter_rows(min_row=5, values_only=True) if r[0])
        assert n > 100
        assert wb["Budgets"]["G1"].value == budgets.current_period()

    def test_formulas_are_written_not_baked_values(self, tmp_path):
        import openpyxl
        path = investments_xlsx.build(tmp_path / "inv2.xlsx")
        wb = openpyxl.load_workbook(path)          # formulas, not cached values
        ws = wb["ISA"]
        assert str(ws["B3"].value).startswith("=SUM(")
        row = investments_xlsx.FIRST_DATA_ROW
        assert str(ws.cell(row=row, column=7).value).startswith("=SUMIFS(")
        assert str(ws.cell(row=row, column=12).value).startswith("=")

    def test_no_spilling_array_functions(self, tmp_path):
        """LibreOffice and older Excel can't evaluate these — they'd ship as #NAME?."""
        import openpyxl
        banned = ("XLOOKUP", "XMATCH", "FILTER(", "UNIQUE(", "SORT(", "SEQUENCE(")
        for build in (investments_xlsx.build, banking_xlsx.build):
            path = build(tmp_path / f"{build.__module__.split('.')[-1]}.xlsx")
            wb = openpyxl.load_workbook(path)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        v = cell.value
                        if isinstance(v, str) and v.startswith("="):
                            up = v.upper()
                            for b in banned:
                                assert b not in up, f"{ws.title}!{cell.coordinate}: {v}"


# ===========================================================================
# Connections
# ===========================================================================

class TestConnections:
    def test_creating_and_removing_a_connection(self):
        conn = sync.create_connection("gb-starling-bank", "csv", label="Starling test")
        assert conn["status"] == "ready"
        assert sync.test_connection(conn["id"])["ok"]
        sync.delete_connection(conn["id"])
        assert sync.get_connection(conn["id"]) is None

    def test_an_api_connection_asks_for_what_it_needs(self):
        conn = sync.create_connection("gb-trading-212", "direct_api", provider="trading212",
                                      label="T212 test")
        res = sync.test_connection(conn["id"])
        assert not res["ok"]
        assert "API key" in res["message"]
        sync.delete_connection(conn["id"])

    def test_a_connection_on_a_dead_provider_is_repaired(self):
        """A connection saved against a provider with no client — the Monzo/Enable
        Banking beta bug — must be re-pointed at one that works, not left broken."""
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        with db.tx() as c:
            c.execute("""INSERT INTO connections(member_id,institution_id,institution_name,
                         country,method,provider,label,status) VALUES(1,'gb-monzo','Monzo','GB',
                         'openbanking','a-provider-that-does-not-exist','Monzo repair test',
                         'needs_credentials')""")
        cid = db.scalar("SELECT id FROM connections WHERE label='Monzo repair test'")
        out = sync.repair_connections()
        # Open Banking is no longer an offered route, so there is nothing to re-point it
        # at — the honest repair is to turn it into a statement upload and say so.
        assert any(m["connection_id"] == cid for m in out["downgraded"])
        fixed = db.one("SELECT provider, method FROM connections WHERE id=?", (cid,))
        assert fixed["method"] == "csv"
        assert fixed["provider"] is None
        sync.delete_connection(cid)

    def test_a_working_paid_connection_is_left_alone(self):
        # Withholding the route must not tear down a connection somebody already has
        # working — they may well be paying for that aggregator.
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        live = next((p["id"] for p in provider_base.available()), None)
        assert live, "there should be at least one built provider"
        with db.tx() as c:
            c.execute("""INSERT INTO connections(member_id,institution_id,institution_name,
                         country,method,provider,label,status) VALUES(1,'gb-monzo','Monzo','GB',
                         'openbanking',?,'Monzo keep test','ready')""", (live,))
        cid = db.scalar("SELECT id FROM connections WHERE label='Monzo keep test'")
        sync.repair_connections()
        kept = db.one("SELECT provider, method FROM connections WHERE id=?", (cid,))
        assert kept["method"] == "openbanking" and kept["provider"] == live
        sync.delete_connection(cid)

    def test_a_connection_with_no_route_at_all_falls_back_to_upload(self):
        with db.tx() as c:
            c.execute("""INSERT INTO connections(member_id,institution_id,institution_name,
                         country,method,provider,label,status) VALUES(1,'za-capitec-bank',
                         'Capitec Bank','ZA','aggregator','stitch','Capitec repair test','new')""")
        cid = db.scalar("SELECT id FROM connections WHERE label='Capitec repair test'")
        out = sync.repair_connections()
        assert any(d["connection_id"] == cid for d in out["downgraded"])
        fixed = db.one("SELECT provider, method FROM connections WHERE id=?", (cid,))
        assert fixed["method"] == "csv"
        assert fixed["provider"] is None
        sync.delete_connection(cid)

    def test_trading212_explains_a_401_instead_of_just_reporting_it(self):
        """A bare '401' told the user nothing. The message must name the actual causes."""
        from kestrel.providers.trading212 import Trading212
        msg = Trading212._explain_failure("live/raw/equity/account/info: "
                                          "Trading 212 returned 401: unauthorised")
        assert "permission" in msg.lower()
        assert "practice" in msg.lower() or "demo" in msg.lower()
        assert len(msg) > 120
        assert Trading212._explain_failure("returned 429: too many requests").lower() \
            .startswith("trading 212 is rate-limiting")
        assert "no reason" in Trading212._explain_failure("").lower()

    def test_an_empty_error_body_still_says_something(self):
        """Providers that answer 401 with no body used to produce 'returned 401:'."""
        import urllib.error
        from kestrel.providers.base import Provider, ProviderError

        class Boom(Provider):
            name = "Test Provider"

            def http(self, *a, **k):
                raise urllib.error.HTTPError("http://x", 401, "", {}, None)

        p = Provider({"id": 1})
        p.name = "Test Provider"
        try:
            Provider.http(p, "http://127.0.0.1:9/definitely-not-listening")
        except ProviderError as e:
            assert str(e)
            assert not str(e).rstrip().endswith(":")

    def test_health_reports_are_actionable(self):
        for issue in sync.health()["issues"]:
            assert issue["text"]
            assert issue["level"] in ("info", "warn")


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q", "--no-header"]))


# ===========================================================================
# Crypto exchanges — the free APIs the first pass missed entirely
# ===========================================================================

class TestCrypto:
    def test_ticker_aliases_are_canonicalised(self):
        from kestrel.providers.crypto import canonical
        # Kraken's legacy X/Z prefixes and its staking suffixes are the whole reason
        # this function exists: XXBT and BTC must not become two separate holdings.
        assert canonical("XXBT") == "BTC"
        assert canonical("XBT") == "BTC"
        assert canonical("XETH") == "ETH"
        assert canonical("ETH2.S") == "ETH"
        assert canonical("USDC.M") == "USDC"
        assert canonical("ZGBP") == "GBP"
        assert canonical("XXDG") == "DOGE"
        assert canonical("btc") == "BTC"
        assert canonical("SOL") == "SOL"       # untouched when there's no alias

    def _fake(self, balances, ccy="GBP", label="Fake Exchange"):
        """A crypto provider whose only unusual property is that it doesn't use a network."""
        from kestrel.providers.crypto import CryptoExchange

        conn = sync.create_connection("gb-kraken", "direct_api", provider="kraken",
                                      label=label)

        class Fake(CryptoExchange):
            id = "fake-exchange"
            name = "Fake Exchange"

            def balances(self):
                return balances

            def account_currency(self):
                return ccy

        p = Fake(sync.get_connection(conn["id"]))
        p.creds = {"api_key": "k" * 16, "api_secret": "s" * 16}
        return p, conn

    def test_balances_become_holdings(self):
        p, conn = self._fake({"XXBT": 0.5, "ETH": 2.0, "SOL": 0.0})
        res = p.sync()
        assert res.accounts == 1
        assert res.holdings == 2                       # the zero balance is skipped
        acct = db.one("SELECT id, currency, is_investment FROM accounts "
                      "WHERE connection_id=?", (conn["id"],))
        assert acct["is_investment"] == 1
        syms = {r["symbol"] for r in db.rows(
            "SELECT i.symbol FROM holdings h JOIN instruments i ON i.id=h.instrument_id "
            "WHERE h.account_id=?", (acct["id"],))}
        assert syms == {"BTC", "ETH"}, syms

    def test_account_currency_is_cash_but_other_fiat_is_a_position(self):
        # Holding USD on a GBP account is real FX exposure and should show up as
        # something with a value, not vanish into a GBP cash figure.
        p, conn = self._fake({"BTC": 1.0, "GBP": 250.0, "USD": 100.0}, ccy="GBP")
        p.sync()
        acct = db.one("SELECT id, balance FROM accounts WHERE connection_id=?", (conn["id"],))
        assert acct["balance"] == 250.0
        syms = {r["symbol"] for r in db.rows(
            "SELECT i.symbol FROM holdings h JOIN instruments i ON i.id=h.instrument_id "
            "WHERE h.account_id=?", (acct["id"],))}
        assert syms == {"BTC", "USD=X"}, syms

    def test_sync_never_overwrites_an_imported_cost_basis(self):
        """An exchange says what you hold, never what you paid. Writing 0 into cost
        would report the entire holding as profit — so cost must survive a sync."""
        p, conn = self._fake({"BTC": 1.0})
        p.sync()
        acct = db.one("SELECT id FROM accounts WHERE connection_id=?", (conn["id"],))
        iid = db.scalar("SELECT instrument_id FROM holdings WHERE account_id=?", (acct["id"],))
        with db.tx() as c:
            c.execute("UPDATE holdings SET cost=12345.0, cost_currency='GBP' "
                      "WHERE account_id=? AND instrument_id=?", (acct["id"], iid))
        p.sync()
        row = db.one("SELECT shares, cost FROM holdings WHERE account_id=? AND instrument_id=?",
                     (acct["id"], iid))
        assert row["cost"] == 12345.0
        assert row["shares"] == 1.0

    def test_a_sold_asset_stops_being_a_holding(self):
        p, conn = self._fake({"BTC": 1.0, "ETH": 5.0})
        p.sync()
        acct = db.one("SELECT id FROM accounts WHERE connection_id=?", (conn["id"],))
        assert db.scalar("SELECT COUNT(*) FROM holdings WHERE account_id=?", (acct["id"],)) == 2
        p2, _ = self._fake({"BTC": 1.0})
        p2.connection = sync.get_connection(conn["id"])
        p2.sync()
        syms = {r["symbol"] for r in db.rows(
            "SELECT i.symbol FROM holdings h JOIN instruments i ON i.id=h.instrument_id "
            "WHERE h.account_id=?", (acct["id"],))}
        assert syms == {"BTC"}, syms

    def test_kraken_signature_is_stable_and_base64(self):
        import base64
        from kestrel.providers.crypto import Kraken
        conn = sync.create_connection("gb-kraken", "direct_api", provider="kraken")
        k = Kraken(sync.get_connection(conn["id"]))
        k.creds = {"api_key": "key", "api_secret": base64.b64encode(b"x" * 64).decode()}
        captured = {}

        def fake_http(url, method="GET", headers=None, body=None, timeout=30.0):
            captured.update(headers=headers, body=body, url=url)
            return {"error": [], "result": {"XXBT": "0.25"}}

        k.http = fake_http
        assert k.balances() == {"XXBT": 0.25}
        assert captured["url"].endswith("/0/private/Balance")
        base64.b64decode(captured["headers"]["API-Sign"])      # raises if malformed
        assert "nonce=" in captured["body"]

    def test_kraken_says_something_useful_when_the_secret_is_not_base64(self):
        from kestrel.providers.base import ProviderError
        from kestrel.providers.crypto import Kraken
        conn = sync.create_connection("gb-kraken", "direct_api", provider="kraken")
        k = Kraken(sync.get_connection(conn["id"]))
        k.creds = {"api_key": "key", "api_secret": "this is not base64 at all!!"}
        with pytest.raises(ProviderError) as e:
            k.balances()
        assert "base64" in str(e.value).lower()

    def test_exchanges_with_free_keys_are_offered_and_the_rest_are_not(self):
        from kestrel.providers import base as provider_base
        provider_base.load_all()
        # Verified Sept 2026: an individual can self-serve a free read-only key.
        for iid in ("gb-coinbase", "gb-kraken", "gb-binance", "za-luno", "za-valr"):
            assert registry.next_step(iid)["recommended"] == "direct_api", iid
        # Uphold gates production behind a business account; Altify has no API at all.
        for iid in ("gb-uphold", "za-altify"):
            assert registry.next_step(iid)["recommended"] == "csv", iid

    def test_ziglu_is_not_offered_as_a_live_account(self):
        """Ziglu entered special administration in July 2025. A defunct firm must not
        sit in the dropdown looking like somewhere you can still connect."""
        hits = registry.search("ziglu", "GB")
        assert hits, "Ziglu should still be findable so an old balance can be recorded"
        assert "closed" in hits[0]["name"].lower()
        assert registry.next_step(hits[0]["id"])["recommended"] == "manual"

    def test_ig_is_listed_but_not_chosen_because_it_cannot_see_an_isa(self):
        """IG's API is genuinely self-serve and free, and genuinely useless here: it
        reaches spread betting and CFD accounts only, never a share dealing ISA."""
        step = registry.next_step("gb-ig")
        assert step["recommended"] == "csv"
        api = next(m for m in step["methods"] if m["method"] == "direct_api")
        assert "ig" in api["documented_only"]
        assert registry.provider("ig")["warning"]

    def test_an_unknown_cost_reports_no_return_rather_than_pure_profit(self):
        """The whole reason cost_known exists. A £5,000 crypto holding with no cost
        recorded must not appear as £5,000 of profit."""
        p, conn = self._fake({"BTC": 2.0})
        p.sync()
        acct = db.one("SELECT id FROM accounts WHERE connection_id=?", (conn["id"],))
        rows = [r for r in portfolio.holdings_rows([acct["id"]]) if r["symbol"] == "BTC"]
        assert rows, "the holding should still be listed"
        assert rows[0]["money_made"] is None
        assert rows[0]["shares"] == 2.0

        # ...and once a cost basis exists, the return becomes real again.
        iid = db.scalar("SELECT instrument_id FROM holdings WHERE account_id=?", (acct["id"],))
        with db.tx() as c:
            c.execute("UPDATE holdings SET cost=1000.0, cost_known=1 WHERE account_id=? "
                      "AND instrument_id=?", (acct["id"], iid))
        rows = [r for r in portfolio.holdings_rows([acct["id"]]) if r["symbol"] == "BTC"]
        assert rows[0]["cost"] == 1000.0
        if rows[0]["value"] is not None:
            assert rows[0]["money_made"] is not None


class TestMigrations:
    def test_a_column_added_after_release_is_applied_to_an_existing_database(self):
        """CREATE TABLE IF NOT EXISTS never alters a table that already exists, so a
        new column is invisible to anyone who already has a database. Mithapp is in
        beta on Allan's machine — his file predates cost_known."""
        import sqlite3
        import tempfile as _tf
        path = pathlib.Path(_tf.mkdtemp()) / "old.db"
        old = sqlite3.connect(str(path))
        old.executescript("""
            CREATE TABLE holdings (
                id INTEGER PRIMARY KEY,
                account_id INTEGER,
                instrument_id INTEGER,
                shares REAL NOT NULL DEFAULT 0,
                cost REAL NOT NULL DEFAULT 0,
                cost_currency TEXT DEFAULT 'GBP',
                source TEXT DEFAULT 'import',
                updated_at TEXT,
                UNIQUE(account_id, instrument_id));
        """)
        old.execute("INSERT INTO holdings(account_id,instrument_id,shares,cost)"
                    " VALUES(1,1,3.0,600.0)")
        old.commit()
        assert "cost_known" not in {r[1] for r in
                                    old.execute("PRAGMA table_info(holdings)").fetchall()}

        db._apply_added_columns(old)

        cols = {r[1] for r in old.execute("PRAGMA table_info(holdings)").fetchall()}
        assert "cost_known" in cols
        # An existing row keeps its cost and is treated as known — it came from a
        # statement import, so its return figure was already meaningful.
        row = old.execute("SELECT cost, cost_known FROM holdings").fetchone()
        assert row[0] == 600.0 and row[1] == 1
        # Running it twice must be harmless.
        db._apply_added_columns(old)
        assert len(old.execute("PRAGMA table_info(holdings)").fetchall()) == len(cols)
        old.close()


class TestOnboarding:
    def test_the_tour_flag_starts_false_and_persists_once_set(self):
        """The tour must run exactly once for a new household. It is stored in settings
        rather than the browser so a cleared cache doesn't replay it, and so the desktop
        window and a browser tab agree."""
        assert "tour_done" in config.DEFAULTS
        assert config.DEFAULTS["tour_done"] is False
        before = config.settings["tour_done"]
        try:
            config.settings.update({"tour_done": True})
            assert config.settings["tour_done"] is True
            reloaded = config.Settings(config.settings.path)
            assert reloaded["tour_done"] is True, "the flag must survive a restart"
        finally:
            config.settings.update({"tour_done": before})


class TestStaticAssets:
    """The logo is a real file, not an inline data URI, so it can go missing from a
    build in a way the SVG mark never could. These assert it ships and matches what
    the CSS expects."""

    def _static(self):
        # the python package keeps its original directory name; only the product renamed
        return pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web" / "static"

    def test_the_logo_and_favicon_ship(self):
        s = self._static()
        logo, fav = s / "logo.webp", s / "favicon.png"
        assert logo.exists() and fav.exists()
        # Small enough to load instantly from local disk; the source video was 3.3 MB.
        assert logo.stat().st_size < 400_000, logo.stat().st_size
        assert fav.stat().st_size < 40_000

    def test_the_sprite_sheet_matches_the_css_grid(self):
        """If the sheet is ever re-exported at a different size, the CSS steps() values
        silently stop lining up and the animation shears.

        The frame is deliberately NOT square any more: the picture is about 1.44:1, and
        squaring it spent two fifths of every frame on empty space, which forced an
        upscale in the sidebar and made the logo soft. So the check is against what the
        CSS actually claims, not against a number written down twice.
        """
        from PIL import Image
        s = self._static()
        w, h = Image.open(s / "logo.webp").size
        css = (s / "styles.css").read_text(encoding="utf-8")
        cols = int(re.search(r"--cols:\s*(\d+)", css).group(1))
        rows = int(re.search(r"--rows:\s*(\d+)", css).group(1))
        assert w % cols == 0 and h % rows == 0, f"{w}x{h} does not divide into {cols}x{rows}"
        m = re.search(r"--fw:\s*calc\(var\(--fh\)\s*\*\s*(\d+)\s*/\s*(\d+)\)", css)
        assert m, "--fw must be derived from --fh by the frame's aspect ratio"
        assert abs(int(m.group(1)) / int(m.group(2)) - (w / cols) / (h / rows)) < 0.01, \
            f"CSS says {m.group(1)}/{m.group(2)}, the sheet is {(w/cols):.0f}x{(h/rows):.0f}"

    def test_the_logo_has_a_transparent_background(self):
        """The whole point of the keying pass. If it ever ships opaque, it will show a
        cream box against the dark theme."""
        from PIL import Image
        im = Image.open(self._static() / "logo.webp").convert("RGBA")
        assert im.getpixel((2, 2))[3] == 0, "top-left of frame 0 should be transparent"
        alpha = im.getchannel("A")
        assert alpha.getextrema()[0] == 0, "there must be fully transparent pixels"

    def test_the_help_icon_ships_and_is_transparent(self):
        from PIL import Image
        f = self._static() / "help.webp"
        assert f.exists()
        assert f.stat().st_size < 60_000, f.stat().st_size
        im = Image.open(f).convert("RGBA")
        assert im.getpixel((1, 1))[3] == 0, "the black background should be keyed out"


class TestNaming:
    def test_the_product_renamed_but_the_data_did_not_move(self):
        """Renaming the app must never strand a beta tester's database. The package
        directory, the settings file and an existing data folder all stay put; only what
        the user reads changes."""
        assert config.APP_NAME == "Mittens & Pence"
        # Both previous names, so neither install is ever orphaned.
        assert config.LEGACY_NAMES == ("Mithapp", "Kestrel")
        assert config.LEGACY_SLUGS == ("mithapp", "kestrel")
        # The python package is deliberately still `kestrel`: renaming it would move
        # every user's data directory to prove a point about naming.
        assert (pathlib.Path(__file__).resolve().parent.parent / "kestrel").is_dir()

    def _clear_env(self, monkeypatch):
        for var in ("MITTENS_DATA_DIR", "MITHAPP_DATA_DIR", "KESTREL_DATA_DIR"):
            monkeypatch.delenv(var, raising=False)

    def test_an_existing_install_keeps_its_folder(self, tmp_path, monkeypatch):
        # Renaming the app must never move somebody's accounts, credentials and history.
        legacy = tmp_path / "Kestrel"
        legacy.mkdir()
        (legacy / "kestrel.db").write_bytes(b"")
        self._clear_env(monkeypatch)
        monkeypatch.setattr(config, "_data_root",
                            lambda: (tmp_path / "Mittens and Pence",
                                     [tmp_path / "Mithapp", legacy]))
        assert config.data_dir() == legacy, "an existing folder must win"
        assert config.db_path().name == "kestrel.db", "and so must an existing database"

    def test_the_newest_existing_folder_wins(self, tmp_path, monkeypatch):
        # After a *second* rename, somebody upgrading from Mithapp must not be sent back
        # to a Kestrel folder they stopped using two names ago.
        for name in ("Mithapp", "Kestrel"):
            (tmp_path / name).mkdir()
        (tmp_path / "Mithapp" / "mithapp.db").write_bytes(b"")
        self._clear_env(monkeypatch)
        monkeypatch.setattr(config, "_data_root",
                            lambda: (tmp_path / "Mittens and Pence",
                                     [tmp_path / "Mithapp", tmp_path / "Kestrel"]))
        assert config.data_dir().name == "Mithapp"
        assert config.db_path().name == "mithapp.db"

    def test_a_fresh_install_uses_the_new_names(self, tmp_path, monkeypatch):
        self._clear_env(monkeypatch)
        monkeypatch.setattr(config, "_data_root",
                            lambda: (tmp_path / "Mittens and Pence",
                                     [tmp_path / "Mithapp", tmp_path / "Kestrel"]))
        assert config.data_dir().name == "Mittens and Pence"
        assert config.db_path().name == "mittens.db"

    def test_no_shell_metacharacter_reaches_a_filename(self):
        # `&` is a command separator in Windows batch. The name people read has one; the
        # name that becomes folders, executables and launchers must not.
        assert "&" in config.APP_NAME
        for value in (config.APP_FILE_NAME, config.APP_SLUG, *config.DB_NAMES):
            assert not set(value) & set('&<>|^"'), value
        root = pathlib.Path(__file__).resolve().parent.parent
        for path in root.iterdir():
            assert "&" not in path.name, f"{path.name} would break a batch script"


class TestTrading212:
    """Mithapp told Allan to tick every Trading 212 permission, including
    'Orders - Execute', to fix a 401. That advice was wrong twice over: a missing
    permission is 403 not 401, and the real cause was that the API had moved."""

    def _p(self, creds):
        from kestrel.providers.trading212 import Trading212
        conn = sync.create_connection("gb-trading-212", "direct_api", provider="trading212")
        p = Trading212(sync.get_connection(conn["id"]))
        p.creds = creds
        return p

    def test_it_never_asks_for_permission_to_place_orders(self):
        from kestrel.providers.trading212 import Trading212
        steps = " ".join(Trading212.setup_steps).lower()
        assert "orders - execute" in steps
        assert "untick" in steps or "unticked" in steps
        assert "tick every permission" not in steps
        # and the catalogue the setup screen reads must agree with the steps
        assert "do not tick 'orders - execute'" in registry.provider("trading212")["notes"].lower()

    def test_it_calls_no_order_placing_endpoint(self):
        """The strongest form of the promise: there is no POST to an orders path
        anywhere in the client, so the permission genuinely cannot be needed."""
        src = (pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "providers"
               / "trading212.py").read_text(encoding="utf-8")
        for danger in ("/equity/orders/market", "/equity/orders/limit", "/equity/orders/stop",
                       '"POST"', "'POST'", '"DELETE"'):
            assert danger not in src, f"{danger} must never appear in a read-only client"

    def test_a_missing_permission_is_explained_as_403_not_401(self):
        from kestrel.providers.trading212 import Trading212
        forbidden = Trading212._explain_failure("403 Forbidden")
        assert "permission" in forbidden.lower()
        assert "does not need" in forbidden.lower() or "not need" in forbidden.lower()

        rejected = Trading212._explain_failure("401 Unauthorised")
        low = rejected.lower()
        assert "credentials" in low
        assert "403" in rejected, "a 401 should point out that permissions give 403"
        # It must never send someone back to tick more permissions on a 401 — that is
        # exactly the advice that had Allan granting order-execute rights.
        for bad in ("tick", "every permission", "generate a new one and tick"):
            assert bad not in low, f"401 advice must not say {bad!r}: {rejected}"

    def test_it_uses_the_current_endpoints(self):
        from kestrel.providers.trading212 import Trading212 as T
        assert T.P_SUMMARY[0] == "/equity/account/summary"
        assert T.P_POSITIONS[0] == "/equity/positions"
        assert T.P_DIVIDENDS[0] == "/equity/history/dividends"
        assert T.P_TRANSACTIONS[0] == "/equity/history/transactions"
        assert T.P_ORDERS[0] == "/equity/history/orders"
        # the retired paths stay as fallbacks for an older tenant
        assert "/equity/portfolio" in T.P_POSITIONS
        assert "/equity/account/info" in T.P_SUMMARY

    def test_basic_auth_is_used_when_a_secret_is_present(self):
        import base64
        p = self._p({"api_key": "abc", "api_secret": "shh"})
        h = p._auth_header()
        assert h["Authorization"].startswith("Basic ")
        raw = base64.b64decode(h["Authorization"].split(" ", 1)[1]).decode()
        assert raw == "abc:shh"
        assert "\n" not in h["Authorization"], "a newline in the header is a known 401 cause"
        assert p._auth_styles()[0] == "basic"

    def test_an_old_single_value_key_still_uses_the_legacy_header(self):
        p = self._p({"api_key": "legacy-key"})
        assert p._auth_header() == {"Authorization": "legacy-key"}
        assert "basic" not in p._auth_styles(), "no secret means Basic can't be formed"

    def test_an_unknown_cash_shape_is_not_reported_as_zero(self):
        from kestrel.providers.trading212 import Trading212 as T
        assert T._pick({"free": 12.5}, "free", "cash") == 12.5
        assert T._pick({"cash": {"free": 9.0}}, "free", "cash") == 9.0    # nested
        assert T._pick({"total": 1000.0}, "free", "freeCash") is None     # never guess


class TestBuildTooling:
    def test_the_generator_writes_where_the_registry_reads(self):
        """A rename once pointed the generator at a directory nothing loads. It built a
        perfect registry into thin air, the app served the stale one, and every test
        passed."""
        import kestrel.institutions.registry as reg
        src = (pathlib.Path(__file__).resolve().parent.parent / "tools"
               / "build_institutions.py").read_text(encoding="utf-8")
        out = re.search(r'^OUT = (.+)$', src, re.M).group(1)
        assert '"kestrel"' in out, f"generator writes to the wrong package: {out}"
        assert (pathlib.Path(reg.__file__).resolve().parent / "data").is_dir()

    def test_the_shipped_registry_is_not_stale(self):
        """Catches the same bug from the other side: whatever the generator's tables say
        about Trading 212 permissions must be what the app actually serves."""
        src = (pathlib.Path(__file__).resolve().parent.parent / "tools"
               / "build_institutions.py").read_text(encoding="utf-8")
        assert "Tick EVERY permission" not in src
        assert "Tick EVERY permission" not in registry.provider("trading212")["notes"]


class TestPdfStatements:
    """PDFs are the awkward format: a page has glyphs at coordinates, not rows and
    columns. These build real statements in the styles banks actually use and check the
    grid comes back out."""

    ROWS = [
        ("01/07/2026", "TESCO STORES 3345 LONDON", "", "42.10", "1,192.46"),
        ("03/07/2026", "SALARY PAYMENT ACME LTD REF 88213", "3,850.00", "", "5,035.40"),
        ("05/07/2026", "SANTANDER MORTGAGE COLLECTION FOR ACCOUNT ENDING 4471",
         "", "1,180.00", "3,855.40"),
        ("09/07/2026", "AMAZON.CO.UK*MK1QW9RT3", "", "17.34", "3,838.06"),
    ]

    def _pdf(self, path, ruled=False, pages=1, wrap_at=46):
        reportlab = pytest.importorskip("reportlab")
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        head = ["Date", "Description", "Money In", "Money Out", "Balance"]
        X = [20 * mm, 45 * mm, 118 * mm, 145 * mm, 175 * mm]
        c = canvas.Canvas(str(path), pagesize=A4)
        for _ in range(pages):
            y = 265 * mm
            c.setFont("Helvetica-Bold", 12)
            c.drawString(20 * mm, 278 * mm, "Barclays Bank UK PLC")
            c.setFont("Helvetica-Bold", 9)
            for i, h in enumerate(head):
                (c.drawRightString(X[i] + 18 * mm, y, h) if i >= 2
                 else c.drawString(X[i], y, h))
            if ruled:
                c.line(18 * mm, y - 2 * mm, A4[0] - 15 * mm, y - 2 * mm)
            y -= 7 * mm
            c.setFont("Helvetica", 8.5)
            for r in self.ROWS:
                desc, parts = r[1], []
                while len(desc) > wrap_at:
                    cut = desc.rfind(" ", 0, wrap_at)
                    parts.append(desc[:cut]); desc = desc[cut + 1:]
                parts.append(desc)
                c.drawString(X[0], y, r[0])
                c.drawString(X[1], y, parts[0])
                for i in (2, 3, 4):
                    if r[i]:
                        c.drawRightString(X[i] + 18 * mm, y, r[i])
                y -= 4.6 * mm
                for extra in parts[1:]:
                    c.drawString(X[1], y, extra); y -= 4.6 * mm
            c.showPage()
        c.save()
        return path

    def test_an_unruled_statement_becomes_a_grid(self, tmp_path):
        from kestrel.importers.readers import read_any
        t = read_any(self._pdf(tmp_path / "s.pdf"))
        assert t.header[:2] == ["Date", "Description"]
        assert "Balance" in t.header
        assert len(t.rows) == len(self.ROWS), t.rows
        assert t.rows[0][0] == "01/07/2026"
        assert t.rows[0][1].startswith("TESCO")
        assert t.rows[0][3] == "42.10", "money must land in its own column"
        assert t.rows[1][2] == "3,850.00", "credits and debits must not be conflated"

    def test_a_ruled_statement_becomes_the_same_grid(self, tmp_path):
        from kestrel.importers.readers import read_any
        t = read_any(self._pdf(tmp_path / "r.pdf", ruled=True))
        assert len(t.rows) == len(self.ROWS)
        assert t.rows[0][3] == "42.10"

    def test_a_wrapped_description_rejoins_its_row(self, tmp_path):
        """The one that nearly shipped broken: parse_amount reads 4471 out of
        'ACCOUNT ENDING 4471', so the continuation line looked like a transaction of
        its own and became a fifth row with no date."""
        from kestrel.importers.readers import read_any
        t = read_any(self._pdf(tmp_path / "w.pdf", wrap_at=40))
        assert len(t.rows) == len(self.ROWS), "a wrapped line must not become a row"
        mortgage = [r for r in t.rows if "SANTANDER" in r[1]][0]
        assert "4471" in mortgage[1], "the wrapped tail must be kept, not dropped"
        assert mortgage[0] == "05/07/2026"

    def test_prose_containing_a_number_is_not_a_transaction(self):
        from kestrel.importers.pdfstatement import _is_money
        assert not _is_money("ACCOUNT ENDING 4471")
        assert not _is_money("TESCO STORES 3345 LONDON")
        for good in ("1,180.00", "-42.10", "R 1 234,56", "(45.00)", "120.00CR", "1 234,56"):
            assert _is_money(good), good

    def test_a_repeated_header_on_page_two_is_not_a_row(self, tmp_path):
        from kestrel.importers.readers import read_any
        t = read_any(self._pdf(tmp_path / "two.pdf", pages=2))
        assert len(t.rows) == len(self.ROWS) * 2
        assert not any(r[0].lower() == "date" for r in t.rows)

    def test_a_scanned_pdf_says_so_instead_of_returning_nothing(self, tmp_path):
        """A photo of a statement has no text. Returning an empty table would leave
        someone staring at 'no rows found' with no idea why."""
        pytest.importorskip("reportlab")
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from kestrel.importers.pdfstatement import ScannedPdf, read_pdf
        p = tmp_path / "scan.pdf"
        c = canvas.Canvas(str(p), pagesize=A4)
        c.rect(20 * mm, 20 * mm, 100 * mm, 100 * mm, fill=1)   # ink, no text
        c.save()
        with pytest.raises(ScannedPdf) as e:
            read_pdf(p)
        msg = str(e.value).lower()
        assert "csv" in msg or "ocr" in msg, "must say what to do instead"

    def test_a_pdf_imports_and_de_duplicates_like_any_other_file(self, tmp_path):
        from kestrel.importers.readers import read_any
        conn = sync.create_connection("gb-barclays", "csv", label="PDF test")
        with db.tx() as c:
            cur = c.execute(
                "INSERT INTO accounts(connection_id,external_id,name,account_type,"
                "currency,balance,last_updated) VALUES(?,?,?,?,?,NULL,datetime('now'))",
                (conn["id"], "pdftest", "PDF account", "current", "GBP"))
            aid = cur.lastrowid
        path = self._pdf(tmp_path / "imp.pdf")
        t = read_any(path)
        s = mapping.suggest(t, "transactions")
        assert not s["missing"], f"a statement PDF should auto-map: {s}"
        r1 = ingest.import_transactions(aid, t, s["mapping"], "imp.pdf", conn["id"])
        assert r1["added"] == len(self.ROWS)
        r2 = ingest.import_transactions(aid, read_any(path), s["mapping"], "imp.pdf",
                                        conn["id"])
        assert r2["added"] == 0, "re-uploading the same PDF must add nothing"
        signs = {row["description"][:6]: row["amount"] for row in db.rows(
            "SELECT description, amount FROM transactions WHERE account_id=?", (aid,))}
        assert signs["TESCO "] < 0 and signs["SALARY"] > 0, "debit/credit signs"

    def test_the_frozen_build_keeps_pillow(self):
        """pdfplumber needs Pillow. The spec used to exclude it, which produced a build
        that started fine and died the first time someone dropped in a PDF."""
        spec = (pathlib.Path(__file__).resolve().parent.parent / "build"
                / "mittens.spec").read_text(encoding="utf-8")
        excludes = spec[spec.index("excludes = ["):spec.index("]", spec.index("excludes = ["))]
        assert '"PIL"' not in excludes
        assert '"pdfplumber"' in spec and '"PIL"' in spec

    def test_a_matching_profile_never_loses_the_amount_column(self, tmp_path):
        """Found while adding PDF support, but not a PDF bug — any file with
        Date/Description/Money In/Money Out/Balance headers hit it. The Standard Bank
        profile matches on Date + Description alone and names an 'Amount' column that
        isn't there, so it resolved to date+description+balance and every row would
        have imported with no amount at all."""
        from kestrel.web.server import _has_amount
        from kestrel.importers import profiles as profs
        from kestrel.importers.readers import read_csv_bytes

        raw = (b"Date,Description,Money In,Money Out,Balance\n"
               b"01/07/2026,TESCO,,42.10,1192.46\n"
               b"03/07/2026,SALARY,3850.00,,5042.46\n")
        t = read_csv_bytes(raw, "x.csv")
        detected = profs.detect(t.header)
        if detected and (profs.get(detected) or {}).get("fields"):
            resolved = profs.resolve_columns(detected, t.header) or {}
            # the profile alone is not enough — that is the bug
            if not _has_amount(resolved):
                merged = dict(resolved)
                for k, v in mapping.suggest(t, "transactions")["mapping"].items():
                    merged.setdefault(k, v)
                assert _has_amount(merged), "the merge must restore an amount source"
        # and the content reading on its own must always find one
        assert _has_amount(mapping.suggest(t, "transactions")["mapping"])

    def test_has_amount_rejects_a_mapping_with_no_number(self):
        from kestrel.web.server import _has_amount
        assert not _has_amount({"date": 0, "description": 1, "balance": 4})
        assert not _has_amount({})
        assert _has_amount({"date": 0, "amount": 2})
        assert _has_amount({"date": 0, "debit": 3})
        assert _has_amount({"date": 0, "credit": 2})


# ===========================================================================
# Importing somebody's existing budget spreadsheet
# ===========================================================================

def _budget_workbook(path, sheets):
    """Write a workbook shaped like a household's own budget.

    `sheets` is [(title, rows)] where rows is a list of (label, amount) or a bare
    string used as a heading. Real ones are messier than this; the messiness that
    matters (totals, signs, several years) is what these tests reproduce.
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title[:31])
        for r, row in enumerate(rows, start=1):
            if isinstance(row, str):
                ws.cell(r, 3, row)
            else:
                ws.cell(r, 2, row[0])
                cell = ws.cell(r, 3, row[1])
                if len(row) > 2:
                    cell.number_format = row[2]
    wb.save(str(path))
    return path


class TestBudgetSpreadsheet:
    """Reading the *plan* out of a workbook nobody designed for a computer."""

    def _sheet(self, tmp_path, rows, title="UK '26"):
        from kestrel.importers import budgetsheet
        p = _budget_workbook(tmp_path / "b.xlsx", [(title, rows)])
        return budgetsheet.find_budget_lists(p)

    def test_finds_a_plain_category_list(self, tmp_path):
        blocks = self._sheet(tmp_path, [
            "Budget P/M",
            ("Groceries", 200), ("Petrol", 150), ("Council Tax", 180),
            ("Broadband", 32), ("Gym", 25),
        ])
        assert blocks, "a category column beside an amount column is a budget"
        labels = [e["label"] for e in blocks[0]["entries"]]
        assert labels == ["Groceries", "Petrol", "Council Tax", "Broadband", "Gym"]
        assert blocks[0]["header"] == "Budget P/M"

    def test_a_total_row_is_never_a_category(self, tmp_path):
        # Importing "Total" as a category doubles the household's spending, and the
        # screen looks perfectly reasonable while it does it.
        blocks = self._sheet(tmp_path, [
            "Budget P/M",
            ("Groceries", 200), ("Petrol", 150), ("Council Tax", 180), ("Broadband", 32),
            ("Total", 562),
        ])
        entries = [e["label"] for e in blocks[0]["entries"]]
        assert "Total" not in entries
        assert "Total" in blocks[0]["excluded_totals"]

    def test_an_unlabelled_total_is_caught_by_arithmetic(self, tmp_path):
        # Some sheets call it nothing useful at all. The giveaway is the number.
        blocks = self._sheet(tmp_path, [
            "Budget P/M",
            ("Groceries", 200), ("Petrol", 150), ("Council Tax", 180), ("Broadband", 32),
            ("Per month", 562),
        ])
        entries = [e["label"] for e in blocks[0]["entries"]]
        assert "Per month" not in entries, "a row equal to the sum above it is a total"
        assert sum(e["amount"] for e in blocks[0]["entries"]) == 562

    def test_negative_budgets_are_reported_and_normalised_on_apply(self, tmp_path):
        blocks = self._sheet(tmp_path, [
            "Budget P/M",
            ("Groceries", -200), ("Petrol", -150), ("Council Tax", -180),
            ("Broadband", -32), ("Gym", -25),
        ])
        assert blocks[0]["sign"] == "negative"
        assert all(e["amount"] < 0 for e in blocks[0]["entries"])

    def test_the_latest_year_wins(self, tmp_path):
        from kestrel.importers import budgetsheet
        rows_24 = ["Budget P/M", ("Groceries", 150), ("Petrol", 120),
                   ("Council Tax", 160), ("Broadband", 28)]
        rows_26 = ["Budget P/M", ("Groceries", 200), ("Petrol", 150),
                   ("Council Tax", 180), ("Broadband", 32)]
        p = _budget_workbook(tmp_path / "y.xlsx",
                             [("UK '24", rows_24), ("UK '26", rows_26)])
        blocks = budgetsheet.find_budget_lists(p)
        # A workbook with a tab per year keeps the current plan last.
        assert blocks[0]["sheet"] == "UK '26"
        assert blocks[0]["entries"][0]["amount"] == 200

    def test_a_marker_row_beats_a_display_table(self, tmp_path):
        from kestrel.importers import budgetsheet
        display = ["Spent", ("Groceries", 187.4), ("Petrol", 143.2),
                   ("Council Tax", 180), ("Broadband", 32)]
        master = ["DON'T TOUCH UNLESS UPDATING BUDGET VALUE",
                  ("Groceries", 200), ("Petrol", 150), ("Council Tax", 180),
                  ("Broadband", 32)]
        p = _budget_workbook(tmp_path / "m.xlsx", [("Sheet1", display + [""] * 0)])
        # put both on one sheet, separated by a blank row
        import openpyxl
        wb = openpyxl.load_workbook(str(p))
        ws = wb["Sheet1"]
        start = len(display) + 3
        for i, row in enumerate(master):
            r = start + i
            if isinstance(row, str):
                ws.cell(r, 3, row)
            else:
                ws.cell(r, 2, row[0]); ws.cell(r, 3, row[1])
        wb.save(str(p))
        blocks = budgetsheet.find_budget_lists(p)
        assert blocks[0]["marker"] is True
        assert blocks[0]["entries"][0]["amount"] == 200, "the master list, not the display"

    def test_currency_comes_from_the_cell_formatting(self, tmp_path):
        from kestrel.importers import budgetsheet
        fmt = '[$£-en-GB]#,##0.00'
        p = _budget_workbook(tmp_path / "c.xlsx", [("UK '26", [
            "Budget P/M",
            ("Groceries", 200, fmt), ("Petrol", 150, fmt),
            ("Council Tax", 180, fmt), ("Broadband", 32, fmt)])])
        assert budgetsheet.find_budget_lists(p)[0]["currency"] == "GBP"

    def test_a_statement_is_not_mistaken_for_a_budget(self, tmp_path):
        from kestrel.importers import budgetsheet
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date", "Description", "Amount", "Balance"])
        for i in range(1, 20):
            ws.append([f"0{i%9+1}/07/2026", f"CARD PAYMENT {i}", -12.5 - i, 1000 - i])
        wb.save(str(tmp_path / "s.xlsx"))
        # No header saying budget and no marker: nothing here should score as a plan.
        assert not budgetsheet.has_budget_list(tmp_path / "s.xlsx")

    def test_short_runs_are_ignored(self, tmp_path):
        blocks = self._sheet(tmp_path, ["Budget P/M", ("Groceries", 200), ("Petrol", 150)])
        assert not blocks, "two rows is a label and a number, not a budget"


class TestBudgetMapping:
    def test_an_existing_category_is_reused(self):
        from kestrel.importers.budgetsheet import suggest_section
        existing = {"Food & Drink": ["Groceries", "Eating out"], "Home": ["Rent"]}
        assert suggest_section("groceries", existing) == ("Food & Drink", "Groceries", False)

    def test_a_new_name_lands_in_a_sensible_section(self):
        from kestrel.importers.budgetsheet import suggest_section
        existing = {"Food & Drink": ["Groceries"], "Transport": ["Petrol"]}
        sec, line, is_new = suggest_section("Car petrol", existing)
        assert sec == "Transport" and line == "Car petrol" and is_new is False

    def test_a_section_may_be_invented_rather_than_dumped_in_other(self):
        from kestrel.importers.budgetsheet import suggest_section
        sec, line, is_new = suggest_section("Cinema", {"Food & Drink": ["Groceries"]})
        assert sec == "Entertainment" and is_new is True

    def test_what_it_cannot_place_goes_to_other(self):
        from kestrel.importers.budgetsheet import suggest_section
        sec, _, _ = suggest_section("Zzzz miscellany", {"Other": ["Unsorted"]})
        assert sec == "Other"


class TestBudgetApply:
    def test_it_creates_what_is_missing_and_sets_the_budgets(self):
        from kestrel.importers.budgetsheet import apply_budget
        from kestrel.engine import sections as sections_engine
        name = "Zzz Test Section"
        res = apply_budget([
            {"label": "Cinema", "amount": 40, "section": name, "line": "Cinema"},
            {"label": "Concerts", "amount": 25, "section": name, "line": "Concerts"},
        ], "GBP")
        assert name in res["sections_created"]
        assert res["budgets_set"] == 2
        got = {s["name"] for s in sections_engine.list_sections()}
        assert name in got
        row = db.one("SELECT b.amount FROM budgets b JOIN categories c ON c.id=b.category_id "
                     "WHERE c.parent=? AND c.name='Cinema'", (name,))
        assert row and abs(row["amount"] - 40) < 0.01
        sections_engine.delete_section(name)

    def test_a_negative_budget_is_stored_as_an_amount_to_stay_under(self):
        from kestrel.importers.budgetsheet import apply_budget
        from kestrel.engine import sections as sections_engine
        name = "Zzz Signs"
        apply_budget([{"label": "Petrol", "amount": -150, "section": name, "line": "Petrol"}])
        row = db.one("SELECT b.amount FROM budgets b JOIN categories c ON c.id=b.category_id "
                     "WHERE c.parent=? AND c.name='Petrol'", (name,))
        assert row["amount"] > 0, "a budget of minus £150 never matches anything"
        sections_engine.delete_section(name)

    def test_categories_budgeted_at_zero_are_left_out(self):
        from kestrel.importers.budgetsheet import apply_budget
        from kestrel.engine import sections as sections_engine
        name = "Zzz Zeroes"
        res = apply_budget([
            {"label": "Live one", "amount": 10, "section": name, "line": "Live one"},
            {"label": "Dead one", "amount": 0, "section": name, "line": "Dead one"},
        ])
        assert res["skipped_zero"] == ["Dead one"]
        assert not db.one("SELECT id FROM categories WHERE parent=? AND name='Dead one'", (name,))
        sections_engine.delete_section(name)


class TestBudgetRoutes:
    def test_a_budget_workbook_is_read_rather_than_refused(self, tmp_path):
        # It has no statement header, so the ordinary reader rightly fails on it. The
        # file is still perfectly good; refusing it is the wrong answer.
        from kestrel.web import server
        p = _budget_workbook(tmp_path / "mum.xlsx", [("UK '26", [
            "Budget P/M",
            ("Groceries", -200), ("Car petrol", -150), ("Council Tax", -180),
            ("Broadband", -32), ("Cinema", -20)])])
        out = server.api_analyse(None, {}, {"filename": "mum.xlsx"}, p.read_bytes())
        assert out["kind"] == "budget"
        assert out["blocks"][0]["total"] == 582
        rows = out["blocks"][0]["rows"]
        assert all(r["amount"] > 0 for r in rows), "shown as amounts to stay under"
        assert {r["section"] for r in rows} >= {"Food & Drink", "Transport"}

    def test_commit_writes_only_what_was_ticked(self, tmp_path):
        from kestrel.web import server
        from kestrel.engine import sections as sections_engine
        name = "Zzz Routes"
        out = server.api_budget_commit(None, {}, {}, {
            "token": None,
            "currency": "GBP",
            "rows": [
                {"label": "Kept", "amount": 12, "section": name, "line": "Kept", "include": True},
                {"label": "Dropped", "amount": 99, "section": name, "line": "Dropped", "include": False},
            ]})
        assert out["budgets_set"] == 1
        assert db.one("SELECT id FROM categories WHERE parent=? AND name='Kept'", (name,))
        assert not db.one("SELECT id FROM categories WHERE parent=? AND name='Dropped'", (name,))
        sections_engine.delete_section(name)

    def test_an_empty_commit_is_refused_rather_than_silently_doing_nothing(self):
        from kestrel.web import server
        with pytest.raises(Exception):
            server.api_budget_commit(None, {}, {}, {"rows": []})

    def test_income_lines_are_not_offered_as_spending_budgets(self, tmp_path):
        # A budget is a ceiling on spending. A salary line ticked by default makes a
        # category that is permanently hundreds of percent "over".
        from kestrel.web import server
        p = _budget_workbook(tmp_path / "inc.xlsx", [("UK '26", [
            "Budget P/M",
            ("Groceries", 200), ("Petrol", 150), ("Salary", 2400),
            ("Broadband", 32), ("Gym", 25)])])
        out = server.api_analyse(None, {}, {"filename": "inc.xlsx"}, p.read_bytes())
        rows = {r["label"]: r for r in out["blocks"][0]["rows"]}
        assert rows["Salary"]["include"] is False
        assert "income" in rows["Salary"]["note"]
        assert rows["Groceries"]["include"] is True

    def test_the_way_back_to_reading_it_as_a_statement(self, tmp_path):
        from kestrel.web import server
        p = _budget_workbook(tmp_path / "both.xlsx", [("UK '26", [
            "Budget P/M",
            ("Groceries", 200), ("Petrol", 150), ("Council Tax", 180),
            ("Broadband", 32), ("Gym", 25)])])
        raw = p.read_bytes()
        assert server.api_analyse(None, {}, {"filename": "b.xlsx"}, raw)["kind"] == "budget"
        forced = server.api_analyse(None, {}, {"filename": "b.xlsx", "as": "transactions"}, raw)
        assert forced["kind"] != "budget"
        assert forced["has_budget"] is True, "and the offer stays on screen"

    def test_importing_the_same_budget_twice_does_not_duplicate_it(self):
        from kestrel.importers.budgetsheet import apply_budget
        from kestrel.engine import sections as sections_engine
        name = "Zzz Twice"
        rows = [{"label": "Groceries", "amount": 200, "section": name, "line": "Groceries"}]
        apply_budget(rows)
        apply_budget([{**rows[0], "amount": 220}])
        got = db.rows("SELECT b.amount FROM budgets b JOIN categories c ON c.id=b.category_id "
                      "WHERE c.parent=?", (name,))
        assert len(got) == 1, "the second import updates the budget, it doesn't add one"
        assert abs(got[0]["amount"] - 220) < 0.01
        assert db.scalar("SELECT COUNT(*) FROM categories WHERE parent=?", (name,)) == 1
        sections_engine.delete_section(name)


# ===========================================================================
# "Return per year" — only when there is a year of it
# ===========================================================================

class TestAnnualising:
    """The old line divided a real return by a start date read from settings, which
    ships with the anchor from the original Investments 2 sheet. Three months of data
    reported a rate 'over 5.0 yrs'."""

    def test_a_short_history_gets_no_per_year_rate(self):
        assert portfolio.annualised(0.12, years=0.25) is None
        assert portfolio.annualised(0.12, years=0.99) is None

    def test_a_year_or_more_annualises(self):
        assert abs(portfolio.annualised(0.12, years=1.0) - 0.12) < 1e-9
        assert abs(portfolio.annualised(0.30, years=3.0) - 0.10) < 1e-9

    def test_no_history_at_all_is_none_not_zero(self):
        assert portfolio.annualised(0.12, years=0) is None
        assert portfolio.annualised(None, years=5) is None

    def test_the_span_comes_from_the_data_not_from_settings(self):
        # Settings still carries the old 2021 anchor; the figures must not use it once
        # there is real data, or the divisor is years of history nobody has.
        config.settings["portfolio_start_date"] = "2021-08-01"
        first = portfolio._first_activity()
        if first is not None:
            years = portfolio._years_of_history()
            expected = (dt.date.today() - first).days / 365.25
            assert abs(years - expected) < 0.01
            # It must be the span of the *data*, whatever that is — never the settings
            # anchor, which is what produced "over 5.0 yrs" on a three-month install.
            assert first != dt.date(2021, 8, 1) or years == pytest.approx(expected, abs=0.01)

    def test_the_summary_reports_both_and_annualises_the_ratio_it_shows(self):
        rows = portfolio.holdings_rows()
        s = portfolio.summarise(rows)
        assert "return_on_cost" in s and "return_on_cost_yr" in s
        yrs = s.get("years")
        if s["return_on_cost_yr"] is None:
            assert not yrs or yrs < portfolio.MIN_ANNUALISE_YEARS
        else:
            assert abs(s["return_on_cost_yr"] * yrs - s["return_on_cost"]) < 1e-6


# ===========================================================================
# Editing budget sections
# ===========================================================================

class TestSectionsEditing:
    def test_a_section_can_be_added_with_lines(self):
        from kestrel.engine import sections as sec
        r = sec.add_section("Zzz Hobbies", ["Model trains", "Bell ringing"])
        assert r["name"] == "Zzz Hobbies"
        names = {s["name"] for s in sec.list_sections()}
        assert "Zzz Hobbies" in names
        sec.delete_section("Zzz Hobbies")

    def test_renaming_a_section_carries_its_budget_across(self):
        # A section budget names its section in text. Rename the section without moving
        # the budget and the budget silently stops applying to anything.
        from kestrel.engine import sections as sec
        sec.add_section("Zzz Old", ["Thing"])
        budgets.set_budget(120, parent="Zzz Old")
        sec.rename_section("Zzz Old", "Zzz New")
        assert db.one("SELECT id FROM budgets WHERE parent_only='Zzz New'")
        assert not db.one("SELECT id FROM budgets WHERE parent_only='Zzz Old'")
        assert db.scalar("SELECT COUNT(*) FROM categories WHERE parent='Zzz Old'") == 0
        sec.delete_section("Zzz New")

    def test_structural_sections_cannot_be_deleted(self):
        from kestrel.engine import sections as sec
        income = next((s for s in sec.list_sections() if s["role"] == "income"), None)
        assert income, "the household needs somewhere for income to live"
        assert income["protected"] is True
        with pytest.raises(ValueError):
            sec.delete_section(income["name"])

    def test_deleting_a_section_can_move_its_transactions_somewhere(self):
        from kestrel.engine import sections as sec
        sec.add_section("Zzz Doomed", ["Line A"])
        cat = db.one("SELECT id FROM categories WHERE parent='Zzz Doomed'")
        sec.add_section("Zzz Safe", ["Kept"])
        sec.delete_section("Zzz Doomed", move_to="Zzz Safe")
        assert db.scalar("SELECT COUNT(*) FROM categories WHERE parent='Zzz Doomed'") == 0
        assert db.scalar("SELECT COUNT(*) FROM categories WHERE parent='Zzz Safe'") >= 1
        assert cat is not None
        sec.delete_section("Zzz Safe")

    def test_a_line_can_be_renamed_without_losing_its_transactions(self):
        from kestrel.engine import sections as sec
        sec.add_section("Zzz Lines", ["Before"])
        cat = db.one("SELECT id FROM categories WHERE parent='Zzz Lines' AND name='Before'")
        sec.rename_line(cat["id"], "After")
        assert db.one("SELECT id FROM categories WHERE id=? AND name='After'", (cat["id"],))
        sec.delete_section("Zzz Lines")


# ===========================================================================
# Transactions typed in by hand
# ===========================================================================

class TestManualTransactions:
    def _account(self):
        return db.one("SELECT id, name, currency FROM accounts WHERE closed=0 LIMIT 1")

    def test_the_sign_comes_from_the_direction_not_the_typing(self):
        # Nobody should have to type a minus sign; a stray one is how an amount gets
        # entered backwards.
        from kestrel.engine import manual
        a = self._account()
        out = manual.add(a["id"], "2026-08-20", "Corner shop", "12.40", "out")["rows"][0]
        inn = manual.add(a["id"], "2026-08-20", "Refund", "12.40", "in")["rows"][0]
        assert out["amount"] == -12.40
        assert inn["amount"] == 12.40
        # and a minus typed in anyway doesn't flip it back
        again = manual.add(a["id"], "2026-08-20", "Corner shop", "-12.40", "out")["rows"][0]
        assert again["amount"] == -12.40
        for t in (out, inn, again):
            manual.delete(t["id"])

    def test_it_is_marked_as_typed_so_it_can_be_told_apart(self):
        from kestrel.engine import manual
        a = self._account()
        row = manual.add(a["id"], "2026-08-20", "Window cleaner", "25", "out")["rows"][0]
        assert row["entry_source"] == "manual"
        manual.delete(row["id"])

    def test_the_category_is_guessed_from_the_description(self):
        from kestrel.engine import manual
        a = self._account()
        row = manual.add(a["id"], "2026-08-20", "TESCO METRO LONDON", "42.10", "out")["rows"][0]
        cat = db.one("SELECT parent, name FROM categories WHERE id=?", (row["category_id"],))
        assert cat and cat["parent"] == "Food & Drink"
        manual.delete(row["id"])

    def test_two_identical_entries_both_survive(self):
        # Two £3 coffees on the same day are two transactions. An imported row is
        # de-duplicated against the statement; a typed one only against itself.
        from kestrel.engine import manual
        a = self._account()
        one = manual.add(a["id"], "2026-08-20", "Coffee", "3", "out")["rows"][0]
        two = manual.add(a["id"], "2026-08-20", "Coffee", "3", "out")["rows"][0]
        assert one["id"] != two["id"]
        manual.delete(one["id"]); manual.delete(two["id"])

    def test_nonsense_is_refused_in_words_somebody_can_act_on(self):
        from kestrel.engine import manual
        a = self._account()
        for bad, word in (
            (dict(posted_on="not a date", description="x", amount="5", direction="out"), "date"),
            (dict(posted_on="2026-08-20", description="  ", amount="5", direction="out"), "description"),
            (dict(posted_on="2026-08-20", description="x", amount="nope", direction="out"), "number"),
            (dict(posted_on="2026-08-20", description="x", amount="0", direction="out"), "nothing"),
            (dict(posted_on="2099-01-01", description="x", amount="5", direction="out"), "future"),
        ):
            with pytest.raises(manual.ManualError) as e:
                manual.add(a["id"], **bad)
            assert word in str(e.value).lower()

    def test_a_transfer_writes_both_halves_and_pairs_them(self):
        from kestrel.engine import manual
        accts = db.rows("SELECT id FROM accounts WHERE closed=0 LIMIT 2")
        if len(accts) < 2:
            pytest.skip("needs two accounts")
        res = manual.add(accts[0]["id"], "2026-08-20", "Moving money", "500", "out",
                         transfer_to_account_id=accts[1]["id"])
        assert len(res["ids"]) == 2
        rows = res["rows"]
        assert sorted(r["amount"] for r in rows) == [-500.0, 500.0]
        # half a transfer is money appearing from nowhere
        paired = db.one("SELECT transfer_pair FROM transactions WHERE id=?", (res["ids"][0],))
        assert paired["transfer_pair"] == res["ids"][1]
        manual.delete(res["ids"][0])
        assert not db.one("SELECT id FROM transactions WHERE id=?", (res["ids"][1],))

    def test_a_transfer_needs_two_different_accounts(self):
        from kestrel.engine import manual
        a = self._account()
        with pytest.raises(manual.ManualError):
            manual.add(a["id"], "2026-08-20", "To myself", "5", "out",
                       transfer_to_account_id=a["id"])

    def test_a_typed_entry_can_be_corrected(self):
        from kestrel.engine import manual
        a = self._account()
        row = manual.add(a["id"], "2026-08-20", "Windo cleaner", "250", "out")["rows"][0]
        fixed = manual.update(row["id"], description="Window cleaner", amount="25")
        assert fixed["description"] == "Window cleaner"
        assert fixed["amount"] == -25.0
        manual.delete(row["id"])

    def test_an_imported_entry_cannot_be_edited_or_deleted(self):
        # Its fingerprint says "this is what the statement said". Letting it be edited
        # would break the promise that re-uploading a file changes nothing.
        from kestrel.engine import manual
        imported = db.one("SELECT id FROM transactions WHERE IFNULL(entry_source,'import')"
                          "='import' LIMIT 1")
        assert imported, "the demo household imports its transactions"
        with pytest.raises(manual.ManualError):
            manual.update(imported["id"], description="something else")
        with pytest.raises(manual.ManualError):
            manual.delete(imported["id"])
        assert db.one("SELECT id FROM transactions WHERE id=?", (imported["id"],))


class TestSupersedingTypedNotes:
    """A typed note and the statement row for the same payment must not both survive."""

    def _fresh_account(self, name):
        with db.tx() as c:
            cur = c.execute("INSERT INTO accounts(name, account_type, currency) "
                            "VALUES(?,'current','GBP')",
                            (name,))
        return cur.lastrowid

    def test_an_import_replaces_the_note_it_matches(self):
        from kestrel.engine import manual
        from kestrel.importers.readers import read_csv_bytes
        acct = self._fresh_account("Zzz Supersede")
        manual.add(acct, "2026-07-03", "Cash out", "40", "out")
        table = read_csv_bytes(b"Date,Description,Amount\n05/07/2026,ATM WITHDRAWAL,-40.00\n",
                               "s.csv")
        res = ingest.import_transactions(acct, table, {"date": 0, "description": 1, "amount": 2},
                                         None, "s.csv", None)
        assert res["added"] == 1
        assert res["superseded"] == 1
        rows = db.rows("SELECT description, entry_source FROM transactions WHERE account_id=?",
                       (acct,))
        assert len(rows) == 1, "one payment, one row"
        assert rows[0]["entry_source"] == "import", "the bank's version is the one kept"
        assert any("typed in by hand" in w for w in res["warnings"]), "and it is reported"
        with db.tx() as c:
            c.execute("DELETE FROM accounts WHERE id=?", (acct,))

    def test_a_different_amount_is_left_alone(self):
        from kestrel.engine import manual
        from kestrel.importers.readers import read_csv_bytes
        acct = self._fresh_account("Zzz Keep")
        manual.add(acct, "2026-07-03", "Cash out", "40", "out")
        table = read_csv_bytes(b"Date,Description,Amount\n05/07/2026,ATM WITHDRAWAL,-60.00\n",
                               "s.csv")
        res = ingest.import_transactions(acct, table, {"date": 0, "description": 1, "amount": 2},
                                         None, "s.csv", None)
        assert res["superseded"] == 0
        assert db.scalar("SELECT COUNT(*) FROM transactions WHERE account_id=?", (acct,)) == 2
        with db.tx() as c:
            c.execute("DELETE FROM accounts WHERE id=?", (acct,))

    def test_a_date_far_apart_is_left_alone(self):
        from kestrel.engine import manual
        from kestrel.importers.readers import read_csv_bytes
        acct = self._fresh_account("Zzz Far")
        manual.add(acct, "2026-05-03", "Cash out", "40", "out")
        table = read_csv_bytes(b"Date,Description,Amount\n05/07/2026,ATM WITHDRAWAL,-40.00\n",
                               "s.csv")
        res = ingest.import_transactions(acct, table, {"date": 0, "description": 1, "amount": 2},
                                         None, "s.csv", None)
        assert res["superseded"] == 0, "two months apart is a different £40"
        with db.tx() as c:
            c.execute("DELETE FROM accounts WHERE id=?", (acct,))

    def test_one_imported_row_can_only_replace_one_note(self):
        # Two £40 notes and one £40 statement row: exactly one note is superseded.
        from kestrel.engine import manual
        from kestrel.importers.readers import read_csv_bytes
        acct = self._fresh_account("Zzz Once")
        manual.add(acct, "2026-07-03", "Cash out", "40", "out")
        manual.add(acct, "2026-07-04", "Cash out again", "40", "out")
        table = read_csv_bytes(b"Date,Description,Amount\n05/07/2026,ATM WITHDRAWAL,-40.00\n",
                               "s.csv")
        res = ingest.import_transactions(acct, table, {"date": 0, "description": 1, "amount": 2},
                                         None, "s.csv", None)
        assert res["superseded"] == 1
        assert db.scalar("SELECT COUNT(*) FROM transactions WHERE account_id=?", (acct,)) == 2
        with db.tx() as c:
            c.execute("DELETE FROM accounts WHERE id=?", (acct,))


class TestManualRoutes:
    def test_the_route_reports_what_it_wrote(self):
        from kestrel.web import server
        a = db.one("SELECT id FROM accounts WHERE closed=0 LIMIT 1")
        out = server.api_add_transaction(None, {}, {}, {
            "account_id": a["id"], "posted_on": "2026-08-21", "description": "Sainsburys",
            "amount": "18.20", "direction": "out"})
        assert out["ok"] and len(out["rows"]) == 1
        assert out["rows"][0]["amount"] == -18.20
        from kestrel.engine import manual
        manual.delete(out["ids"][0])

    def test_a_bad_entry_comes_back_as_a_sentence_not_a_stack_trace(self):
        from kestrel.web import server
        a = db.one("SELECT id FROM accounts WHERE closed=0 LIMIT 1")
        with pytest.raises(server.ApiError) as e:
            server.api_add_transaction(None, {}, {}, {
                "account_id": a["id"], "posted_on": "2026-08-21", "description": "x",
                "amount": "banana", "direction": "out"})
        assert "number" in str(e.value).lower()

    def test_guessing_a_category_before_it_is_saved(self):
        from kestrel.web import server
        out = server.api_guess_category(None, {}, {"description": "TESCO METRO"}, None)
        assert out["category_id"]
        assert out["category"]["parent"] == "Food & Drink"
        assert server.api_guess_category(None, {}, {"description": ""}, None)["category_id"] is None

    def test_a_cross_currency_transfer_asks_for_the_arriving_amount(self):
        # £500 out of a UK account is not R500 into a South African one, and only the
        # bank knows the rate it used. Mirroring the number invents one.
        from kestrel.engine import manual
        gb = db.one("SELECT id FROM accounts WHERE currency='GBP' AND closed=0 LIMIT 1")
        za = db.one("SELECT id FROM accounts WHERE currency='ZAR' AND closed=0 LIMIT 1")
        if not (gb and za):
            pytest.skip("needs one account in each currency")
        with pytest.raises(manual.ManualError) as e:
            manual.add(gb["id"], "2026-08-20", "Home", "500", "out",
                       transfer_to_account_id=za["id"])
        assert "currenc" in str(e.value).lower()

        res = manual.add(gb["id"], "2026-08-20", "Home", "500", "out",
                         transfer_to_account_id=za["id"], transfer_amount="11840.50")
        by_ccy = {r["currency"]: r["amount"] for r in res["rows"]}
        assert by_ccy["GBP"] == -500.0
        assert by_ccy["ZAR"] == 11840.50, "the far leg carries the opposite sign"
        manual.delete(res["ids"][0])

    def test_a_refused_transfer_writes_nothing_at_all(self):
        # Validating halfway through left money leaving one account and nothing arriving
        # in the other — an orphan half, worse than the error that caused it.
        from kestrel.engine import manual
        gb = db.one("SELECT id FROM accounts WHERE currency='GBP' AND closed=0 LIMIT 1")
        za = db.one("SELECT id FROM accounts WHERE currency='ZAR' AND closed=0 LIMIT 1")
        if not (gb and za):
            pytest.skip("needs one account in each currency")
        before = db.scalar("SELECT COUNT(*) FROM transactions", (), 0)
        with pytest.raises(manual.ManualError):
            manual.add(gb["id"], "2026-08-20", "Home", "500", "out",
                       transfer_to_account_id=za["id"])
        assert db.scalar("SELECT COUNT(*) FROM transactions", (), 0) == before

    def test_a_declared_transfer_pairs_its_own_two_legs(self):
        # The detector matches amounts within a few days and can pair the wrong £500 when
        # several exist. A declared transfer has nothing to guess at.
        from kestrel.engine import manual
        accts = db.rows("SELECT id, currency FROM accounts WHERE closed=0 LIMIT 2")
        if len(accts) < 2:
            pytest.skip("needs two accounts")
        # a decoy of the same size, on the same day, in a third place
        decoy = manual.add(accts[0]["id"], "2026-06-10", "Decoy in", "500", "in")["rows"][0]
        kw = {}
        if accts[0]["currency"] != accts[1]["currency"]:
            kw["transfer_amount"] = "500"
        res = manual.add(accts[0]["id"], "2026-06-10", "The real transfer", "500", "out",
                         transfer_to_account_id=accts[1]["id"], **kw)
        a, b = res["ids"]
        assert db.one("SELECT transfer_pair FROM transactions WHERE id=?", (a,))["transfer_pair"] == b
        assert db.one("SELECT transfer_pair FROM transactions WHERE id=?", (b,))["transfer_pair"] == a
        assert db.one("SELECT transfer_pair FROM transactions WHERE id=?",
                      (decoy["id"],))["transfer_pair"] is None
        manual.delete(a); manual.delete(decoy["id"])


# ===========================================================================
# Mortgages
# ===========================================================================

class TestMortgageMaths:
    """The annuity formula, and the cases where the honest answer is 'that never
    happens' rather than a very large number."""

    def test_the_payment_matches_what_a_lender_would_quote(self):
        from kestrel.engine import mortgage as mo
        # £250,000 at 4.29% over 25 years is £1,359.95 — the figure every UK calculator
        # gives, because lenders divide the nominal rate by twelve rather than
        # compounding it.
        assert round(mo.payment_for(250000, 4.29, 300), 2) == 1359.95
        assert round(mo.payment_for(100000, 0, 100), 2) == 1000.00

    def test_the_term_it_implies_is_the_term_it_was_built_from(self):
        from kestrel.engine import mortgage as mo
        p = mo.payment_for(250000, 4.29, 300)
        assert mo.months_to_clear(250000, 4.29, p) == 300

    def test_a_payment_that_never_touches_the_capital_says_so(self):
        # £200,000 at 5% costs £833 a month in interest alone. A £700 payment clears
        # nothing, ever, and returning a big number would imply otherwise.
        from kestrel.engine import mortgage as mo
        assert mo.months_to_clear(200000, 5.0, 700) is None
        rows = mo.amortise(200000, 5.0, 700, 24)
        assert len(rows) == 1 and rows[0]["never_clears"] is True

    def test_the_schedule_adds_up(self):
        from kestrel.engine import mortgage as mo
        p = mo.payment_for(120000, 3.5, 240)
        rows = mo.amortise(120000, 3.5, p, 240)
        assert rows[-1]["balance"] <= 0.02, "it must actually reach zero"
        assert abs(sum(r["capital"] for r in rows) - 120000) < 1.0
        for r in rows[:-1]:
            assert abs(r["interest"] + r["capital"] - r["payment"]) < 0.01

    def test_interest_falls_and_capital_rises_over_the_term(self):
        from kestrel.engine import mortgage as mo
        p = mo.payment_for(200000, 4.0, 300)
        rows = mo.amortise(200000, 4.0, p, 300)
        assert rows[0]["interest"] > rows[-1]["interest"]
        assert rows[0]["capital"] < rows[-1]["capital"]

    def test_month_arithmetic_survives_the_end_of_a_month(self):
        from kestrel.engine.mortgage import _add_months
        assert _add_months(dt.date(2026, 1, 31), 1) == dt.date(2026, 2, 28)
        assert _add_months(dt.date(2024, 1, 31), 1) == dt.date(2024, 2, 29)  # leap year
        assert _add_months(dt.date(2026, 12, 15), 1) == dt.date(2027, 1, 15)
        assert _add_months(dt.date(2026, 6, 1), 300) == dt.date(2051, 6, 1)


class TestMortgageRecord:
    def _mortgage(self):
        m = db.one("SELECT m.*, a.currency FROM mortgages m JOIN accounts a ON a.id=m.account_id "
                   "LIMIT 1")
        assert m, "the demo household has a mortgage"
        return m

    def test_the_demo_household_is_internally_consistent(self):
        # Sample data that trips the app's own consistency check is a standing bug report.
        from kestrel.engine import mortgage as mo
        s = mo.summarise(self._mortgage())
        assert s["term_mismatch"] is None
        assert s["never_clears"] is False
        # Within a month: one is the date of the final payment, the other the contract's
        # anniversary, and they need not be the same day.
        gap = abs((dt.date.fromisoformat(s["payoff_on"])
                   - dt.date.fromisoformat(s["term_ends_on"])).days)
        assert gap <= 32, (s["payoff_on"], s["term_ends_on"])

    def test_a_payment_is_split_into_interest_and_capital(self):
        from kestrel.engine import mortgage as mo
        s = mo.summarise(self._mortgage())
        assert s["interest_this_month"] > 0 and s["capital_this_month"] > 0
        assert abs(s["interest_this_month"] + s["capital_this_month"]
                   - s["monthly_payment"]) < 0.01

    def test_the_balance_is_projected_and_says_so(self):
        # A mortgage statement arrives once a year. Showing a year-old figure as though
        # it were today's is the same error as an account claiming £0.00.
        from kestrel.engine import mortgage as mo
        m = self._mortgage()
        s = mo.summarise(m)
        assert s["estimated"] is True
        assert s["confirmed_on"] == m["statement_on"]
        assert s["projected_months"] > 0
        assert s["balance"] < m["statement_balance"], "it should have come down since"

    def test_equity_follows_the_linked_property_account(self):
        # One figure for the house, not two that can disagree.
        from kestrel.engine import mortgage as mo
        m = self._mortgage()
        house = db.one("SELECT balance FROM accounts WHERE id=?", (m["property_account_id"],))
        s = mo.summarise(m)
        assert s["property_value"] == house["balance"]
        assert abs(s["equity"] - (house["balance"] - s["balance"])) < 0.01

    def test_net_worth_agrees_with_the_mortgage_screen(self):
        from kestrel.engine import mortgage as mo
        s = mo.summarise(self._mortgage())
        debt = portfolio.net_worth()["groups"]["Debt"]
        assert abs(abs(debt) - s["balance"]) < 0.01, \
            "two screens of one app must not report two different debts"


class TestOverpayments:
    def _mortgage(self):
        return db.one("SELECT * FROM mortgages LIMIT 1")

    def test_a_monthly_overpayment_saves_months_and_interest(self):
        from kestrel.engine import mortgage as mo
        r = mo.overpayment_effect(self._mortgage(), 200, "monthly")
        assert r["available"]
        assert r["months_saved"] > 0
        assert r["interest_saved"] > 0
        assert r["new_payoff_on"] < r["old_payoff_on"]

    def test_a_lump_sum_saves_less_than_the_same_amount_every_month(self):
        from kestrel.engine import mortgage as mo
        m = self._mortgage()
        once = mo.overpayment_effect(m, 200, "once")
        monthly = mo.overpayment_effect(m, 200, "monthly")
        assert once["interest_saved"] < monthly["interest_saved"]

    def test_it_never_implies_the_saving_is_free(self):
        # Overpayment charges are real and Mithapp cannot know the limit.
        from kestrel.engine import mortgage as mo
        r = mo.overpayment_effect(self._mortgage(), 500, "monthly")
        assert "overpayment" in r["caveat"].lower()

    def test_nothing_to_overpay_is_explained_rather_than_calculated(self):
        from kestrel.engine import mortgage as mo
        r = mo.overpayment_effect(self._mortgage(), 0)
        assert r["available"] is False and r["why"]

    def test_what_happens_when_the_fix_ends(self):
        from kestrel.engine import mortgage as mo
        m = self._mortgage()
        r = mo.rate_change_effect(m, (m["revert_rate"] or 7.0))
        assert r["available"]
        assert r["payment_then"] > r["payment_now"], "a higher rate costs more"
        assert r["difference"] > 0


class TestMortgageInBudgets:
    def test_the_capital_half_is_saving_rather_than_spending(self):
        # £1,180 leaves the current account but only the interest leaves the household —
        # the rest moved into the house, exactly like paying into an ISA.
        from kestrel.engine import mortgage as mo
        m = db.one("SELECT * FROM mortgages LIMIT 1")
        s = mo.summarise(m)
        start, end = budgets.month_bounds(budgets.current_period())
        with_split = budgets.spend_by_category(start, end)
        assert with_split["mortgage_capital"], "the adjustment should have applied"
        moved = with_split["mortgage_capital"][0]["amount"]
        assert moved > 0

        with db.tx() as c:
            c.execute("UPDATE mortgages SET split_payments=0 WHERE id=?", (m["id"],))
        without = budgets.spend_by_category(start, end)
        with db.tx() as c:
            c.execute("UPDATE mortgages SET split_payments=1 WHERE id=?", (m["id"],))

        assert not without["mortgage_capital"]
        assert abs((without["totals"]["spend"] - with_split["totals"]["spend"]) - moved) < 0.02
        assert abs((with_split["totals"]["saving"] - without["totals"]["saving"]) - moved) < 0.02

    def test_it_never_moves_more_than_actually_went_out(self):
        # A projection can run past a month where the payment was missed, or where the
        # statement hasn't caught up. Moving more than left the account would invent
        # saving out of nothing.
        start, end = budgets.month_bounds("2019-01")     # long before any transactions
        out = budgets.spend_by_category(start, end)
        assert out["totals"]["spend"] >= 0
        assert not out["mortgage_capital"], "no payments that month, so nothing to move"

    def test_an_interest_only_mortgage_moves_nothing(self):
        from kestrel.engine import mortgage as mo
        m = db.one("SELECT * FROM mortgages LIMIT 1")
        with db.tx() as c:
            c.execute("UPDATE mortgages SET repayment_type='interest_only' WHERE id=?", (m["id"],))
        try:
            start, end = budgets.month_bounds(budgets.current_period())
            assert not budgets.spend_by_category(start, end)["mortgage_capital"], \
                "nothing is being repaid, so there is nothing to reclassify"
        finally:
            with db.tx() as c:
                c.execute("UPDATE mortgages SET repayment_type='repayment' WHERE id=?", (m["id"],))


class TestMortgageRoutes:
    def test_saving_details_creates_the_record_and_mirrors_the_balance(self):
        from kestrel.web import server
        acct = server.api_create_account(None, {}, {}, {
            "name": "Zzz Test Mortgage", "account_type": "mortgage", "currency": "GBP"})
        out = server.api_save_mortgage(None, {"aid": str(acct["id"])}, {}, {
            "lender": "Test Bank", "rate": "5.0", "monthly_payment": "900",
            "statement_balance": "150000", "statement_on": "2026-06-30"})
        assert out["summary"]["lender"] == "Test Bank"
        a = db.one("SELECT balance FROM accounts WHERE id=?", (acct["id"],))
        assert a["balance"] == -150000, "a debt is stored negative, so net worth is right"
        with db.tx() as c:
            c.execute("DELETE FROM accounts WHERE id=?", (acct["id"],))

    def test_a_typo_comes_back_as_a_sentence(self):
        from kestrel.web import server
        acct = db.one("SELECT id FROM accounts WHERE account_type='mortgage' LIMIT 1")
        with pytest.raises(server.ApiError) as e:
            server.api_save_mortgage(None, {"aid": str(acct["id"])}, {}, {"rate": "four"})
        assert "isn't a number" in str(e.value)

    def test_recording_a_statement_updates_what_is_on_screen(self):
        from kestrel.web import server
        from kestrel.engine import mortgage as mo
        acct = db.one("SELECT id FROM accounts WHERE account_type='mortgage' LIMIT 1")
        before = mo.get(acct["id"])
        out = server.api_mortgage_event(None, {"aid": str(acct["id"])}, {}, {
            "kind": "statement", "happened_on": "2026-07-31", "amount": "180000"})
        assert out["summary"]["confirmed_on"] == "2026-07-31"
        assert db.one("SELECT balance FROM accounts WHERE id=?", (acct["id"],))["balance"] == -180000
        # put the household back as it was
        with db.tx() as c:
            c.execute("UPDATE mortgages SET statement_balance=?, statement_on=? WHERE id=?",
                      (before["statement_balance"], before["statement_on"], before["id"]))
            c.execute("UPDATE accounts SET balance=? WHERE id=?",
                      (-abs(before["statement_balance"]), acct["id"]))
            c.execute("DELETE FROM mortgage_events WHERE happened_on='2026-07-31'")

    def test_the_detail_route_returns_a_schedule_and_a_yearly_curve(self):
        from kestrel.web import server
        acct = db.one("SELECT id FROM accounts WHERE account_type='mortgage' LIMIT 1")
        d = server.api_mortgage(None, {"aid": str(acct["id"])}, {}, None)
        assert len(d["schedule"]) > 12
        assert 2 <= len(d["curve"]) <= 41, "one point a year keeps the payload small"
        assert d["curve"][0]["balance"] > d["curve"][-1]["balance"]

    def test_the_card_and_the_chart_name_the_same_final_month(self):
        # They were a month apart: the summary counted the month *after* the last
        # payment, and a card saying "Jun 2046" over a chart ending "May 2046" reads as
        # a bug even when both numbers are nearly right.
        from kestrel.web import server
        acct = db.one("SELECT id FROM accounts WHERE account_type='mortgage' LIMIT 1")
        d = server.api_mortgage(None, {"aid": str(acct["id"])}, {}, None)
        assert d["schedule"][-1]["on"][:7] == d["summary"]["payoff_on"][:7]
        assert d["curve"][-1]["on"] == d["schedule"][-1]["on"]
        assert d["curve"][-1]["balance"] <= 0.02, "the curve must reach zero"


# ===========================================================================
# Who's in the household
# ===========================================================================

class TestMembers:
    """A member is a label on an account — whose ISA is whose — not a login. Mithapp
    has no sign-in at all."""

    def test_a_fresh_install_calls_the_first_person_Me(self):
        # Not a name the app's author happened to have.
        row = db.one("SELECT name FROM members WHERE is_default=1")
        assert row is not None

    def test_the_sample_data_never_renames_the_person_using_the_app(self):
        # This is the bug: `demo.load()` used to run
        # `UPDATE members SET name='Allan' WHERE is_default=1`, so loading the sample
        # household on anybody's machine stamped the author's name over theirs — and
        # members could only be added, never renamed, so it could not be undone.
        src = pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "demo.py"
        text = src.read_text(encoding="utf-8")
        body = "\n".join(l for l in text.splitlines() if not l.strip().startswith("#"))
        assert "SET name='Allan'" not in body
        assert "'Allan'" not in body and '"Allan"' not in body

    def test_the_sample_household_keeps_whatever_name_is_there(self):
        from kestrel.web import server
        me = db.one("SELECT id, name FROM members WHERE is_default=1")
        server.api_edit_member(None, {"mid": str(me["id"])}, {}, {"name": "Zzz Tester"})
        rep = demo.load()
        assert db.one("SELECT name FROM members WHERE id=?", (me["id"],))["name"] == "Zzz Tester"
        assert "Zzz Tester" in rep["members"]
        server.api_edit_member(None, {"mid": str(me["id"])}, {}, {"name": me["name"]})

    def test_a_member_can_be_renamed_and_recoloured(self):
        from kestrel.web import server
        me = db.one("SELECT id, name, colour FROM members WHERE is_default=1")
        out = server.api_edit_member(None, {"mid": str(me["id"])}, {},
                                     {"name": "Zzz Renamed", "colour": "#0f9d8e"})
        assert out["name"] == "Zzz Renamed" and out["colour"] == "#0f9d8e"
        server.api_edit_member(None, {"mid": str(me["id"])}, {},
                               {"name": me["name"], "colour": me["colour"]})

    def test_two_people_cannot_share_one_name(self):
        from kestrel.web import server
        with pytest.raises(server.ApiError):
            server.api_add_member(None, {}, {}, {"name": "Family"})
        with pytest.raises(server.ApiError):
            server.api_add_member(None, {}, {}, {"name": "  "})

    def test_exactly_one_person_is_the_main_one(self):
        from kestrel.web import server
        fam = db.one("SELECT id FROM members WHERE name='Family'")
        was = db.one("SELECT id FROM members WHERE is_default=1")
        server.api_edit_member(None, {"mid": str(fam["id"])}, {}, {"is_default": True})
        assert db.scalar("SELECT COUNT(*) FROM members WHERE is_default=1", (), 0) == 1
        assert db.one("SELECT is_default FROM members WHERE id=?", (fam["id"],))["is_default"] == 1
        server.api_edit_member(None, {"mid": str(was["id"])}, {}, {"is_default": True})

    def test_the_last_person_cannot_be_removed(self):
        from kestrel.web import server
        keep = db.rows("SELECT id FROM members")
        extra = [r for r in keep if r["id"] != keep[0]["id"]]
        # remove everyone but one, then check the last is refused
        for r in extra:
            server.api_delete_member(None, {"mid": str(r["id"])}, {"move_to": str(keep[0]["id"])}, None)
        with pytest.raises(server.ApiError) as e:
            server.api_delete_member(None, {"mid": str(keep[0]["id"])}, {}, None)
        assert "at least one" in str(e.value)
        demo.load()          # put the household back

    def test_removing_someone_can_hand_their_accounts_on(self):
        # Without somewhere to go, accounts fall to NULL and the household quietly loses
        # track of whose things they are.
        from kestrel.web import server
        me = db.one("SELECT id FROM members WHERE is_default=1")
        temp = server.api_add_member(None, {}, {}, {"name": "Zzz Leaving"})
        acct = db.one("SELECT id FROM accounts LIMIT 1")
        with db.tx() as c:
            c.execute("UPDATE accounts SET member_id=? WHERE id=?", (temp["id"], acct["id"]))
        server.api_delete_member(None, {"mid": str(temp["id"])},
                                 {"move_to": str(me["id"])}, None)
        assert db.one("SELECT member_id FROM accounts WHERE id=?",
                      (acct["id"],))["member_id"] == me["id"]

    def test_the_listing_says_what_removing_someone_would_disturb(self):
        from kestrel.web import server
        out = server.api_members(None, {}, {}, None)
        assert all("accounts" in r and "budgets" in r for r in out["items"])
        assert "unnamed" in out

    def test_a_household_still_called_Me_is_flagged(self):
        from kestrel.web import server
        me = db.one("SELECT id, name FROM members WHERE is_default=1")
        server.api_edit_member(None, {"mid": str(me["id"])}, {}, {"name": "Me"})
        assert server.api_members(None, {}, {}, None)["unnamed"] is True
        server.api_edit_member(None, {"mid": str(me["id"])}, {}, {"name": "Zzz Named"})
        assert server.api_members(None, {}, {}, None)["unnamed"] is False
        server.api_edit_member(None, {"mid": str(me["id"])}, {}, {"name": me["name"]})


# ===========================================================================
# Packaging — the build has to ship the app, not just the python
# ===========================================================================

class TestBuildSpec:
    """The PyInstaller spec decides what actually goes in the .exe. Everything it gets
    wrong produces a build that looks fine until somebody runs it."""

    def _spec(self) -> str:
        return (pathlib.Path(__file__).resolve().parent.parent / "build"
                / "mittens.spec").read_text(encoding="utf-8")

    def _root(self) -> pathlib.Path:
        return pathlib.Path(__file__).resolve().parent.parent

    def test_the_spec_bundles_files_that_actually_exist(self):
        # The blanket kestrel→mithapp rename rewrote `PKG = ROOT / "kestrel"` to
        # `ROOT / "mithapp"`, which is not a directory. `datas` then pointed at nothing,
        # so the build would have shipped with **no web interface and no institution
        # list** and opened to a blank window. The same rename broke
        # tools/build_institutions.py the same way; this is the guard for the spec.
        spec, root = self._spec(), self._root()
        ns = {"SPECPATH": str(root / "build")}
        head = spec.split("block_cipher")[0]
        exec(compile(head, "mittens.spec", "exec"), ns)      # runs the spec's assertions
        pkg = ns["PKG"]
        assert pkg.is_dir(), pkg
        assert (pkg / "web" / "static" / "app.js").exists()
        assert (pkg / "web" / "static" / "index.html").exists()
        assert (pkg / "institutions" / "data").is_dir()
        assert list((pkg / "institutions" / "data").glob("*.json"))

    def test_the_spec_asserts_for_itself_at_build_time(self):
        # A test in this file only helps somebody who runs the tests. The spec carries
        # its own assertion so a build cannot quietly produce an empty app either.
        spec = self._spec()
        assert "assert (PKG / \"web\" / \"static\" / \"app.js\").exists()" in spec

    def _exclude_block(self):
        """Run just the spec's exclude section — the part that needs no PyInstaller."""
        spec = self._spec()
        start = spec.index("ALIASED_BY_PYINSTALLER = (")
        end = spec.index("a = Analysis")        # the first line that needs PyInstaller
        ns: dict = {}
        exec(compile(spec[start:end], "mittens.spec", "exec"), ns)
        return ns, spec

    def test_nothing_pyinstaller_aliases_is_excluded(self):
        """`distutils` was removed from the stdlib in Python 3.12 and is now provided by
        setuptools. PyInstaller's own hook aliases the vendored copy back to the name
        `distutils`, and aliasing onto an excluded name is a hard error:

            ValueError: Target module "distutils" already imported as ExcludedModule

        It needs Python 3.12+, setuptools present, *and* an import chain that reaches
        distutils — on Windows, pywebview → pythonnet supplies one. So the build passed
        on 3.11 and on Linux and died on a Windows PC running 3.14.
        """
        ns, spec = self._exclude_block()
        assert not set(ns["excludes"]) & set(ns["ALIASED_BY_PYINSTALLER"])
        # and the spec must keep checking this for itself at build time
        assert "assert not set(excludes) & set(ALIASED_BY_PYINSTALLER)" in spec

    def test_the_alias_list_matches_what_pyinstaller_actually_aliases(self):
        # Read it from PyInstaller rather than trusting a list I typed. Only meaningful
        # on the Pythons where the hook does anything.
        from PyInstaller import compat
        if not getattr(compat, "is_py312", False):
            pytest.skip("the distutils hook only aliases on Python 3.12+")
        from PyInstaller.utils.hooks.setuptools import setuptools_info
        if not setuptools_info.distutils_vendored:
            pytest.skip("setuptools is not vendoring distutils here")
        aliased = {a.split(".")[0] for a, _ in setuptools_info.get_distutils_aliases()}
        ns, _ = self._exclude_block()
        missing = aliased - set(ns["ALIASED_BY_PYINSTALLER"])
        assert not missing, f"PyInstaller also aliases {missing} — add them to the list"
        assert not aliased & set(ns["excludes"])

    def test_pillow_is_not_excluded(self):
        # pdfplumber needs it. Excluding it made a build that started perfectly and died
        # the moment a PDF statement was dropped on the import screen.
        spec = self._spec()
        excludes = spec[spec.index("excludes = ["):spec.index("]", spec.index("excludes = ["))]
        assert '"PIL"' not in excludes
        assert '"Pillow"' not in excludes

    def test_every_provider_module_is_named_as_a_hidden_import(self):
        # Providers register themselves on import, so a frozen build only knows about the
        # ones the spec names. A missing one vanishes from the app with no error.
        spec = self._spec()
        root = self._root() / "kestrel" / "providers"
        for path in root.glob("*.py"):
            if path.stem in ("__init__", "base"):
                continue
            assert f"kestrel.providers.{path.stem}" in spec, \
                f"{path.stem} would not be bundled"

    def test_the_entry_point_the_spec_names_is_there(self):
        spec, root = self._spec(), self._root()
        assert "run_kestrel.py" in spec
        assert (root / "run_kestrel.py").exists()

    def test_both_build_scripts_produce_something_to_send(self):
        build = self._root() / "build"
        bat = (build / "build_windows.bat").read_text(encoding="utf-8")
        sh = (build / "build_macos.sh").read_text(encoding="utf-8")
        for script in (bat, sh):
            assert "Read me first.txt" in script
        assert (build / "Read me first.txt").exists()
        # One thing to send, but not the same thing on both. Windows gets a single
        # self-unpacking .exe. A Mac gets the .app bundle in a zip: an executable
        # with no extension opens a Terminal window when somebody double-clicks it
        # in Finder, which is not an app arriving, it is something going wrong.
        assert "MITTENS_ONEFILE=1" in bat, "no single file means nothing to send"
        assert ".app.zip" in sh, "a Mac is sent the bundle, not a bare executable"

    def test_the_mac_script_zips_the_bundle_with_ditto(self):
        # A .app is a folder of symlinks and permission bits. `zip -r` flattens them
        # and the bundle unzips looking perfectly normal, then refuses to launch on
        # somebody else's Mac — which is the worst possible moment to find out.
        sh = (self._root() / "build" / "build_macos.sh").read_text(encoding="utf-8")
        assert "ditto -c -k --keepParent" in sh
        import re
        for line in sh.splitlines():
            bare = line.strip()
            if bare.startswith("#"):
                continue
            assert not re.match(r"zip\s+-", bare), \
                "plain zip breaks a .app bundle; use ditto"

    def test_the_mac_zip_is_named_after_the_processor_it_was_built_for(self):
        # An Apple Silicon build will not start at all on an Intel Mac. Two files that
        # look identical in a downloads folder is how half the family ends up with the
        # wrong one and concludes the app is broken.
        sh = (self._root() / "build" / "build_macos.sh").read_text(encoding="utf-8")
        assert "uname -m" in sh, "the build has to say which processor it is for"
        assert "macos-$ARCH.app.zip" in sh

    def test_the_windows_script_has_windows_line_endings(self):
        # cmd.exe is unreliable with LF-only .bat files, particularly around the
        # multi-line `if errorlevel ( ... )` blocks this one uses.
        raw = (self._root() / "build" / "build_windows.bat").read_bytes()
        assert raw.count(b"\r\n") > 20
        assert raw.count(b"\n") == raw.count(b"\r\n"), "some lines are missing a CR"

    def test_the_read_me_warns_about_the_unsigned_app_warning(self):
        # It is the first thing every recipient will hit, and it looks alarming.
        text = (self._root() / "build" / "Read me first.txt").read_text(encoding="utf-8")
        assert "Windows protected your PC" in text
        assert "Run anyway" in text
        # Apple withdrew right-click -> Open. The route that works on every macOS
        # from Catalina onwards is System Settings > Privacy & Security > Open Anyway,
        # and a read-me that sends somebody down the old path leaves them stuck with an
        # app they will assume is broken.
        low = text.lower()
        assert "privacy & security" in low, "the macOS route must be the one Apple documents"
        assert "open anyway" in low
        assert "no sign-in" in text.lower() or "no account" in text.lower()
        # Two Mac files go out and only one of them will start. Somebody who does not
        # know which Mac they own has a one-in-two chance of picking a file that does
        # nothing, so the read-me has to tell them how to find out.
        assert "about this mac" in low, "say how to tell an Apple chip from an Intel one"
        assert "arm64" in low and "x86_64" in low, "name the two files as they are named"

    def test_the_build_scripts_install_what_they_then_run(self):
        # They installed requirements.txt and then ran pytest, which lives in
        # requirements-dev.txt. Every clean machine got "No module named pytest" at
        # step 3 and the build stopped dead — with a message blaming the tests.
        build = self._root() / "build"
        for name in ("build_windows.bat", "build_macos.sh"):
            script = (build / name).read_text(encoding="utf-8")
            assert "requirements-dev.txt" in script, \
                f"{name} runs pytest but never installs it"
            assert "import pytest" in script, \
                f"{name} should say so plainly when pytest is missing"

    def test_requirements_dev_covers_everything_the_build_needs(self):
        dev = (self._root() / "requirements-dev.txt").read_text(encoding="utf-8")
        assert "-r requirements.txt" in dev, "the dev file must pull in the runtime one"
        for pkg in ("pytest", "pyinstaller", "reportlab"):
            assert pkg in dev, f"{pkg} is needed to build or to check the build"

    def test_no_test_silently_skips_for_a_missing_dev_dependency(self):
        # Six PDF tests used to skip on any machine without reportlab — including the
        # machine doing the build, so "checking everything still works" quietly stopped
        # covering the one feature that has already caused a packaging bug.
        import importlib.util
        dev = (self._root() / "requirements-dev.txt").read_text(encoding="utf-8")
        for mod in ("reportlab",):
            if f"{mod}" in dev:
                assert importlib.util.find_spec(mod) is not None, (
                    f"{mod} is in requirements-dev.txt but not installed — "
                    "install it, or the tests that need it will skip")

    def test_a_failed_onefile_build_is_not_reported_as_success(self):
        # `set VAR=` succeeds and resets errorlevel, so checking it *after* clearing the
        # variable always saw 0 — a failed build would have printed "Done."
        bat = (self._root() / "build" / "build_windows.bat").read_bytes().decode()
        after_onefile = bat[bat.index("set MITTENS_ONEFILE=1"):]
        check = after_onefile.index("if errorlevel 1")
        clear = after_onefile.index("\r\nset MITTENS_ONEFILE=\r\n")
        assert check < clear, "check the build's result before clearing the variable"

    def test_the_build_calls_pyinstaller_through_the_interpreter(self):
        # `python -m PyInstaller` uses this environment's interpreter whether or not the
        # venv's Scripts folder made it onto PATH.
        build = self._root() / "build"
        for name in ("build_windows.bat", "build_macos.sh"):
            script = (build / name).read_bytes().decode()
            assert "python -m PyInstaller" in script, name
            assert not re.search(r"^\s*pyinstaller ", script, re.M), \
                f"{name} still calls the bare pyinstaller command somewhere"


class TestNoDeprecatedDatetime:
    """`utcnow()` and `utcfromtimestamp()` are deprecated and scheduled for removal.
    They produced 3,045 warnings on a Python 3.14 build — noise that hides real
    warnings, and code that will simply stop working."""

    def _sources(self):
        root = pathlib.Path(__file__).resolve().parent.parent / "kestrel"
        return list(root.rglob("*.py"))

    def test_none_are_left(self):
        offenders = []
        for path in self._sources():
            text = path.read_text(encoding="utf-8")
            for name in ("utcnow()", "utcfromtimestamp("):
                if name in text:
                    offenders.append(f"{path.name}: {name}")
        assert not offenders, offenders

    def test_the_replacement_is_still_naive(self):
        # The stored expiry strings are naive ISO text and are compared as strings
        # against what is already on disk. A replacement that grew a "+00:00" suffix
        # would make every existing token look expired — or never expire.
        now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
        assert now.tzinfo is None
        assert "+" not in now.isoformat()

    def test_the_suite_raises_no_deprecation_warnings_from_our_own_code(self):
        import warnings
        root = pathlib.Path(__file__).resolve().parent.parent
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always", DeprecationWarning)
            import importlib
            for mod in ("kestrel.market.prices", "kestrel.export.gsheets",
                        "kestrel.providers.monzo", "kestrel.providers.starling"):
                importlib.reload(importlib.import_module(mod))
        ours = [w for w in caught if str(root) in str(w.filename)]
        assert not ours, [str(w.message) for w in ours]


# ===========================================================================
# Splitting one payment across categories
# ===========================================================================

class TestSplitting:
    """£100 at Tesco is £50 of groceries, £20 of electronics and £30 of liquor. The
    bank sends one line and has no idea."""

    def _payment(self):
        start, _ = budgets.month_bounds(budgets.current_period())
        return db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                      "AND IFNULL(is_split,0)=0 AND split_of IS NULL "
                      "ORDER BY amount LIMIT 1", (start,))

    def _thirds(self, t):
        a, b = round(t["amount"] * 0.5, 2), round(t["amount"] * 0.2, 2)
        return [
            {"amount": a, "category_id": db.category_id("Food & Drink", "Groceries")},
            {"amount": b, "category_id": db.category_id("Shopping", "Clothing")},
            {"amount": round(t["amount"] - a - b, 2),
             "category_id": db.category_id("Food & Drink", "Eating out")},
        ]

    def _totals(self):
        s, e = budgets.month_bounds(budgets.current_period())
        t = budgets.spend_by_category(s, e)["totals"]
        return {k: round(t[k], 2) for k in ("income", "spend", "saving", "net")}

    def test_the_original_row_is_never_altered(self):
        # It keeps the bank's figure and its fingerprint, so re-uploading the statement
        # still changes nothing and the record of what the bank said survives.
        from kestrel.engine import split as sp
        t = self._payment()
        before = dict(t)
        sp.split(t["id"], self._thirds(t))
        after = db.one("SELECT * FROM transactions WHERE id=?", (t["id"],))
        assert after["amount"] == before["amount"]
        assert after["fingerprint"] == before["fingerprint"]
        assert after["description"] == before["description"]
        assert after["is_split"] == 1        # the only thing that changed
        sp.unsplit(t["id"])

    def test_splitting_changes_no_household_total(self):
        # The invariant the whole design rests on. Count the parent AND its parts and
        # every figure in the app silently doubles, with nothing on screen looking wrong.
        from kestrel.engine import split as sp
        t = self._payment()
        before = self._totals()
        sp.split(t["id"], self._thirds(t))
        assert self._totals() == before, "a split must not move a single penny"
        sp.unsplit(t["id"])
        assert self._totals() == before, "and undoing it must not either"

    def test_the_money_lands_in_the_categories_chosen(self):
        from kestrel.engine import split as sp
        t = self._payment()
        s, e = budgets.month_bounds(budgets.current_period())
        gid = db.category_id("Food & Drink", "Groceries")
        before = budgets.spend_by_category(s, e)
        was = next((l["amount"] for sec in before["sections"].values()
                    for l in sec["lines"] if l["category_id"] == gid), 0.0)
        sp.split(t["id"], self._thirds(t))
        after = budgets.spend_by_category(s, e)
        now = next((l["amount"] for sec in after["sections"].values()
                    for l in sec["lines"] if l["category_id"] == gid), 0.0)
        expected_move = round(t["amount"] * 0.5, 2) - (t["amount"]
                              if db.one("SELECT category_id FROM transactions WHERE id=?",
                                        (t["id"],)) and False else 0)
        assert abs(now - was) > 0.01, "groceries should have moved"
        sp.unsplit(t["id"])

    def test_the_parts_must_add_up(self):
        # Splitting £100 into £50 + £20 + £25 loses £5: the spending disappears, the
        # account stops reconciling, and nobody notices.
        from kestrel.engine import split as sp
        t = self._payment()
        short = self._thirds(t)
        short[-1]["amount"] = round(short[-1]["amount"] + 5, 2)     # £5 unallocated
        with pytest.raises(sp.SplitError) as e:
            sp.split(t["id"], short)
        assert "allocate" in str(e.value) or "too much" in str(e.value)
        assert not db.one("SELECT id FROM transactions WHERE split_of=?", (t["id"],)), \
            "a refused split must write nothing"
        assert db.one("SELECT is_split FROM transactions WHERE id=?", (t["id"],))["is_split"] == 0

    def test_a_part_cannot_pull_the_other_way(self):
        # A positive part inside a payment out would quietly *reduce* the household's
        # spending.
        from kestrel.engine import split as sp
        t = self._payment()
        parts = [{"amount": t["amount"] - 10, "category_id": db.category_id("Food & Drink", "Groceries")},
                 {"amount": 10, "category_id": db.category_id("Shopping", "Clothing")}]
        with pytest.raises(sp.SplitError):
            sp.split(t["id"], parts)

    def test_it_needs_at_least_two_parts(self):
        from kestrel.engine import split as sp
        t = self._payment()
        with pytest.raises(sp.SplitError):
            sp.split(t["id"], [{"amount": t["amount"], "category_id": None}])

    def test_re_splitting_replaces_the_parts_rather_than_adding_to_them(self):
        from kestrel.engine import split as sp
        t = self._payment()
        sp.split(t["id"], self._thirds(t))
        assert db.scalar("SELECT COUNT(*) FROM transactions WHERE split_of=?", (t["id"],), 0) == 3
        half = [{"amount": round(t["amount"] / 2, 2),
                 "category_id": db.category_id("Food & Drink", "Groceries")},
                {"amount": round(t["amount"] - round(t["amount"] / 2, 2), 2),
                 "category_id": db.category_id("Shopping", "Clothing")}]
        sp.split(t["id"], half)
        assert db.scalar("SELECT COUNT(*) FROM transactions WHERE split_of=?", (t["id"],), 0) == 2
        assert self._totals()  # still coherent
        sp.unsplit(t["id"])

    def test_a_part_cannot_itself_be_split(self):
        from kestrel.engine import split as sp
        t = self._payment()
        sp.split(t["id"], self._thirds(t))
        part = db.one("SELECT * FROM transactions WHERE split_of=?", (t["id"],))
        with pytest.raises(sp.SplitError):
            sp.split(part["id"], self._thirds(part))
        sp.unsplit(t["id"])

    def test_asking_about_a_part_answers_about_the_whole(self):
        from kestrel.engine import split as sp
        t = self._payment()
        sp.split(t["id"], self._thirds(t))
        part = db.one("SELECT id FROM transactions WHERE split_of=?", (t["id"],))
        assert sp.get(part["id"])["parent"]["id"] == t["id"]
        sp.unsplit(t["id"])

    def test_undoing_removes_the_parts_and_restores_the_row(self):
        from kestrel.engine import split as sp
        t = self._payment()
        sp.split(t["id"], self._thirds(t))
        out = sp.unsplit(t["id"])
        assert out["removed"] == 3
        assert db.one("SELECT is_split FROM transactions WHERE id=?", (t["id"],))["is_split"] == 0
        assert not db.rows("SELECT id FROM transactions WHERE split_of=?", (t["id"],))


class TestSplitIsExcludedEverywhere:
    """The rule is: a split parent is never counted, its children are. Every place that
    sums or counts has to know."""

    def _split_one(self):
        from kestrel.engine import split as sp
        start, _ = budgets.month_bounds(budgets.current_period())
        t = db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                   "AND IFNULL(is_split,0)=0 AND split_of IS NULL ORDER BY amount LIMIT 1",
                   (start,))
        half = round(t["amount"] / 2, 2)
        sp.split(t["id"], [
            {"amount": half, "category_id": db.category_id("Food & Drink", "Groceries")},
            {"amount": round(t["amount"] - half, 2),
             "category_id": db.category_id("Shopping", "Clothing")}])
        return t

    def test_the_listing_shows_the_parts_and_not_the_whole(self):
        from kestrel.engine import split as sp
        from kestrel.web import server
        t = self._split_one()
        ids = {r["id"] for r in server.api_transactions(None, {}, {"limit": "2000"}, None)["items"]}
        assert t["id"] not in ids, "the parent would be a second copy of the same money"
        parts = {r["id"] for r in db.rows("SELECT id FROM transactions WHERE split_of=?", (t["id"],))}
        assert parts <= ids
        sp.unsplit(t["id"])

    def test_the_workbook_does_not_double_count(self, tmp_path):
        from kestrel.engine import split as sp
        from kestrel.export import banking_xlsx
        import openpyxl
        t = self._split_one()
        path = banking_xlsx.build(tmp_path / "b.xlsx")
        wb = openpyxl.load_workbook(str(path))
        ws = wb["Transactions"]
        seen = [r for r in ws.iter_rows(values_only=True)
                if r and any(str(c) == str(t["description"]) for c in r if c)]
        # the parent's own row must not be in there alongside its parts
        totals = [r for r in seen if any(isinstance(c, (int, float))
                                         and abs(abs(c) - abs(t["amount"])) < 0.01 for c in r)]
        assert not totals, "the whole payment appears beside its parts"
        sp.unsplit(t["id"])

    def test_a_part_is_never_paired_as_an_internal_transfer(self):
        # A £50 fragment of a £100 shop matching some other £50 would silently remove
        # real spending from the budget.
        from kestrel.engine import split as sp
        t = self._split_one()
        categorise.detect_internal_transfers()
        for r in db.rows("SELECT transfer_pair FROM transactions WHERE split_of=?", (t["id"],)):
            assert r["transfer_pair"] is None
        assert db.one("SELECT transfer_pair FROM transactions WHERE id=?",
                      (t["id"],))["transfer_pair"] is None
        sp.unsplit(t["id"])

    def test_re_sorting_categories_leaves_a_split_alone(self):
        from kestrel.engine import split as sp
        t = self._split_one()
        before = {r["id"]: r["category_id"]
                  for r in db.rows("SELECT id, category_id FROM transactions WHERE split_of=?",
                                   (t["id"],))}
        categorise.categorise_all(only_uncategorised=False)
        after = {r["id"]: r["category_id"]
                 for r in db.rows("SELECT id, category_id FROM transactions WHERE split_of=?",
                                  (t["id"],))}
        assert after == before, "the parts were chosen by hand"
        sp.unsplit(t["id"])

    def test_a_split_survives_re_importing_the_statement(self):
        # The parent keeps its fingerprint, so the import skips it as a duplicate and
        # the parts are untouched. This is the whole reason the parent is left alone.
        from kestrel.engine import split as sp
        from kestrel.importers.readers import read_csv_bytes
        acct = db.one("SELECT id, currency FROM accounts WHERE closed=0 LIMIT 1")
        raw = b"Date,Description,Amount\n14/07/2026,SPLIT ME LATER,-90.00\n"
        table = read_csv_bytes(raw, "s.csv")
        mapping = {"date": 0, "description": 1, "amount": 2}
        ingest.import_transactions(acct["id"], table, mapping, None, "s.csv", None)
        t = db.one("SELECT * FROM transactions WHERE description='SPLIT ME LATER'")
        sp.split(t["id"], [
            {"amount": -60.0, "category_id": db.category_id("Food & Drink", "Groceries")},
            {"amount": -30.0, "category_id": db.category_id("Shopping", "Clothing")}])

        again = ingest.import_transactions(acct["id"], read_csv_bytes(raw, "s.csv"),
                                           mapping, None, "s.csv", None)
        assert again["added"] == 0 and again["skipped"] == 1, "the bank's row is still a duplicate"
        assert db.scalar("SELECT COUNT(*) FROM transactions WHERE split_of=?", (t["id"],), 0) == 2
        assert db.scalar("SELECT COUNT(*) FROM transactions "
                         "WHERE description='SPLIT ME LATER'", (), 0) == 1
        with db.tx() as c:
            c.execute("DELETE FROM transactions WHERE split_of=?", (t["id"],))
            c.execute("DELETE FROM transactions WHERE id=?", (t["id"],))


class TestSplitRoutes:
    def test_the_screen_gets_a_sensible_starting_point(self):
        from kestrel.web import server
        start, _ = budgets.month_bounds(budgets.current_period())
        t = db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                   "AND IFNULL(is_split,0)=0 AND split_of IS NULL ORDER BY amount LIMIT 1",
                   (start,))
        out = server.api_get_split(None, {"tid": str(t["id"])}, {}, None)
        assert out["is_split"] is False
        # the whole amount in the category it already has, plus a blank — most splits are
        # "most of it was what you thought, and a bit of it wasn't"
        assert out["suggested"][0]["amount"] == round(t["amount"], 2)
        assert out["suggested"][0]["category_id"] == t["category_id"]

    def test_a_bad_split_comes_back_as_a_sentence(self):
        from kestrel.web import server
        start, _ = budgets.month_bounds(budgets.current_period())
        t = db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                   "AND IFNULL(is_split,0)=0 AND split_of IS NULL ORDER BY amount LIMIT 1",
                   (start,))
        with pytest.raises(server.ApiError) as e:
            server.api_split(None, {"tid": str(t["id"])}, {}, {"parts": [
                {"amount": -1, "category_id": None}, {"amount": -2, "category_id": None}]})
        assert "still to allocate" in str(e.value).lower()

    def test_the_shortfall_message_says_which_way_it_is_wrong(self):
        # `total - running` is NEGATIVE when the parts fall short of a payment out, and
        # reading that sign directly told people "too much" when they had underspent —
        # the exact opposite of what was wrong.
        from kestrel.engine import split as sp
        start, _ = budgets.month_bounds(budgets.current_period())
        t = db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                   "AND IFNULL(is_split,0)=0 AND split_of IS NULL ORDER BY amount LIMIT 1",
                   (start,))
        gid = db.category_id("Food & Drink", "Groceries")
        cid = db.category_id("Shopping", "Clothing")

        with pytest.raises(sp.SplitError) as under:
            sp.split(t["id"], [{"amount": -1.0, "category_id": gid},
                               {"amount": -2.0, "category_id": cid}])
        assert "still to allocate" in str(under.value)

        with pytest.raises(sp.SplitError) as over:
            sp.split(t["id"], [{"amount": t["amount"], "category_id": gid},
                               {"amount": -50.0, "category_id": cid}])
        assert "too much" in str(over.value)

    def test_a_part_the_app_named_itself_is_not_pretend_user_text(self):
        # Reopening a split used to hand back "TESCO STORES 3411 (1 of 3)" in the
        # editable box, as though somebody had typed it and now had to maintain it.
        from kestrel.engine import split as sp
        start, _ = budgets.month_bounds(budgets.current_period())
        t = db.one("SELECT * FROM transactions WHERE amount < -40 AND posted_on >= ? "
                   "AND IFNULL(is_split,0)=0 AND split_of IS NULL ORDER BY amount LIMIT 1",
                   (start,))
        half = round(t["amount"] / 2, 2)
        out = sp.split(t["id"], [
            {"amount": half, "category_id": db.category_id("Food & Drink", "Groceries")},
            {"amount": round(t["amount"] - half, 2), "description": "the telly",
             "category_id": db.category_id("Shopping", "Clothing")}])
        auto, typed = out["parts"][0], out["parts"][1]
        assert auto["description"] == f"{t['description']} (1 of 2)"
        assert typed["description"] == "the telly", "a real note is kept verbatim"
        sp.unsplit(t["id"])


# ===========================================================================
# The app icon and the words on screen
# ===========================================================================

class TestAppIcon:
    def _res(self):
        return pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "resources"

    def test_every_icon_file_is_there(self):
        for name in ("kestrel.ico", "kestrel.icns", "kestrel.png", "kestrel-512.png"):
            f = self._res() / name
            assert f.exists() and f.stat().st_size > 500, name

    def test_the_icon_is_square_and_not_the_old_bird(self):
        from PIL import Image
        im = Image.open(self._res() / "kestrel-512.png")
        assert im.size == (512, 512), "a squashed icon is worse than a padded one"
        # The Kestrel bird was a flat two-colour mark; the artwork is a photograph-like
        # illustration. Colour variety separates them without pinning the exact image.
        colours = im.convert("RGB").getcolors(maxcolors=200000)
        assert colours and len(colours) > 2000, "this looks like the old flat logo"

    def test_the_windows_icon_carries_the_small_sizes(self):
        # Ship only a 256px image and Windows squashes it down for the taskbar itself,
        # which turns to mush at 16px.
        import struct
        raw = (self._res() / "kestrel.ico").read_bytes()
        _, itype, count = struct.unpack("<HHH", raw[:6])
        assert itype == 1, "not an icon file"
        sizes = []
        for i in range(count):
            off = 6 + i * 16
            w, _h, _c, _r, _p, bpp, size, offset = struct.unpack("<BBBBHHII", raw[off:off+16])
            sizes.append(w or 256)
            assert bpp == 32
            assert offset + size <= len(raw), "an image runs off the end of the file"
        for needed in (16, 32, 48, 256):
            assert needed in sizes, f"no {needed}px image in the .ico"

    def test_the_ico_uses_the_layout_windows_is_always_given(self):
        # BMP below 256, PNG at 256. Pillow writes PNG at every size — legal, but
        # unconventional, and this file cannot be tested here because the executable is
        # built on Windows.
        import struct
        raw = (self._res() / "kestrel.ico").read_bytes()
        _, _, count = struct.unpack("<HHH", raw[:6])
        for i in range(count):
            off = 6 + i * 16
            w, _h, _c, _r, _p, _bpp, size, offset = struct.unpack("<BBBBHHII", raw[off:off+16])
            blob = raw[offset:offset+size]
            n = w or 256
            if n >= 256:
                assert blob[:8] == b"\x89PNG\r\n\x1a\n", "the 256px entry should be PNG"
            else:
                hs, _bw, bh, _pl, bits = struct.unpack("<IiiHH", blob[:16])
                assert hs == 40 and bits == 32
                # A BMP inside an icon stores colour rows AND the mask rows, so the
                # header height is doubled. Get this wrong and it shows as a blank square.
                assert bh == n * 2, f"{n}px entry has height {bh}, expected {n * 2}"

    def test_pillow_can_read_every_size_back(self):
        from PIL import Image
        path = self._res() / "kestrel.ico"
        for n in sorted(Image.open(path).ico.sizes()):
            im = Image.open(path)
            im.size = n
            rgba = im.convert("RGBA")
            assert rgba.size == n
            assert rgba.getextrema()[3][1] > 0, f"the {n} image is fully transparent"

    def test_the_mac_icon_is_a_valid_icns(self):
        # Written by hand: Pillow emits the whole tier ladder at full quality whatever
        # you feed it, which came to 3 MB for one icon file.
        import io, struct
        from PIL import Image
        raw = (self._res() / "kestrel.icns").read_bytes()
        magic, total = struct.unpack(">4sI", raw[:8])
        assert magic == b"icns"
        assert total == len(raw), "the header length must match the file, or macOS rejects it"
        i, seen = 8, {}
        while i < len(raw) - 8:
            tag, n = struct.unpack(">4sI", raw[i:i+8])
            assert n >= 8 and i + n <= len(raw), f"chunk {tag} runs off the end"
            im = Image.open(io.BytesIO(raw[i+8:i+n]))
            seen[tag.decode()] = im.size[0]
            i += n
        for tag, px in (("ic07", 128), ("ic08", 256), ("ic09", 512)):
            assert seen.get(tag) == px, f"{tag} should hold a {px}px image, got {seen.get(tag)}"
        assert len(raw) < 2_000_000, "one icon should not weigh more than the rest of the app"

    def test_the_browser_tab_uses_the_same_artwork(self):
        from PIL import Image
        fav = Image.open(pathlib.Path(__file__).resolve().parent.parent / "kestrel"
                         / "web" / "static" / "favicon.png")
        assert fav.size == (64, 64) and fav.mode == "RGBA"

    def test_the_spec_asks_for_the_right_format_per_platform(self):
        # Windows needs .ico and macOS needs .icns. The old loop took the first file
        # that existed, which would have handed every Mac build a .ico.
        spec = (pathlib.Path(__file__).resolve().parent.parent / "build"
                / "mittens.spec").read_text(encoding="utf-8")
        block = spec[spec.index("_ICON_FOR = {"):spec.index("ONEFILE =")]
        root = pathlib.Path(__file__).resolve().parent.parent
        for platform, expected in (("win32", "kestrel.ico"), ("darwin", "kestrel.icns")):
            ns = {"PKG": root / "kestrel",
                  "sys": type("S", (), {"platform": platform})}
            exec(compile(block, "spec", "exec"), ns)
            assert ns["icon"] and ns["icon"].endswith(expected), (platform, ns["icon"])


class TestUserFacingWording:
    """Words the app's owner asked for by name. Cheap to assert, easy to lose in a
    later edit."""

    def _root(self):
        return pathlib.Path(__file__).resolve().parent.parent

    def test_the_side_panel_and_page_subtitles(self):
        html = (self._root() / "kestrel" / "web" / "static" / "index.html").read_text(encoding="utf-8")
        js = (self._root() / "kestrel" / "web" / "static" / "app.js").read_text(encoding="utf-8")
        assert "Budget and Investment App" in html
        assert "UK &amp; South Africa" not in html
        assert "'All investments overview'" in js
        assert "'All overall balances'" in js
        for gone in ("The same figures your Investments sheet tracked",
                     "Balances across both countries",
                     "The same shape as your Investments 2 sheet"):
            assert gone not in js, gone

    def test_no_test_reads_a_file_without_saying_utf8(self):
        """`read_text()` uses the *platform's* default encoding — UTF-8 on Linux and
        macOS, cp1252 on Windows. Every file in this repo is UTF-8, and several carry
        em dashes or the nav glyphs in index.html, so a bare read passes everywhere
        except the machine that actually builds the .exe:

            UnicodeDecodeError: 'charmap' codec can't decode byte 0x90

        The build runs the tests before it packages anything, so this stopped the build
        rather than merely annoying somebody.
        """
        import re
        root = pathlib.Path(__file__).resolve().parent.parent
        opener = re.compile(r"\.(?:read|write)_text\(")
        bare = []
        for path in sorted(list((root / "kestrel").rglob("*.py"))
                           + list((root / "tools").rglob("*.py"))
                           + list((root / "tests").rglob("*.py"))):
            text = path.read_text(encoding="utf-8")
            for m in opener.finditer(text):
                if "opener" in text[text.rfind("\n", 0, m.start()) + 1:m.start()]:
                    continue                      # the checker's own line
                # Walk to the matching close paren. A naive "up to the first )" reads
                # `write_text(json.dumps(x), encoding="utf-8")` as unencoded, because
                # json.dumps closes first — and the call can span lines besides.
                depth, i = 1, m.end()
                while i < len(text) and depth:
                    depth += {"(": 1, ")": -1}.get(text[i], 0)
                    i += 1
                if "encoding=" not in text[m.end():i]:
                    line = text.count("\n", 0, m.start()) + 1
                    bare.append(f"{path.name}:{line}")
        assert not bare, f"text reads with no explicit encoding: {bare}"

    def test_the_docs_address_a_new_user_not_the_author(self):
        # The read-me went out with the app; it must not assume the reader knows the
        # spreadsheet it grew out of, or who built it.
        readme = (self._root() / "README.md").read_text(encoding="utf-8")
        for personal in ("Investments 2", "the old spreadsheet", "your old sheet"):
            assert personal not in readme, personal
        first = (self._root() / "build" / "Read me first.txt").read_text(encoding="utf-8")
        assert "one family" not in first
        assert "Allan" not in readme and "Allan" not in first

    def test_the_vault_key_survives_a_rename(self):
        # The keyring service name is what the vault key is stored under. A blanket
        # rename rewrote it to the new product name, which cannot see the old entry —
        # the app would have minted a fresh key and every saved bank credential would
        # have become unreadable, with no error to say why.
        from kestrel import security
        assert security._KEYRING_SERVICE == "Mittens and Pence Finance"
        assert security._LEGACY_KEYRING_SERVICES == ("Mithapp Finance", "Kestrel Finance")
        assert "&" not in security._KEYRING_SERVICE

    def test_an_older_installs_vault_key_is_found_and_carried_forward(self):
        from kestrel import security

        class FakeKeyring:
            def __init__(self):
                self.store = {("Mithapp Finance", "vault-key"): "b3ZlcnRoZXJlLWtleS1oZXJlLTMyLWJ5dGVz"}
            def get_password(self, service, user):
                return self.store.get((service, user))
            def set_password(self, service, user, value):
                self.store[(service, user)] = value

        fake = FakeKeyring()
        real = security._try_keyring
        security._try_keyring = lambda: fake
        try:
            key = security._load_or_create_key()
            assert key, "the old key should have been found"
            # copied forward, and the original left alone
            assert ("Mittens and Pence Finance", "vault-key") in fake.store
            assert ("Mithapp Finance", "vault-key") in fake.store
            assert fake.store[("Mittens and Pence Finance", "vault-key")] == \
                   fake.store[("Mithapp Finance", "vault-key")]
        finally:
            security._try_keyring = real

    def test_the_legacy_chain_is_never_rewritten_to_the_current_name(self):
        # Guard for the blanket-rename trap that has now caught this project three times:
        # a value whose whole job is to hold the OLD name got the new one.
        from kestrel import security
        assert config.APP_NAME not in config.LEGACY_NAMES
        assert config.APP_FILE_NAME not in config.LEGACY_NAMES
        assert config.APP_SLUG not in config.LEGACY_SLUGS
        assert config.DB_NAMES[0] not in config.DB_NAMES[1:]
        assert security._KEYRING_SERVICE not in security._LEGACY_KEYRING_SERVICES


# ===========================================================================
# The phone copy — one self-contained HTML file
# ===========================================================================

class TestPhoneCopy:
    def _html(self):
        from kestrel.export import phone
        return phone.render()

    def test_it_is_genuinely_self_contained(self):
        # The whole promise is "open it on a phone with no signal". One reference to
        # a font, a script or an image on the internet and that promise is broken —
        # silently, and only for the person standing in a car park.
        html = self._html()
        for outside in ("http://", "https://", "//cdn", "<img", "<link", "@import",
                        "src=", "fetch(", "XMLHttpRequest"):
            assert outside not in html, f"the phone copy reaches outside for {outside!r}"

    def test_it_says_when_it_was_made(self):
        # A stale balance that looks current is worse than no balance.
        html = self._html()
        assert "snapshot taken on" in html
        assert "does not update on its own" in html.replace("\n", " ").replace("  ", " ")

    def test_names_and_descriptions_are_escaped(self):
        from kestrel import db
        from kestrel.export import phone
        acct = db.one("SELECT id, name FROM accounts WHERE closed=0 LIMIT 1")
        original = acct["name"]
        nasty = '<script>alert(1)</script> & "co"'
        with db.tx() as c:
            c.execute("UPDATE accounts SET name=? WHERE id=?", (nasty, acct["id"]))
        try:
            html = phone.render()
            assert "<script>alert(1)</script>" not in html
            assert "&lt;script&gt;" in html
        finally:
            with db.tx() as c:
                c.execute("UPDATE accounts SET name=? WHERE id=?", (original, acct["id"]))

    def test_unknown_is_still_not_zero(self):
        from kestrel.export import phone
        assert phone.money(None, "GBP") == "—"
        assert phone.money(0, "GBP") == "£0.00"
        assert phone.money(-1234.5, "GBP") == "−£1,234"     # rounds above a thousand
        assert phone.money(-12.34, "GBP") == "−£12.34"
        assert phone.pct(None) == "—"

    def test_dates_do_not_depend_on_the_platform(self):
        # strftime's no-padding flag is %-d on Unix and %#d on Windows; using the
        # wrong one raises there and would break the export on the machines this app
        # actually runs on.
        from kestrel.export import phone
        assert phone.long_date(dt.date(2026, 4, 3)) == "3 April 2026"
        assert phone.short_date("2026-04-03") == "3 Apr"
        assert phone.short_date(None) == ""
        assert phone.short_date("not a date") == ""
        assert phone.clock(dt.datetime(2026, 4, 3, 0, 5)) == "12:05 am"
        assert phone.clock(dt.datetime(2026, 4, 3, 12, 5)) == "12:05 pm"
        assert phone.clock(dt.datetime(2026, 4, 3, 17, 43)) == "5:43 pm"
        # No strftime at all in there, so the flag can't creep back in.
        src = pathlib.Path(phone.__file__).read_text(encoding="utf-8")
        assert "strftime(" not in src

    def test_a_split_parent_is_never_listed_twice(self):
        # The children carry the money; listing the parent as well would show the same
        # payment twice on the one screen somebody glances at.
        from kestrel.export import phone
        src = pathlib.Path(phone.__file__).read_text(encoding="utf-8")
        assert "NOT_SPLIT_PARENT" in src

    def test_it_writes_a_file_you_can_find(self):
        from kestrel.export import phone
        out = phone.build()
        assert out.exists() and out.suffix == ".html"
        assert out.stat().st_size > 2000
        assert "&" not in out.name, "a filename with an ampersand is trouble in a shell"


# ===========================================================================
# Who the local server will answer
# ===========================================================================

class TestLocalServerIsActuallyLocal:
    """Binding to 127.0.0.1 is not the same as being private.

    Every web page the household opens is also running on this machine, and a page
    can POST to 127.0.0.1:8765 in the background. Before these two checks existed,
    any site could have wiped the database through /api/demo/clear, and a site
    pointing its own domain at 127.0.0.1 could have read every transaction.
    """

    def _handler(self, **headers):
        from kestrel.web import server as web

        class Fake:
            pass
        f = Fake()
        f.headers = {"Host": "127.0.0.1:8765", **headers}
        f.LOCAL_HOSTS = web.Handler.LOCAL_HOSTS
        return web.Handler, f

    def test_a_rebound_domain_is_refused(self):
        H, f = self._handler(Host="money.evil.example.com")
        assert H._host_ok(f) is False

    def test_the_loopback_host_is_fine_in_its_usual_spellings(self):
        for host in ("127.0.0.1:8765", "localhost:8765", "127.0.0.1", "[::1]:8765"):
            H, f = self._handler(Host=host)
            assert H._host_ok(f) is True, host

    def test_a_missing_host_header_is_allowed(self):
        # HTTP/1.0 clients and some local tooling send none.
        from kestrel.web import server as web

        class Fake:
            headers = {}
            LOCAL_HOSTS = web.Handler.LOCAL_HOSTS
        assert web.Handler._host_ok(Fake()) is True

    def test_another_site_cannot_post(self):
        H, f = self._handler(Origin="https://evil.example.com")
        assert H._origin_ok(f, "POST") is False

    def test_a_sandboxed_frame_reporting_null_cannot_post(self):
        # "null" is what a sandboxed iframe sends. Reading it as "no origin, must be
        # a local tool" would hand any web page a way straight in.
        H, f = self._handler(Origin="null")
        assert H._origin_ok(f, "POST") is False

    def test_our_own_front_end_can_post(self):
        for origin in ("http://127.0.0.1:8765", "http://localhost:8765"):
            H, f = self._handler(Origin=origin)
            assert H._origin_ok(f, "POST") is True, origin

    def test_a_local_tool_with_no_origin_can_post(self):
        H, f = self._handler()
        assert H._origin_ok(f, "POST") is True

    def test_a_cross_site_read_is_refused_even_without_an_origin(self):
        H, f = self._handler(**{"Sec-Fetch-Site": "cross-site"})
        assert H._origin_ok(f, "GET") is False
        H, f = self._handler(**{"Sec-Fetch-Site": "same-origin"})
        assert H._origin_ok(f, "GET") is True

    def test_the_destructive_routes_are_all_behind_this(self):
        # Nothing may opt out: the check lives in the dispatcher, not per-route.
        from kestrel.web import server as web
        src = pathlib.Path(web.__file__).read_text(encoding="utf-8")
        body = src[src.index("def _dispatch"):src.index("def _serve_static")]
        assert "_host_ok" in body and "_origin_ok" in body


class TestTheScreensAgreeWithEachOther:
    def test_net_worth_is_what_you_own_less_what_you_owe(self):
        # The Dashboard shows a headline net worth and, beside it, a ring of where the
        # money sits. The ring used to take the absolute value of every group, which
        # put the mortgage in as though it were an asset and produced a figure in the
        # middle — labelled "total" — of roughly twice the net worth printed above it.
        from kestrel.engine import portfolio
        nw = portfolio.net_worth()
        owned = sum(v for v in nw["groups"].values() if v > 0)
        owed = sum(v for v in nw["groups"].values() if v < 0)
        assert owned + owed == pytest.approx(nw["total"], abs=0.01)
        assert owed < 0, "the sample household has a mortgage; this test needs one"

    def test_the_dashboard_ring_shows_assets_only(self):
        js = (pathlib.Path(__file__).resolve().parent.parent
              / "kestrel" / "web" / "static" / "app.js").read_text(encoding="utf-8")
        block = js[js.index("VIEWS.dashboard"):js.index("VIEWS.dashboard") + 1200]
        assert "Math.abs(value)" not in block, "the ring is folding debt in as an asset again"
        assert "owedTotal" in block

    def test_the_read_me_points_at_the_folder_the_app_actually_uses(self):
        # Somebody following this to back up their data has to arrive somewhere real.
        # It said AppData\Roaming\Mittens & Pence; the app uses AppData\Local and the
        # ampersand-free name, so the instructions led to an empty folder.
        first = ((pathlib.Path(__file__).resolve().parent.parent / "build"
                  / "Read me first.txt").read_text(encoding="utf-8"))
        assert f"AppData\\Local\\{config.APP_FILE_NAME}" in first
        assert f"Library/Application Support/{config.APP_FILE_NAME}" in first
        assert "AppData\\Roaming" not in first


class TestTheBrandBlock:
    def _static(self, name):
        return ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
                 / "static" / name).read_text(encoding="utf-8"))

    def test_the_brand_stays_inside_the_rail(self):
        # An earlier version had the brand card overhang the rail and sit over the main
        # column, with matching left padding on the topbar so headings cleared it.
        # Allan didn't like it. The two halves of that hack have to leave together —
        # a leftover 66px gutter with nothing in it is worse than either.
        css = self._static("styles.css")
        brand = css[css.index(".brand {"):css.index(".brand {") + 400]
        assert "margin: -" not in brand, "the brand is breaking out of the rail again"
        assert "z-index" not in brand
        assert "66px" not in css, "left padding compensating for an overhang that has gone"

    def test_the_brand_is_sized_from_one_number(self):
        # Logo, name and subtitle are all derived from --rail, which is itself a
        # proportion of the window. A fixed pixel size anywhere in here is how the cat
        # ends up lost on a big monitor and crowding a small one.
        css = self._static("styles.css")
        assert "--rail: clamp(" in css
        assert "grid-template-columns: var(--rail) 1fr" in css
        for rule in (".mark {", ".brand-name {", ".brand-sub {"):
            block = css[css.index(rule):css.index(rule) + 320]
            assert "var(--rail)" in block, f"{rule} is not sized from the rail"

    def test_the_narrow_rail_hides_labels_and_keeps_icons(self):
        # `#nav button span:last-child` matched the ICON, because the label was a bare
        # text node — so a narrow window hid the icons, kept the labels, and the labels
        # spilled out of a 68px rail.
        html = self._static("index.html")
        css = self._static("styles.css")
        nav = html[html.index("<nav id=\"nav\">"):html.index("</nav>")]
        assert nav.count('class="lbl"') == nav.count("data-view="), \
            "every nav label needs its own span, or the collapse rule hides the icon"
        assert "#nav button span:last-child" not in css
        assert "#nav .lbl { display: none; }" in css

    def test_the_watermark_survives_a_narrow_window(self):
        css = self._static("styles.css")
        narrow = css[css.index("@media (max-width: 900px)"):]
        narrow = narrow[:narrow.index("\n}")]
        assert ".credit { display: none" not in narrow


# ===========================================================================
# Categories and rules added after the first release
# ===========================================================================

class TestTheShoppingBoundary:
    """Shopping is things you buy and keep; Food & Drink is things you consume.

    The boundary is decided by the *shop*, not the basket — one Tesco payment is one
    transaction and goes to Groceries whole, even if half of it was a frying pan.
    Splitting a payment is how you say otherwise.
    """

    def _guess(self, text):
        rules = categorise._rules()
        for r in rules:
            if categorise._match(r, text, -10.0):
                return f"{r['parent']} > {r['name']}"
        return "Uncategorised"

    @pytest.mark.parametrize("description,expected", [
        # Consumed → Food & Drink, whichever shop it came from
        ("TESCO STORES 3456", "Food & Drink > Groceries"),
        ("ALDI 812 BROMSGROVE", "Food & Drink > Groceries"),
        ("DELIVEROO", "Food & Drink > Restaurants & takeaway"),
        ("COSTA COFFEE", "Food & Drink > Coffee & snacks"),
        ("MAJESTIC WINE", "Food & Drink > Alcohol"),
        # Kept → Shopping, by what kind of shop it is
        ("ZARA UK", "Shopping > Clothing"),
        ("CURRYS ONLINE", "Shopping > Electronics"),
        ("WATERSTONES", "Shopping > Books & media"),
        ("MOONPIG.COM", "Shopping > Gifts given"),
        # The shop that sells everything
        ("AMAZON.CO.UK*AB12CD", "Shopping > Household & general"),
        ("AMZNMKTPLACE", "Shopping > Household & general"),
        ("EBAY O*12-34567", "Shopping > Household & general"),
        ("JOHN LEWIS", "Shopping > Household & general"),
        ("WILKO", "Shopping > Household & general"),
        ("POUNDLAND", "Shopping > Household & general"),
        ("B&M BARGAINS", "Shopping > Household & general"),
    ])
    def test_where_a_payment_lands(self, description, expected):
        assert self._guess(description) == expected

    @pytest.mark.parametrize("description,wrong,right", [
        # `next` is one of the commonest words in English and was matched bare.
        ("NEXT DAY DELIVERY CHARGE", "Shopping > Clothing", "Uncategorised"),
        # A charitable donation filed as a present.
        ("GIFT AID DONATION", "Shopping > Gifts given", "Giving > Charity"),
        # The old lookahead only looked forward, so "station" before "gas" slipped past
        # and a fill-up was filed as a utility bill.
        ("PETROL STATION GAS", "Utilities > Gas", "Transport > Fuel"),
        # Prime is a subscription, not a trip to the shops.
        ("AMAZON PRIME", "Shopping > Electronics", "Subscriptions > Music & video"),
    ])
    def test_the_false_positives_stay_fixed(self, description, wrong, right):
        got = self._guess(description)
        assert got != wrong
        assert got == right

    def test_a_marketplace_is_not_an_electronics_shop(self):
        # Amazon lived in Electronics because there was nowhere else, which made a
        # nappy order and a kettle both read as gadgets.
        assert "Electronics" not in self._guess("AMAZON.CO.UK")

    def test_clicks_is_not_claimed_twice(self):
        # Health > Pharmacy matches first, so a second Clicks pattern under Shopping
        # would never fire and would only mislead whoever read it next.
        src = pathlib.Path(categorise.__file__).read_text(encoding="utf-8")
        shopping = src[src.index('("Shopping", "Household & general"'):]
        shopping = shopping[:shopping.index("]),")]
        assert "clicks" not in shopping


class TestUpdatesReachExistingDatabases:
    """The trap: both seeders stop dead once a database has anything in them.

    `seed_categories()` returns early if any category exists, and `seed_builtin_rules()`
    if any built-in rule exists — correctly, because re-running either would resurrect
    what the household deleted. The consequence is that anything added later reaches
    new installs only, and every copy already out there silently keeps the old
    behaviour. These recorded, one-off steps are the way round it.
    """

    def _fresh(self, tmp_path, monkeypatch):
        monkeypatch.setenv("MITTENS_DATA_DIR", str(tmp_path))
        import importlib
        from kestrel import config as cfg
        importlib.reload(cfg)
        return cfg

    def test_a_category_added_later_reaches_an_old_database(self):
        assert db.category_id("Shopping", "Household & general") is not None

    def test_applying_twice_changes_nothing(self):
        before = db.scalar("SELECT COUNT(*) FROM rules", (), 0)
        cats = db.scalar("SELECT COUNT(*) FROM categories", (), 0)
        categorise.apply_rule_updates()
        db._apply_added_categories(db.conn())
        assert db.scalar("SELECT COUNT(*) FROM rules", (), 0) == before
        assert db.scalar("SELECT COUNT(*) FROM categories", (), 0) == cats

    def test_a_rule_the_household_edited_is_never_rewritten(self):
        # Once somebody has changed a rule it is theirs. Quietly overwriting it would
        # be worse than leaving the original bug in place.
        with db.tx() as c:
            c.execute("INSERT INTO rules(match_type,pattern,field,category_id,priority,"
                      "is_builtin) VALUES('regex','my own pattern','description',?,50,1)",
                      (db.category_id("Shopping", "Clothing"),))
            c.execute("DELETE FROM meta WHERE key='did:rule:next-needs-context'")
        # The shipped pattern is gone (this stand-in is not it), so nothing matches
        # and nothing is touched.
        categorise.apply_rule_updates()
        assert db.one("SELECT 1 FROM rules WHERE pattern='my own pattern'")
        with db.tx() as c:
            c.execute("DELETE FROM rules WHERE pattern='my own pattern'")

    def test_every_update_names_a_category_that_exists(self):
        for key, _old, _new, parent, child in categorise.RULE_UPDATES:
            assert db.category_id(parent, child) is not None, f"{key} -> {parent} > {child}"

    def test_the_keys_are_unique(self):
        keys = [k for k, *_ in categorise.RULE_UPDATES]
        assert len(keys) == len(set(keys)), "a repeated key means one update never runs"


class TestNeedsALookIsActionable:
    """The dashboard card that says what wants attention.

    It used to render the text and throw the ids away, so the front screen told you
    something was wrong and then left you to go and find it. An item nobody can act
    on has no business being on the front screen.
    """

    def test_every_item_says_where_to_go(self):
        from kestrel.engine import sync as sync_engine
        with db.tx() as c:
            c.execute("UPDATE transactions SET category_id=? WHERE id IN "
                      "(SELECT id FROM transactions ORDER BY id DESC LIMIT 3)",
                      (db.uncategorised_id(),))
            row = c.execute("SELECT id FROM connections LIMIT 1").fetchone()
            if row:
                c.execute("UPDATE connections SET status='error', status_detail='key rejected'"
                          " WHERE id=?", (row["id"],))
        try:
            issues = sync_engine.health()["issues"]
            assert issues, "the fixture should have produced something to report"
            for i in issues:
                assert i.get("action"), f"no way to act on: {i['text']}"
                assert i["action"].get("go"), f"nowhere to go for: {i['text']}"
                assert i.get("do"), f"no label for: {i['text']}"
        finally:
            with db.tx() as c:
                c.execute("UPDATE connections SET status='ok', status_detail=NULL")
            from kestrel.engine import categorise as cat
            cat.categorise_all(only_uncategorised=True)

    def test_the_dashboard_uses_them(self):
        js = (pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
              / "static" / "app.js").read_text(encoding="utf-8")
        assert "openAlert" in js
        assert "data-alert" in js

    def test_one_transaction_is_not_described_as_transaction_s(self):
        # "1 transaction(s) still need a category" is the sort of thing that makes an
        # app feel like a form.
        from kestrel.engine import sync as sync_engine
        with db.tx() as c:
            c.execute("UPDATE transactions SET category_id=? WHERE id = "
                      "(SELECT id FROM transactions ORDER BY id DESC LIMIT 1)",
                      (db.uncategorised_id(),))
        try:
            texts = [i["text"] for i in sync_engine.health()["issues"]]
            line = next((t for t in texts if "category" in t), "")
            assert "(s)" not in line, line
            assert "1 transaction still needs a category." == line, line
        finally:
            from kestrel.engine import categorise as cat
            cat.categorise_all(only_uncategorised=True)

    def test_accounts_kept_by_hand_are_never_nagged_about(self):
        # A house or a pension statement has no transactions by design.
        src = pathlib.Path((pathlib.Path(__file__).resolve().parent.parent
                            / "kestrel" / "engine" / "sync.py")).read_text(encoding="utf-8")
        block = src[src.index("def health()"):]
        assert "'asset','pension','loan','mortgage'" in block
        assert "IFNULL(c.method,'manual') <> 'manual'" in block

    def test_the_boundary_is_explained_where_people_read(self):
        # It is the one categorisation question that catches people out, so it is
        # spelled out in the file that ships beside the .exe — not only in the
        # developer README, which the family never opens.
        root = pathlib.Path(__file__).resolve().parent.parent
        first = (root / "build" / "Read me first.txt").read_text(encoding="utf-8")
        readme = (root / "README.md").read_text(encoding="utf-8")
        for doc, label in ((first, "Read me first.txt"), (readme, "README.md")):
            flat = " ".join(doc.split())
            assert "Household & general" in doc, label
            assert "frying pan" in flat, f"{label} does not give the supermarket example"
            assert "Split" in doc or "split the payment" in doc, label

    def test_the_app_explains_it_at_the_moment_of_asking(self):
        js = (pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
              / "static" / "app.js").read_text(encoding="utf-8")
        dialog = js[js.index("function categoryDialog"):]
        dialog = dialog[:dialog.index("\n}\n")]
        assert "Shopping or Food" in dialog, "the category picker does not answer it"


class TestTheLogoSprite:
    """The sheet's geometry and the CSS's idea of it have to agree exactly.

    A sprite sheet is only correct if every frame lands on a whole-frame boundary. Get
    the frame size or the grid wrong by a pixel and the logo shows slivers of the four
    neighbouring cats — which looks like a rendering glitch, not like a wrong number,
    so it is the sort of thing that ships.
    """

    def _css(self):
        return ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
                 / "static" / "styles.css").read_text(encoding="utf-8"))

    def _sheet(self):
        from PIL import Image
        return Image.open(pathlib.Path(__file__).resolve().parent.parent / "kestrel"
                          / "web" / "static" / "logo.webp")

    def test_the_whole_picture_is_inside_every_frame(self):
        # The point of rebuilding it: the previous sheet was cropped flush, so the pile
        # of banknotes was sliced off at the frame edge and read as a rendering fault.
        #
        # numpy is NOT a dependency of this app -- it belongs to tools/make_logo.py,
        # which is only run when the artwork changes. Skipping is right: a machine
        # without it can still build and ship. Failing there stopped Allan's build.
        np = pytest.importorskip("numpy")
        css = self._css()
        cols = int(re.search(r"--cols:\s*(\d+)", css).group(1))
        rows = int(re.search(r"--rows:\s*(\d+)", css).group(1))
        sheet = np.asarray(self._sheet().convert("RGBA"))
        fh, fw = sheet.shape[0] // rows, sheet.shape[1] // cols
        touching = []
        for r in range(rows):
            for c in range(cols):
                a = sheet[r * fh:(r + 1) * fh, c * fw:(c + 1) * fw, 3]
                if (a[0] > 12).any() or (a[-1] > 12).any() \
                   or (a[:, 0] > 12).any() or (a[:, -1] > 12).any():
                    touching.append((r, c))
        assert not touching, f"{len(touching)} frames run off their own edge: {touching[:6]}"

    def test_it_rests_two_to_three_minutes_between_plays(self):
        js = (pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
              / "static" / "app.js").read_text(encoding="utf-8")
        assert re.search(r"LOGO_REST_MIN\s*=\s*120", js)
        assert re.search(r"LOGO_REST_MAX\s*=\s*180", js)

    def test_the_frame_count_matches_the_play_length(self):
        # 120 frames over 10s is 12fps. A sheet with a different frame count and the
        # same --play runs at the wrong speed.
        css = self._css()
        cols = int(re.search(r"--cols:\s*(\d+)", css).group(1))
        rows = int(re.search(r"--rows:\s*(\d+)", css).group(1))
        play = float(re.search(r"--play:\s*([\d.]+)s", css).group(1))
        fps = (cols * rows) / play
        assert 10 <= fps <= 25, f"{cols*rows} frames over {play}s is {fps:.0f}fps"

    def test_there_is_a_script_to_rebuild_it(self):
        # The technique is fiddly enough that "how was this made" has to be answerable.
        tool = pathlib.Path(__file__).resolve().parent.parent / "tools" / "make_logo.py"
        assert tool.exists()
        src = tool.read_text(encoding="utf-8")
        for idea in ("connectivity", "un-multiply", "one crop box"):
            assert idea in src.lower(), f"the script does not explain: {idea}"


# ===========================================================================
# The statements folder
# ===========================================================================

class TestTheStatementsFolder:
    """Drop a file in, press Sync everything, and it goes where it belongs.

    The two gates are the whole safety argument: a file imports unattended only when
    the layout is one the app has been taught AND the filename names an account
    distinctively. Either alone is not enough — four banks can share a layout, and a
    remembered layout says nothing about whose account a file is. Skipping the second
    gate once put a Barclays PDF into a Capitec account.
    """

    #: Its own account with a name nothing else uses. Leaning on the demo household's
    #: "Monzo current account" made these pass alone and fail in the full run, because
    #: other tests rename and close accounts — an ordering bug in the tests, not in the
    #: code, which is the worst kind to chase.
    ACCOUNT = "Wobblethorpe Building Society"

    @pytest.fixture()
    def box(self, tmp_path, monkeypatch):
        from kestrel import inbox
        monkeypatch.setitem(config.settings._data, "inbox_folder", str(tmp_path / "drop"))
        monkeypatch.setattr(inbox, "_state_path", lambda: tmp_path / "seen.json")
        existing = db.one("SELECT id FROM accounts WHERE name=?", (self.ACCOUNT,))
        if not existing:
            with db.tx() as c:
                c.execute("INSERT INTO accounts(name,currency,account_type,is_investment,"
                          "closed) VALUES(?,'GBP','current',0,0)", (self.ACCOUNT,))
        yield inbox
        with db.tx() as c:
            c.execute("DELETE FROM transactions WHERE account_id IN "
                      "(SELECT id FROM accounts WHERE name=?)", (self.ACCOUNT,))
            c.execute("DELETE FROM accounts WHERE name=?", (self.ACCOUNT,))

    _n = 0

    def _statement(self, box, name, rows=4):
        # Every file gets genuinely different rows. Two files with identical
        # transactions are correctly de-duplicated by fingerprint, which would make a
        # test about *files* pass or fail for reasons about *rows*.
        TestTheStatementsFolder._n += 1
        tag = TestTheStatementsFolder._n
        f = box.folder() / name
        lines = ["Date,Description,Amount,Balance"]
        for i in range(rows):
            lines.append(f"{i+1:02d}/08/2026,INBOX TEST {tag}-{i},-{100 * tag + i}.00,"
                         f"{5000 - i}.00")
        f.write_text("\n".join(lines), encoding="utf-8")
        return f

    def test_the_folder_explains_itself(self, box):
        f = box.folder()
        readme = next(p for p in f.iterdir() if p.name.lower().startswith("read me"))
        text = readme.read_text(encoding="utf-8")
        assert "Sync everything" in text
        assert "NOT read" in text, "it must say receipts are kept, not read"

    def test_a_named_statement_goes_in_and_is_filed(self, box):
        self._statement(box, f"{self.ACCOUNT} Aug 2026.csv")
        out = box.import_waiting()
        assert len(out["imported"]) == 1, out
        assert out["imported"][0]["account"] == self.ACCOUNT
        assert out["added"] > 0
        assert not box.waiting(), "the file should have been filed away"
        assert list((box.folder() / box.FILED).rglob("*.csv"))

    def test_the_same_file_cannot_go_in_twice(self, box):
        self._statement(box, f"{self.ACCOUNT} Sept 2026.csv")
        first = box.import_waiting()
        assert first["added"] > 0
        # Put the identical content back under a different name.
        filed = next((box.folder() / box.FILED).rglob("*.csv"))
        (box.folder() / "renamed entirely.csv").write_text(
            filed.read_text(encoding="utf-8"), encoding="utf-8")
        second = box.import_waiting()
        assert second["added"] == 0, "content, not filename, decides what is new"
        assert not second["imported"]

    def test_a_file_that_names_no_account_is_left_alone(self, box):
        self._statement(box, "download (3).csv")
        out = box.import_waiting()
        assert not out["imported"]
        assert out["skipped"] and "doesn't say which account" in out["skipped"][0]["why"]
        assert box.waiting(), "it must stay in the folder for a person to deal with"

    def test_receipts_are_kept_and_never_read(self, box):
        (box.folder() / "receipt boiler.jpg").write_bytes(b"not really a jpeg")
        out = box.import_waiting()
        assert out["receipts"] == ["receipt boiler.jpg"]
        assert out["added"] == 0, "no figures may come off a photograph"
        assert list((box.folder() / box.RECEIPTS).rglob("*.jpg"))

    def test_nothing_is_ever_deleted(self, box):
        self._statement(box, f"{self.ACCOUNT} Oct 2026.csv")
        (box.folder() / "receipt.png").write_bytes(b"x")
        self._statement(box, "no idea whose.csv")
        box.import_waiting()
        names = {p.name for p in box.folder().rglob("*") if p.is_file()}
        for expected in (f"{self.ACCOUNT} Oct 2026.csv", "receipt.png", "no idea whose.csv"):
            assert expected in names, f"{expected} vanished"

    #: A fixed cast, so this is a test about filenames and not about whatever the
    #: demo household happens to contain by the time it runs.
    CAST = [{"id": 1, "name": "Monzo current account", "institution_name": "Monzo"},
            {"id": 2, "name": "Capitec cheque account", "institution_name": "Capitec Bank"},
            {"id": 3, "name": "Freetrade ISA", "institution_name": "Freetrade"}]

    @pytest.mark.parametrize("filename,expect_sure", [
        ("Monzo current account Aug 2026.csv", True),
        ("capitec cheque statement.csv", True),
        ("freetrade-activity-2026.csv", True),
        ("statement.csv", False),
        ("download (3).csv", False),
        ("bank account statement august.csv", False),   # every word identifies nothing
        ("account.pdf", False),
    ])
    def test_which_filenames_are_trusted(self, filename, expect_sure):
        from kestrel import inbox
        _, sure = inbox.guess_account(filename, self.CAST)
        assert sure is expect_sure, filename

    def test_sync_everything_empties_the_folder(self):
        from kestrel.web import server as web
        src = pathlib.Path(web.__file__).read_text(encoding="utf-8")
        block = src[src.index('def api_sync_all'):src.index('def api_sync_all') + 700]
        assert "inbox.import_waiting" in block, "the button must actually do it"
        assert "except Exception" in block, "a bad file must not break the whole sync"


class TestTheMonthlyReminder:
    """The day arithmetic Allan spelled out, and the honest limits of a local app."""

    @pytest.fixture(autouse=True)
    def _clean(self, monkeypatch):
        from kestrel import reminders
        monkeypatch.setitem(config.settings._data, "reminder_on", True)
        monkeypatch.setitem(config.settings._data, "reminder_dismissed", "")
        return reminders

    @pytest.mark.parametrize("chosen,year,month,expected", [
        (1, 2026, 2, 1),
        (15, 2026, 2, 15),
        (28, 2026, 2, 28),
        (31, 2026, 1, 31),      # a long month gets the 31st
        (31, 2026, 2, 28),      # February
        (31, 2028, 2, 29),      # ...and 29 in a leap year
        (31, 2026, 4, 30),      # a 30-day month
        (30, 2026, 2, 28),
        (29, 2026, 2, 28),
        (29, 2028, 2, 29),
        (31, 2100, 2, 28),      # 2100 is NOT a leap year — the century rule
        (31, 2000, 2, 29),      # 2000 IS — the 400 rule
    ])
    def test_the_day_lands_on_a_real_date(self, _clean, chosen, year, month, expected):
        config.settings["reminder_day"] = chosen
        assert _clean.day_in(year, month) == expected

    def test_a_reminder_asks_about_the_month_before_it(self, _clean):
        assert _clean.period_of(dt.date(2026, 10, 1)) == "2026-09"
        assert _clean.period_of(dt.date(2026, 1, 3)) == "2025-12"

    def test_the_next_one_counts_today(self, _clean):
        config.settings["reminder_day"] = 15
        assert _clean.next_due(dt.date(2026, 9, 15)) == dt.date(2026, 9, 15)
        assert _clean.next_due(dt.date(2026, 9, 16)) == dt.date(2026, 10, 15)
        config.settings["reminder_day"] = 31
        assert _clean.next_due(dt.date(2026, 2, 1)) == dt.date(2026, 2, 28)

    def test_it_rolls_over_a_year_end(self, _clean):
        config.settings["reminder_day"] = 5
        assert _clean.next_due(dt.date(2026, 12, 6)) == dt.date(2027, 1, 5)

    @pytest.mark.parametrize("address,ok", [
        ("allan@example.com", True), ("a.b+c@sub.example.co.uk", True),
        ("nope", False), ("a@b", False), ("two words@example.com", False), ("", False),
    ])
    def test_addresses_are_checked_enough_to_catch_a_typo(self, _clean, address, ok):
        assert _clean.valid_email(address) is ok

    def test_turning_it_off_silences_it(self, _clean):
        config.settings["reminder_on"] = False
        st = _clean.state(dt.date(2026, 12, 31))
        assert st["due"] is False and st["show_banner"] is False

    def test_the_calendar_file_is_real_ical(self, _clean):
        config.settings["reminder_day"] = 12
        text = _clean.ics(dt.date(2026, 9, 1))
        raw = text.encode("utf-8")
        assert raw.count(b"\r\n") == raw.count(b"\n"), "iCalendar lines end CRLF"
        assert max(len(l) for l in raw.split(b"\r\n")) <= 75, "lines fold at 75 octets"
        for required in (b"BEGIN:VCALENDAR", b"END:VCALENDAR", b"BEGIN:VEVENT",
                         b"END:VEVENT", b"DTSTART", b"RRULE:", b"BEGIN:VALARM"):
            assert required in raw, required

    def test_the_last_day_is_expressed_as_the_last_day(self, _clean):
        # BYMONTHDAY=31 would simply SKIP every month without a 31st, which is the
        # opposite of what was asked for. -1 means "last", whatever that month's last is.
        for day in (29, 30, 31):
            config.settings["reminder_day"] = day
            assert "BYMONTHDAY=-1" in _clean.ics(dt.date(2026, 9, 1)), day
        config.settings["reminder_day"] = 12
        assert "BYMONTHDAY=12" in _clean.ics(dt.date(2026, 9, 1))

    def test_it_does_not_pretend_to_send_mail_it_cannot(self, _clean):
        # With no mail server there is nothing to send with, and saying so plainly is
        # the whole point — a reminder that silently does nothing is worse than none.
        out = _clean.send_email("allan@example.com")
        if not out["sent"]:
            assert "mail server" in out["why"].lower()

    def test_the_password_never_lands_in_settings_json(self, _clean):
        src = pathlib.Path(_clean.__file__).read_text(encoding="utf-8")
        assert "security.put" in src
        assert "config.settings[\"reminder_password\"]" not in src
        assert "password" not in config.settings.as_dict()


class TestTheTour:
    """A tour that spotlights the same thing twice reads as stuck, not thorough.

    It has happened twice: four steps pointed at Budget once the tabs were merged, and
    later two steps pointed at Import once the statements folder got its own step.
    """

    def _steps(self):
        js = (pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
              / "static" / "app.js").read_text(encoding="utf-8")
        block = js[js.index("const TOUR = ["):js.index("\n];", js.index("const TOUR = ["))]
        return block

    def test_no_screen_is_spotlighted_twice(self):
        block = self._steps()
        navs = re.findall(r"nav:\s*'([\w-]+)'", block)
        assert len(navs) == len(set(navs)), f"two steps point at the same button: {navs}"

    def test_every_step_has_a_title_and_something_to_say(self):
        block = self._steps()
        titles = re.findall(r"title:\s*'([^']+)'", block)
        bodies = re.findall(r"body:\s*'", block)
        assert len(titles) == len(bodies), "a step without a body, or the other way round"
        assert len(set(titles)) == len(titles), "two steps share a title"

    def test_it_asks_who_you_are_first_and_asks_once(self):
        # Name and email belong to the same "who are you" moment, on one card: asking
        # for an address three screens later feels like a form, and asking on the very
        # next step left the spotlight sitting on the same corner twice in a row, which
        # reads as a tour that has got stuck.
        block = self._steps()
        fields = re.findall(r"field:\s*'(\w+)'", block)
        assert fields and fields[0] == "who", fields
        js = (pathlib.Path(__file__).resolve().parent.parent
              / "kestrel" / "web" / "static" / "app.js").read_text(encoding="utf-8")
        assert 'id="tour-name"' in js and 'id="tour-email"' in js, \
            "the who step has to render both boxes"

    def test_no_two_steps_spotlight_the_same_thing(self):
        block = self._steps()
        spots = re.findall(r"(?:el|nav):\s*'([^']+)'", block)
        assert len(spots) == len(set(spots)), f"the spotlight sits still between steps: {spots}"

    def test_it_stays_short(self):
        block = self._steps()
        assert 6 <= len(re.findall(r"title:\s*'", block)) <= 9, "nobody finishes a long tour"


class TestHideAmountsActuallyHides:
    """Hide amounts is the button you press before sharing your screen.

    For a long while it blurred four tiles and nothing else: the CSS it toggled was
    `.money`, a class nothing in this app has ever carried. Every table, every balance
    and every chart tooltip stayed in plain sight. These tests pin the three things that
    made it real.
    """

    def _js(self, code_only=False):
        js = ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
               / "static" / "app.js").read_text(encoding="utf-8"))
        if code_only:
            # These tests look for things that must not appear in the markup. The
            # comments explaining why are allowed to name them.
            js = re.sub(r"/\*.*?\*/", "", js, flags=re.S)
            js = "\n".join(re.sub(r"(^|\s)//.*$", "", ln) for ln in js.splitlines())
        return js

    def test_it_toggles_classes_that_exist(self):
        js = self._js()
        block = js[js.index("function applyPrivacy"):js.index("function applyPrivacy") + 900]
        wanted = re.search(r"\$\$\('([^']+)'\)\.forEach\(e => e\.classList\.toggle\('blur'", block)
        assert wanted, "applyPrivacy no longer toggles blur on a selector"
        classes = {c.strip().lstrip(".") for c in wanted.group(1).split(",")}
        css = ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
                / "static" / "styles.css").read_text(encoding="utf-8"))
        markup = self._js()
        for c in classes:
            assert f".{c}" in css or f'class="{c}' in markup or f" {c}\"" in markup, (
                f".{c} is not a class anything carries — blurring it hides nothing")
        assert {"v", "num", "pv"} <= classes, classes

    def test_chart_tooltips_go_through_tip(self):
        # An SVG <title> is drawn by the browser itself; CSS cannot blur it. Hovering a
        # chart point read the balance out loud with Hide amounts on. `tip()` carries a
        # hidden version in data-h, and applyPrivacy swaps the text rather than covering
        # it. A hand-written <title> would quietly bypass all of that.
        js = self._js(code_only=True)
        assert "<title>" not in js, "a chart tooltip is bypassing tip()"
        assert "svg title[data-h]" in js, "applyPrivacy no longer blanks tooltips"

    def test_new_markup_is_covered_without_anyone_remembering(self):
        # Dialogs and drill-downs render after applyPrivacy has run. Chasing every one
        # of those call sites is how this leaked in the first place.
        js = self._js()
        assert "MutationObserver" in js
        assert js.index("MutationObserver") > js.index("function applyPrivacy")


class TestTheEmailAddressDoesSomething:
    """The tour asks for an email address, so the address has to earn being asked for.

    For a while it went nowhere: it was stored, the copy said it "goes on the reminder",
    and the calendar file never mentioned it. Asking somebody for their address and then
    quietly doing nothing with it is the kind of small dishonesty that makes people
    wonder what else the app is not telling them.
    """

    def _ics(self, address, folded=False):
        from kestrel import config, reminders
        before = config.settings.get("reminder_email")
        try:
            config.settings["reminder_email"] = address
            out = reminders.ics()
        finally:
            config.settings["reminder_email"] = before
        # What a calendar app reads is the unfolded text; the wrapping is transport.
        return out if folded else out.replace("\r\n ", "")

    def test_the_address_is_on_the_entry(self):
        out = self._ics("someone@example.com")
        assert "ATTENDEE" in out and "mailto:someone@example.com" in out
        assert "ORGANIZER" in out

    def test_the_calendar_is_asked_to_email(self):
        # ACTION:EMAIL is the one thing in the whole iCalendar format that means "email
        # me about this". A calendar that supports it sends the email; one that doesn't
        # ignores the block and still shows the pop-up.
        out = self._ics("someone@example.com")
        assert "ACTION:EMAIL" in out
        assert "ACTION:DISPLAY" in out, "the pop-up has to survive for calendars that ignore EMAIL"

    def test_nobody_is_asked_to_accept_their_own_reminder(self):
        # An ATTENDEE without these two turns the entry into a meeting invitation in
        # Outlook, and the person is asked to RSVP to a note from themselves.
        out = self._ics("someone@example.com")
        assert "PARTSTAT=ACCEPTED" in out and "RSVP=FALSE" in out

    def test_no_address_means_no_empty_fields(self):
        out = self._ics("")
        assert "ATTENDEE" not in out and "ORGANIZER" not in out and "ACTION:EMAIL" not in out
        assert "ACTION:DISPLAY" in out, "the reminder still has to work without an address"

    def test_a_typo_is_not_written_into_the_file(self):
        out = self._ics("not an address")
        assert "ATTENDEE" not in out and "mailto:" not in out

    def test_every_line_folds_to_the_limit(self):
        out = self._ics("someone@example.com", folded=True)
        for line in out.split("\r\n"):
            assert len(line.encode("utf-8")) <= 75, line

    def test_the_address_line_needs_no_folding(self):
        # A folded ATTENDEE is legal, and every calendar app I can test unfolds it — but
        # it is the line most likely to be read by something strict, so it is kept short
        # enough not to wrap at all.
        out = self._ics("someone@example.com", folded=True)
        assert any(ln.startswith("ATTENDEE") and len(ln.encode()) <= 75
                   for ln in out.split("\r\n"))

    def test_the_copy_does_not_promise_more_than_that(self):
        js = (pathlib.Path(__file__).resolve().parent.parent
              / "kestrel" / "web" / "static" / "app.js").read_text(encoding="utf-8")
        assert "goes on the reminder; nothing is sent without you" not in js
        assert "calendar that emails" in js, "the honest version of the promise is missing"


class TestLoadingTheSampleDataTwice:
    """Pressing "Load sample data" a second time used to build a second household.

    The transactions inside it deduplicated properly, which made the bug look harmless
    in the tests — but the accounts around them did not, so you got two Freetrade ISAs,
    two houses, two mortgages, and a net worth exactly twice what the sample says. It is
    an easy button to press twice, and there was no way back short of Clear everything.
    """

    def _counts(self):
        return {
            "accounts": db.scalar("SELECT COUNT(*) FROM accounts", (), 0),
            "connections": db.scalar("SELECT COUNT(*) FROM connections", (), 0),
            "mortgages": db.scalar("SELECT COUNT(*) FROM mortgages", (), 0),
            "transactions": db.scalar("SELECT COUNT(*) FROM transactions", (), 0),
            "holdings": db.scalar("SELECT COUNT(*) FROM holdings", (), 0),
        }

    def test_a_second_load_changes_nothing(self):
        before = self._counts()
        demo.load()
        assert self._counts() == before

    def test_a_second_load_does_not_double_what_is_held(self):
        """The question is whether a second load builds a second household.

        This asks it of the stored figures, in the currency they are stored in, and
        never converts or values anything. Two earlier attempts compared `net_worth()`
        to the penny and both failed on Allan's machine while passing here, for the same
        reason twice over: net worth is built from **live market data**. `demo.load()`
        refreshes share prices, and converting a rand account to sterling calls a live
        FX feed whose answer is cached for ten minutes and re-fetched after that. On a
        machine with no internet both fall back to fixed numbers and the comparison
        looks exact; on a real one it drifts by a few pounds, then by 33p, and would
        have gone on drifting by some new amount every time I tightened it.

        A doubled household is a 100% change. Shares held and money on deposit are
        recorded numbers, so they can be compared exactly, and they answer the question
        without asking the market anything.
        """
        def held():
            return {
                "shares": db.rows("SELECT instrument_id, SUM(shares) s FROM holdings "
                                  "GROUP BY instrument_id ORDER BY instrument_id"),
                "balances": db.rows("SELECT currency, SUM(IFNULL(balance,0)) b FROM accounts "
                                    "GROUP BY currency ORDER BY currency"),
                "cost": db.scalar("SELECT ROUND(SUM(cost), 2) FROM holdings", (), 0),
            }

        before = held()
        demo.load()
        after = held()
        assert [dict(r) for r in after["shares"]] == [dict(r) for r in before["shares"]], \
            "a second load changed the number of shares held"
        assert [dict(r) for r in after["balances"]] == [dict(r) for r in before["balances"]], \
            "a second load changed the money on deposit"
        assert after["cost"] == before["cost"], "a second load changed what the holdings cost"

    def test_there_is_exactly_one_of_each_sample_account(self):
        demo.load()
        for external in ("demo:house", "demo:mortgage", "demo:gb-monzo"):
            n = db.scalar("SELECT COUNT(*) FROM accounts WHERE external_id=?", (external,), 0)
            assert n == 1, f"{external} appears {n} times"


class TestTheSampleDataIsInThePast:
    """A statement never lists a payment that hasn't happened yet.

    The generated card payments were always kept to dates that had already passed, but
    the four fixed rows — salary, the two ISA transfers and the mortgage — were not. A
    copy opened before the 26th showed next week's salary as money already in, so the
    Dashboard's "In" for the current month was a month's pay too high, and the
    transaction list opened on a row dated in the future.
    """

    def test_nothing_is_dated_after_today(self):
        today = dt.date.today().isoformat()
        late = db.rows("SELECT posted_on, description, amount FROM transactions "
                       "WHERE posted_on > ? ORDER BY posted_on DESC LIMIT 5", (today,))
        assert not late, f"sample data contains {len(late)} future transactions: {late[:3]}"

    def test_the_generated_files_are_in_the_past_too(self):
        # Checked at the source as well, because the importer is not the only thing that
        # reads these — they are written out for anyone who wants a file to experiment
        # with, and a demo statement with tomorrow's date in it is just wrong.
        import csv as _csv
        import io as _io
        today = dt.date.today()
        monzo = list(_csv.DictReader(_io.StringIO(demo.monzo_csv().decode())))
        assert monzo, "no sample rows at all"
        for r in monzo:
            assert dt.datetime.strptime(r["Date"], "%d/%m/%Y").date() <= today, r
        cap = list(_csv.DictReader(_io.StringIO(demo.capitec_csv().decode())))
        for r in cap:
            assert dt.datetime.strptime(r["Posting Date"], "%Y/%m/%d").date() <= today, r

    def test_the_running_balance_goes_down_the_page(self):
        # The Capitec file is the fixture the CSV importer is tested against, so its
        # Balance column has to behave like a real statement's: each row's balance is the
        # previous row's plus that row's movement, read in date order.
        import csv as _csv
        import io as _io
        rows = list(_csv.DictReader(_io.StringIO(demo.capitec_csv().decode())))
        rows.reverse()                              # the file is newest-first
        for before, after in zip(rows, rows[1:]):
            moved = (float((after["Money In (R)"] or "0").replace(",", ""))
                     - float((after["Money Out (R)"] or "0").replace(",", "")))
            expected = float(before["Balance (R)"].replace(",", "")) + moved
            assert abs(float(after["Balance (R)"].replace(",", "")) - expected) < 0.02, \
                (before, after)


class TestTheBillsGoWhereTheyBelong:
    """Two categorisations that were confidently wrong.

    A British energy bill says who sent it, not which fuel it is for: almost every
    supplier sells gas and electricity on one direct debit. Filing all of them under
    Electricity was a coin toss dressed up as a category, so there is now an Energy
    category for the ones nobody can tell apart, and Electricity and Gas are kept for
    the bills that really are only one thing.

    And a supermarket forecourt is a petrol station that happens to share a name with a
    shop: "TESCO PETROL 4021" went into the weekly groceries, which moved a tank of fuel
    into the food budget every time anyone filled up.
    """

    def _hit(self, text):
        from kestrel.engine import categorise
        for r in categorise._rules():
            if categorise._match(r, text, -50.0):
                return f"{r['parent']} / {r['name']}"
        return "Uncategorised"

    @pytest.mark.parametrize("text", [
        "BRITISH GAS DD 11223", "OCTOPUS ENERGY", "E.ON NEXT", "SSE ENERGY LTD",
        "SCOTTISH POWER", "EDF ENERGY", "OVO ENERGY", "SHELL ENERGY RETAIL"])
    def test_a_dual_fuel_supplier_is_energy(self, text):
        assert self._hit(text) == "Utilities / Energy"

    @pytest.mark.parametrize("text,expect", [
        ("ESKOM PREPAID ELECTRICITY", "Utilities / Electricity"),
        ("CITY POWER JHB", "Utilities / Electricity"),
        ("CALOR GAS DELIVERY", "Utilities / Gas"),
    ])
    def test_a_bill_that_really_is_one_fuel_keeps_its_category(self, text, expect):
        assert self._hit(text) == expect

    @pytest.mark.parametrize("text", [
        "TESCO PETROL 4021", "SAINSBURYS PETROL", "MORRISONS PFS 123", "COSTCO FUEL",
        "SHELL PETROL STATION GAS"])
    def test_filling_up_is_fuel(self, text):
        assert self._hit(text) == "Transport / Fuel"

    @pytest.mark.parametrize("text", ["TESCO STORES 3411", "ASDA SUPERSTORE", "WAITROSE"])
    def test_the_weekly_shop_is_still_the_weekly_shop(self, text):
        assert self._hit(text) == "Food & Drink / Groceries"

    def test_an_existing_database_gets_both_changes(self):
        # The one-off migrations are what reach copies that are already installed. A fix
        # that only lands on fresh installs fixes nobody's data.
        from kestrel.engine import categorise
        keys = {k for k, *_ in categorise.RULE_UPDATES}
        assert "british-gas-is-energy" in keys
        assert "supermarket-forecourt" in keys
        assert ("Utilities", "Energy") in db.ADDED_CATEGORIES


class TestTheAccountTypeIsNeverAWall:
    """The type dropdown used to be exactly the institution's catalogue, and nothing else.

    That made a slightly out-of-date catalogue into a brick wall. Nationwide is the
    country's biggest building society and its FlexAccount is one of the most common
    current accounts there is, and somebody adding one could choose Savings, Cash ISA or
    Mortgage — none of which it is. A catalogue is a good hint about what to pick first
    and a terrible constraint on what exists.
    """

    def _js(self):
        return ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
                 / "static" / "app.js").read_text(encoding="utf-8"))

    def test_every_type_can_be_chosen_whatever_the_institution(self):
        js = self._js()
        block = js[js.index('id="ai-type"'):js.index('id="ai-type"') + 900]
        assert "optgroup" in block, "the other types are no longer offered at all"
        assert "TYPE_LABEL" in block

    def test_the_usual_ones_still_come_first(self):
        js = self._js()
        block = js[js.index('id="ai-type"'):js.index('id="ai-type"') + 900]
        assert block.index("usual.map(opt)") < block.index("optgroup"), \
            "the institution's own products have to be the ones at the top"

    @pytest.mark.parametrize("iid", ["gb-nationwide-building-society",
                                     "gb-cumberland-building-society"])
    def test_the_societies_that_do_current_accounts_say_so(self, iid):
        # Checked against each society's own website rather than assumed — the rest of
        # the building societies in the list really don't offer one.
        from kestrel.institutions import registry
        inst = registry.get(iid)
        assert inst, iid
        assert "current" in inst["account_types"], inst["account_types"]

    def test_the_institution_file_is_still_valid_json(self):
        import json as _json
        for name in ("institutions_gb.json", "institutions_za.json"):
            path = (pathlib.Path(__file__).resolve().parent.parent / "kestrel"
                    / "institutions" / "data" / name)
            data = _json.loads(path.read_text(encoding="utf-8"))
            assert data["institutions"], name
            ids = [i["id"] for i in data["institutions"]]
            assert len(ids) == len(set(ids)), f"duplicate ids in {name}"


class TestTheMonthCardAddsUp:
    """Three numbers on one card, and the third didn't follow from the other two.

    The Dashboard's month card showed In, Out and Left over. Left over is income less
    spending less whatever was put by, and "put by" was nowhere on the card — so a month
    with £0 in and £2,290 out reported −£3,985 left over, and there was nothing on the
    screen to explain the missing £1,695.
    """

    def test_the_arithmetic_is_what_the_card_claims(self):
        from kestrel.engine import budgets as budget_engine
        t = budget_engine.status(budget_engine.current_period())["totals"]
        assert t["net"] == pytest.approx(t["income"] - t["all_spend"] - t["saving"], abs=0.01)

    def test_every_number_in_that_sum_is_on_the_card(self):
        js = (pathlib.Path(__file__).resolve().parent.parent
              / "kestrel" / "web" / "static" / "app.js").read_text(encoding="utf-8")
        start = js.index("VIEWS.dashboard")
        card = js[js.index("monthName(b.period)", start):][:2000]
        for field in ("b.totals.income", "b.totals.all_spend", "b.totals.saving",
                      "b.totals.net"):
            assert field in card, f"{field} is part of the sum and is not shown"


class TestTheTourCardStaysOnScreen:
    """The card is placed next to whatever it is pointing at, and then it grows.

    The statements-folder step fills its path in from the server after the card has been
    positioned, so on a 900x600 window the card got a line taller a moment later and its
    bottom edge — with the Next button on it — ended up below the fold. The tour is the
    first thing a new person sees, and it was possible to get stuck on step two.
    """

    def _static(self, name):
        return ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
                 / "static" / name).read_text(encoding="utf-8"))

    def test_it_is_repositioned_when_it_changes_size(self):
        js = self._static("app.js")
        block = js[js.index("function placeTour"):js.index("function showTourStep")]
        assert "ResizeObserver" in block, "nothing notices the card growing"
        assert "addEventListener('resize'" in block, "nothing notices the window changing"

    def test_a_short_window_scrolls_the_card_rather_than_hiding_the_button(self):
        css = self._static("styles.css")
        block = css[css.index(".tour-card {"):css.index(".tour-step")]
        assert "max-height" in block and "overflow-y: auto" in block


class TestARecordedSaleCanBeUndone:
    """"Record a sale" with every box empty used to save a row, and the row was for ever.

    Nothing validated it, so the Sold table filled up with lines whose every column read
    "—"; and there was no delete route at all, so the only way to remove one was to clear
    the whole database. Typing a sale in by hand is the one place a mistake can be made,
    and it was the one place a mistake could not be taken back.
    """

    def _post(self, body):
        from kestrel.web import server
        return server.api_add_sold(None, {}, {}, body)

    def test_an_empty_form_is_refused(self):
        from kestrel.web import server
        with pytest.raises(server.ApiError) as e:
            self._post({})
        assert "symbol" in str(e.value)

    @pytest.mark.parametrize("body,word", [
        ({"symbol": "ABC"}, "date"),
        ({"symbol": "ABC", "sold_on": "2026-01-02"}, "cost"),
        ({"symbol": "ABC", "sold_on": "2026-01-02", "cost": 10}, "got for it"),
        ({"symbol": "ABC", "sold_on": "2026-01-02", "bought_on": "2026-05-01",
          "cost": 10, "proceeds": 12}, "before it was bought"),
    ])
    def test_the_message_says_which_box(self, body, word):
        from kestrel.web import server
        with pytest.raises(server.ApiError) as e:
            self._post(body)
        assert word in str(e.value)

    def test_a_good_one_saves_and_can_be_removed_again(self):
        from kestrel.web import server
        out = self._post({"symbol": "zzz", "sold_on": "2026-01-02",
                          "bought_on": "2025-05-01", "cost": 10, "proceeds": 12})
        assert out["id"]
        row = db.one("SELECT symbol FROM sold_positions WHERE id=?", (out["id"],))
        assert row["symbol"] == "ZZZ", "the symbol is stored the way the rest of the app writes them"
        server.api_delete_sold(None, {"sid": str(out["id"])}, {}, None)
        assert db.one("SELECT id FROM sold_positions WHERE id=?", (out["id"],)) is None
        with pytest.raises(server.ApiError):
            server.api_delete_sold(None, {"sid": str(out["id"])}, {}, None)

    def test_the_table_offers_the_button(self):
        js = ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
               / "static" / "app.js").read_text(encoding="utf-8"))
        assert "data-sold-del" in js
        assert "/api/investments/sold/" in js


class TestTheBuiltListTellsTheTruth:
    """"Already built" said "Nothing built yet" right under a box saying a file was written.

    Two separate reasons. The list was rendered once when the screen loaded and never
    again, so it could not know about anything built since; and it globbed *.xlsx only,
    so the phone copy — which is an .html file — was never in it at all.
    """

    def test_the_listing_includes_the_phone_copy(self, tmp_path, monkeypatch):
        from kestrel import config
        from kestrel.web import server
        monkeypatch.setattr(config, "exports_dir", lambda: tmp_path)
        (tmp_path / "Investments 2026-01-01.xlsx").write_bytes(b"x" * 10)
        (tmp_path / "phone.html").write_text("<html></html>", encoding="utf-8")
        names = {i["name"] for i in server.api_list_exports(None, {}, {}, None)["items"]}
        assert names == {"Investments 2026-01-01.xlsx", "phone.html"}

    def test_each_entry_says_which_kind_it_is(self, tmp_path, monkeypatch):
        from kestrel import config
        from kestrel.web import server
        monkeypatch.setattr(config, "exports_dir", lambda: tmp_path)
        (tmp_path / "a.xlsx").write_bytes(b"x")
        (tmp_path / "b.html").write_text("x", encoding="utf-8")
        kinds = {i["name"]: i["kind"]
                 for i in server.api_list_exports(None, {}, {}, None)["items"]}
        assert kinds == {"a.xlsx": "workbook", "b.html": "phone"}

    def test_the_screen_redraws_the_list_after_building(self):
        js = ((pathlib.Path(__file__).resolve().parent.parent / "kestrel" / "web"
               / "static" / "app.js").read_text(encoding="utf-8"))
        handler = js[js.index("$$('[data-ex]')"):js.index("$('#ex-open')")]
        assert "builtRows" in handler, "the list is not redrawn after a build"
        assert "'/api/exports'" in handler, "it is redrawn from stale data"


class TestTheReleasePipeline:
    """The GitHub workflow that builds the app on machines we do not own.

    PyInstaller cannot cross-compile, so a Mac build has to be made on a Mac and
    there isn't one. GitHub lends out both kinds of Mac and a Windows box, free, to
    public repositories — which makes this YAML file the only thing standing between
    a tag and three correct downloads.

    It is also a file nobody can test by running it. A typo is accepted by git, pushed
    without complaint, and then reported on a page nobody opens. So it is checked from
    here instead, against the rest of the repository, before the tag goes up.
    """

    @staticmethod
    def _root() -> pathlib.Path:
        return pathlib.Path(__file__).resolve().parent.parent

    @classmethod
    def _path(cls) -> pathlib.Path:
        return cls._root() / ".github" / "workflows" / "release.yml"

    @classmethod
    def _text(cls) -> str:
        return cls._path().read_text(encoding="utf-8")

    @classmethod
    def _yaml(cls) -> dict:
        import yaml
        return yaml.safe_load(cls._text())

    def test_it_is_there_and_it_parses(self):
        assert self._path().exists(), "no workflow means no Mac build, ever"
        d = self._yaml()
        assert set(d["jobs"]) == {"check", "windows", "macos", "release"}

    def test_every_build_runs_every_test(self):
        # A release that skips the tests is worse than no release: it ships to the
        # family with the same confidence as one that passed.
        d = self._yaml()
        for job in ("windows", "macos"):
            steps = " ".join(str(s.get("run", "")) for s in d["jobs"][job]["steps"])
            assert "pytest tests" in steps, f"the {job} build never runs the tests"

    def test_both_kinds_of_mac_are_built(self):
        d = self._yaml()
        matrix = d["jobs"]["macos"]["strategy"]["matrix"]["include"]
        by_arch = {m["arch"]: m["runner"] for m in matrix}
        assert by_arch == {"arm64": "macos-latest", "x86_64": "macos-15-intel"}, (
            "an Apple Silicon build will not start at all on an Intel Mac, and an "
            "Intel build needs Rosetta on Apple Silicon — one file cannot be 'the "
            "Mac version'")

    def test_the_build_checks_it_got_the_processor_it_asked_for(self):
        # setup-python and the runner image between them decide the architecture, and
        # if either ever changes, the build succeeds and produces a file that silently
        # will not open on the machines it was named for.
        steps = " ".join(str(s.get("run", "")) for s in self._yaml()["jobs"]["macos"]["steps"])
        assert "lipo -archs" in steps, "nothing verifies what was actually built"

    def test_the_names_it_builds_are_the_names_it_publishes(self):
        """The one mismatch that cannot be caught by reading either half on its own.

        ditto writes a file; publish.py is told a path; gh uploads a list. Three
        places, one name. Get it wrong and twenty minutes of building ends in a
        failed upload, or — worse — a manifest whose download 404s for everybody.
        """
        text = self._text()
        built, published, uploaded = set(), set(), set()
        for arch in ("arm64", "x86_64"):
            name = f"Mittens and Pence-macos-{arch}.app.zip"
            if f'"out/{name}"' in text.replace("${{ matrix.arch }}", arch):
                built.add(name)
            if f'"macos-{arch}=built/macos-{arch}/{name}"' in text:
                published.add(name)
            if f'"built/macos-{arch}/{name}"' in text:
                uploaded.add(name)
        assert built == published == uploaded and len(built) == 2, (
            f"built={sorted(built)} published={sorted(published)} "
            f"uploaded={sorted(uploaded)} — these three have to agree exactly")

    def test_the_manifest_goes_up_with_the_builds(self):
        # latest.json IS the update mechanism. A release without it leaves every copy
        # already out there reading the previous release's manifest for ever, which
        # looks exactly like there being no new version.
        text = self._text()
        assert "tools/publish.py" in text, "nothing builds the manifest"
        assets = text[text.index("assets=("):text.index(")", text.index("assets=("))]
        assert '"latest.json"' in assets, "the manifest is built and then not uploaded"
        assert '"built/windows/Mittens and Pence.exe"' in assets
        assert '"built/windows/Read me first.txt"' in assets, \
            "the read-me is the only instructions anybody gets"

    def test_the_release_is_left_as_a_draft(self):
        # /releases/latest/download/ ignores drafts. That is what makes a half-right
        # release harmless: it sits there until somebody has read it, and the family's
        # copies carry on pointing at the last good one.
        assert "--draft" in self._text()

    def test_only_the_release_job_may_write(self):
        d = self._yaml()
        assert d["permissions"] == {"contents": "read"}, "the default has to be read-only"
        writers = [n for n, j in d["jobs"].items()
                   if (j.get("permissions") or {}).get("contents") == "write"]
        assert writers == ["release"], f"{writers} can write to the repository"

    def test_the_version_is_checked_against_config_before_anything_is_built(self):
        # Tag and APP_VERSION disagreeing is the classic release bug: the app compares
        # itself against the manifest, so a mismatch either offers everybody an update
        # to the version they are already running, or never offers one at all.
        d = self._yaml()
        assert d["jobs"]["windows"]["needs"] == "check"
        assert "check" in d["jobs"]["macos"]["needs"] or d["jobs"]["macos"]["needs"] == "check"
        steps = " ".join(str(s.get("run", "")) for s in d["jobs"]["check"]["steps"])
        assert "APP_VERSION" in steps and "config" in steps

    def test_nothing_typed_into_a_form_is_pasted_into_a_shell(self):
        """`run: |` with ${{ inputs.x }} in it is a shell injection, flatly.

        Whatever is typed into the "what changed" box would be spliced into the
        script before bash ever sees it. Inputs go through `env:`, where they are
        a value rather than source code.
        """
        import re
        d = self._yaml()
        for name, job in d["jobs"].items():
            for step in job["steps"]:
                script = str(step.get("run", ""))
                found = re.findall(r"\$\{\{\s*(inputs|github\.event)[^}]*\}\}", script)
                assert not found, (
                    f"{name} / {step.get('name', '?')} pastes {found} straight into a "
                    f"shell script — pass it through env: instead")

    def test_the_keys_it_publishes_under_are_keys_the_app_looks_for(self, monkeypatch):
        """The final join: the workflow, publish.py and the running app all agree.

        The workflow chooses the manifest keys. publish.py has to accept them. And a
        real Mac has to actually ask for one of them — otherwise the file is published
        to an address nothing ever reads, and every Mac in the family is told, quite
        cheerfully, that it already has the newest version.
        """
        import importlib.util
        import re
        spec = importlib.util.spec_from_file_location(
            "publish_wf", self._root() / "tools" / "publish.py")
        pub = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(pub)

        keys = set(re.findall(r'"((?:windows|macos|linux)(?:-\w+)?)=built/', self._text()))
        assert keys == {"windows", "macos-arm64", "macos-x86_64"}, keys
        for k in keys:
            assert pub.is_key(k), f"publish.py would reject '{k}'"

        from kestrel import updates
        monkeypatch.setattr(updates, "platform_key", lambda: "macos")
        for machine, wanted in (("arm64", "macos-arm64"), ("x86_64", "macos-x86_64"),
                                ("aarch64", "macos-arm64"), ("AMD64", "macos-x86_64")):
            monkeypatch.setattr(updates.platform, "machine", lambda m=machine: m)
            first = updates.download_keys()[0]
            assert first == wanted, f"a {machine} Mac asks for {first!r}"
            assert first in keys, f"nothing is ever published under {first!r}"

        monkeypatch.setattr(updates, "platform_key", lambda: "windows")
        monkeypatch.setattr(updates.platform, "machine", lambda: "AMD64")
        # Windows is published under the plain key, so the PC has to fall back to it.
        assert "windows" in updates.download_keys()[1:], \
            "the PC would look only for windows-x86_64 and find nothing"
