"""Turn whatever file the bank gave you into a header + list of rows.

Handles CSV/TSV (any common delimiter and encoding), Excel, OFX/QFX and QIF.
Banks put junk above the header often enough that finding the real header row is
its own small problem — `_find_header` does that.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import pathlib
import re
import xml.etree.ElementTree as ET

ENCODINGS = ["utf-8-sig", "utf-8", "cp1252", "latin-1", "utf-16"]


class Table:
    def __init__(self, header: list[str], rows: list[list], source: str = "",
                 preamble: list[list] | None = None):
        self.header = [str(h).strip() if h is not None else "" for h in header]
        self.rows = rows
        self.source = source
        self.preamble = preamble or []

    def dicts(self) -> list[dict]:
        out = []
        for r in self.rows:
            d = {}
            for i, h in enumerate(self.header):
                if h:
                    d[h] = r[i] if i < len(r) else None
            out.append(d)
        return out

    def signature(self) -> str:
        import hashlib
        key = "|".join(re.sub(r"\s+", " ", h.strip().lower()) for h in self.header)
        return hashlib.sha256(key.encode()).hexdigest()[:24]

    def __len__(self):
        return len(self.rows)


def _decode(raw: bytes) -> str:
    # A UTF-16 file decodes "successfully" as cp1252 into mojibake, so look at the
    # byte-order mark and the NUL density before trying the single-byte codecs.
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    head = raw[:2048]
    if head and head.count(b"\x00") > len(head) * 0.2:
        for enc in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    for enc in ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "replace")


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:30])
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except Exception:
        counts = {d: sample.count(d) for d in [",", ";", "\t", "|"]}
        return max(counts, key=counts.get) if max(counts.values()) else ","


HEADER_WORDS = {
    "date", "transaction date", "posted", "value date", "description", "details",
    "narrative", "reference", "amount", "debit", "credit", "money in", "money out",
    "paid in", "paid out", "balance", "type", "category", "merchant", "payee",
    "symbol", "ticker", "quantity", "shares", "price", "currency", "isin", "name",
    "total", "action", "time", "no. of shares", "instrument",
}


def _score_header(cells: list) -> int:
    if not cells:
        return -1
    vals = [str(c or "").strip().lower() for c in cells]
    nonempty = [v for v in vals if v]
    if len(nonempty) < 2:
        return -1
    score = sum(3 for v in nonempty if v in HEADER_WORDS)
    score += sum(1 for v in nonempty if any(w in v for w in HEADER_WORDS))
    # headers rarely look like numbers or dates
    numeric = sum(1 for v in nonempty if re.fullmatch(r"-?[\d.,\s()£$€R]+", v))
    score -= numeric * 2
    score += min(len(nonempty), 8) // 4
    return score


def _find_header(grid: list[list], max_scan: int = 25) -> int:
    best_i, best = 0, -99
    for i, row in enumerate(grid[:max_scan]):
        s = _score_header(row)
        if s > best:
            best, best_i = s, i
    return best_i


def read_csv_bytes(raw: bytes, name: str = "") -> Table:
    text = _decode(raw)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    delim = _sniff_delimiter(text)
    grid = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    grid = [r for r in grid if any(str(c).strip() for c in r)]
    if not grid:
        return Table([], [], name)
    hi = _find_header(grid)
    header = grid[hi]
    width = len(header)
    rows = []
    for r in grid[hi + 1:]:
        if len(r) < width:
            r = r + [None] * (width - len(r))
        rows.append(r[:max(width, len(r))])
    return Table(header, rows, name, preamble=grid[:hi])


def read_excel(path: pathlib.Path, sheet: str | None = None) -> Table:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    grid = []
    for row in ws.iter_rows(values_only=True):
        if row is None:
            continue
        cells = list(row)
        if any(c is not None and str(c).strip() != "" for c in cells):
            grid.append(cells)
    wb.close()
    if not grid:
        return Table([], [], path.name)
    hi = _find_header(grid)
    return Table(grid[hi], grid[hi + 1:], path.name, preamble=grid[:hi])


def excel_sheets(path: pathlib.Path) -> list[str]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


# ---------------------------------------------------------------------------
# OFX / QFX
# ---------------------------------------------------------------------------

def read_ofx(raw: bytes, name: str = "") -> Table:
    text = _decode(raw)
    body = text[text.find("<OFX>"):] if "<OFX>" in text else text
    # OFX 1.x is SGML: close the tags so an XML parser copes.
    if not body.strip().startswith("<?xml"):
        body = re.sub(r"<(\w+)>([^<\r\n]+)", r"<\1>\2</\1>", body)
    body = re.sub(r"&(?!(amp|lt|gt|quot|apos);)", "&amp;", body)
    header = ["date", "description", "memo", "amount", "type", "id", "currency"]
    rows = []
    try:
        root = ET.fromstring(body)
    except ET.ParseError:
        return Table(header, rows, name)

    def txt(node, tag):
        el = node.find(tag)
        return (el.text or "").strip() if el is not None and el.text else ""

    default_ccy = ""
    for cur in root.iter("CURDEF"):
        default_ccy = (cur.text or "").strip()
        break
    for st in root.iter("STMTTRN"):
        d = txt(st, "DTPOSTED")[:8]
        try:
            date = dt.datetime.strptime(d, "%Y%m%d").date().isoformat()
        except Exception:
            date = d
        rows.append([date, txt(st, "NAME"), txt(st, "MEMO"), txt(st, "TRNAMT"),
                     txt(st, "TRNTYPE"), txt(st, "FITID"), default_ccy])
    return Table(header, rows, name)


def read_qif(raw: bytes, name: str = "") -> Table:
    text = _decode(raw)
    header = ["date", "description", "memo", "amount", "category", "cheque"]
    rows, cur = [], {}
    for line in text.splitlines():
        line = line.rstrip()
        if not line:
            continue
        code, val = line[0], line[1:].strip()
        if code == "^":
            if cur:
                rows.append([cur.get("D", ""), cur.get("P", ""), cur.get("M", ""),
                             cur.get("T", cur.get("U", "")), cur.get("L", ""), cur.get("N", "")])
            cur = {}
        elif code in "DPMTULN":
            cur[code] = val
    if cur:
        rows.append([cur.get("D", ""), cur.get("P", ""), cur.get("M", ""),
                     cur.get("T", cur.get("U", "")), cur.get("L", ""), cur.get("N", "")])
    return Table(header, rows, name)


def read_any(path: pathlib.Path, sheet: str | None = None) -> Table:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xltx"):
        return read_excel(path, sheet)
    if suffix == ".pdf":
        from .pdfstatement import read_pdf
        return read_pdf(path, path.name)
    raw = path.read_bytes()
    if suffix in (".ofx", ".qfx"):
        return read_ofx(raw, path.name)
    if suffix == ".qif":
        return read_qif(raw, path.name)
    head = raw[:400].lstrip().upper()
    if raw[:5] == b"%PDF-":
        from .pdfstatement import read_pdf
        return read_pdf(path, path.name)
    if b"OFXHEADER" in head or b"<OFX>" in head:
        return read_ofx(raw, path.name)
    if head.startswith(b"!TYPE"):
        return read_qif(raw, path.name)
    return read_csv_bytes(raw, path.name)


# ---------------------------------------------------------------------------
# Value coercion
# ---------------------------------------------------------------------------

DATE_FORMATS = [
    "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d",
    "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%d/%m/%y", "%m/%d/%y",
    "%Y%m%d", "%d-%b-%Y", "%d-%b-%y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S",
    "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M",
]

_DAY_FIRST_HINT = True  # UK and SA both write 03/04/2026 as 3 April


def parse_date(value, day_first: bool = _DAY_FIRST_HINT) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    s = str(value).strip()
    if not s:
        return None
    s = s.replace("Z", "").strip()
    if "T" in s:
        s = s.split(".")[0]
    fmts = list(DATE_FORMATS)
    if not day_first:
        fmts.insert(0, "%m/%d/%Y")
    for f in fmts:
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            continue
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000 if y < 70 else 1900
        d, mo = (a, b) if day_first else (b, a)
        if mo > 12 and d <= 12:
            d, mo = mo, d
        try:
            return dt.date(y, mo, d).isoformat()
        except ValueError:
            return None
    return None


_NUM_CLEAN = re.compile(r"[^\d\-.,()]")


def parse_amount(value, currency_hint: str | None = None) -> float | None:
    """Copes with '1,234.56', '1.234,56', '(45.00)', '£12.34', 'R 1 234,56', '-'."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s or s in {"-", "--", "n/a", "N/A"}:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.upper().endswith("CR"):
        s = s[:-2].strip()
    elif s.upper().endswith("DR"):
        neg, s = True, s[:-2].strip()
    s = s.replace(" ", " ").replace(" ", "")
    s = _NUM_CLEAN.sub("", s)
    if not s:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):        # 1.234,56
            s = s.replace(".", "").replace(",", ".")
        else:                                   # 1,234.56
            s = s.replace(",", "")
    elif "," in s:
        parts = s.split(",")
        if len(parts[-1]) == 3 and len(parts) > 1 and all(len(p) <= 3 for p in parts[1:]):
            s = s.replace(",", "")              # thousands
        else:
            s = s.replace(",", ".")             # decimal comma
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def parse_text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()
